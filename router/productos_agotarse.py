"""Productos próximos a agotarse: solo productos con buen movimiento (venden casi todos los meses)
cuya existencia disponible (sin bodegas NHE) alcanza para menos de N meses."""
import io
import math
from datetime import datetime, timedelta

from flask import Blueprint, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.db import get_connection
from router.accesos import requiere_acceso
from router.sugerido_compras import a_fecha, buscar_equivalentes, en_lotes
from router.ventas_perdidas import perdidas_por_codigo

productos_agotarse_bp = Blueprint("productos_agotarse", __name__)

HISTORIAL = [3, 6, 12]          # meses para el promedio
COBERTURAS = [1, 2, 3, 4, 6]    # "se agota en menos de N meses"
COMPRAR_PARA = [1, 2, 3, 4, 6]
DIAS_MES = 30.44                # meses móviles hacia atrás desde el último día con ventas
MARCAS_EXCLUIDAS = ["CL"]       # códigos descontinuados (igual que el reporte anterior)
VERDE, NARANJA = "027266", "FE6412"


def entero(valor, permitidos, defecto):
    try:
        v = int(valor)
    except (TypeError, ValueError):
        return defecto
    return v if v in permitidos else defecto


def leer_filtros(args):
    historial = entero(args.get("historial"), HISTORIAL, 6)
    # "vendió en al menos X de los N meses": por defecto 2/3 de los meses (4 de 6)
    meses_venta = entero(args.get("meses_venta"), list(range(1, historial + 1)), max(1, round(historial * 2 / 3)))
    try:
        promedio_min = max(0.0, float(args.get("promedio_min", 1)))
    except ValueError:
        promedio_min = 1.0
    return {"historial": historial, "meses_venta": meses_venta, "promedio_min": promedio_min,
            "cobertura": entero(args.get("cobertura"), COBERTURAS, 3),
            "comprar_para": entero(args.get("comprar_para"), COMPRAR_PARA, 3),
            "marca": (args.get("marca") or "").strip(), "q": (args.get("q") or "").strip(),
            # por defecto las ventas perdidas se suman al sugerido; la casilla desmarcada manda "0"
            "con_perdidas": args.get("con_perdidas", "1") == "1"}


def calcular(f):
    conn = get_connection()
    cursor = conn.cursor()
    try:
        # 1) Ventas diarias por producto (una sola lectura de ventas_por_producto, que no tiene índices).
        #    Se lee un margen extra hacia atrás y el último día con ventas se toma de las mismas filas,
        #    en vez de consultar MAX(Fecha) aparte.
        dias = round(f["historial"] * DIAS_MES)
        desde = datetime.today().date() - timedelta(days=dias + 60)
        cursor.execute("""
            SELECT Codigo, Fecha, SUM(`Cantidad Vendida`) FROM ventas_por_producto
            WHERE Fecha >= %s AND Estatus <> 'Anulado' GROUP BY Codigo, Fecha
        """, (desde.isoformat(),))
        diarias = [(cod, a_fecha(fecha), float(cant or 0)) for cod, fecha, cant in cursor.fetchall()]
        fin = max((d for _, d, _ in diarias), default=None)
        if fin is None:  # sin ventas recientes cargadas: usar el último día registrado
            cursor.execute("SELECT MAX(Fecha) FROM ventas_historicas")
            fin = a_fecha(cursor.fetchone()[0]) or datetime.today().date()
        inicio = fin - timedelta(days=dias - 1)
        if inicio < desde:  # datos con mucho atraso: volver a leer el período completo
            cursor.execute("""
                SELECT Codigo, Fecha, SUM(`Cantidad Vendida`) FROM ventas_por_producto
                WHERE Fecha BETWEEN %s AND %s AND Estatus <> 'Anulado' GROUP BY Codigo, Fecha
            """, (inicio.isoformat(), fin.isoformat()))
            diarias = [(cod, a_fecha(fecha), float(cant or 0)) for cod, fecha, cant in cursor.fetchall()]

        ventas, ultima_venta = {}, {}
        for cod, fecha, cant in diarias:
            if not (inicio <= fecha <= fin):
                continue
            # mes móvil: 0 = los últimos ~30 días, 1 = los anteriores, ...
            mes = min(int((fin - fecha).days // DIAS_MES), f["historial"] - 1)
            ventas.setdefault(cod, [0.0] * f["historial"])[mes] += cant
            if cant > 0 and (cod not in ultima_venta or fecha > ultima_venta[cod]):
                ultima_venta[cod] = fecha

        # 2) Buen movimiento: vendió en suficientes meses y con promedio mínimo
        candidatos = {}
        for cod, meses in ventas.items():
            promedio = sum(meses) / f["historial"]
            con_venta = sum(1 for m in meses if m > 0)
            if con_venta >= f["meses_venta"] and promedio >= f["promedio_min"] and promedio > 0:
                candidatos[cod] = {"meses": list(reversed(meses)), "promedio": promedio, "meses_con_venta": con_venta}
        if not candidatos:
            return [], {"inicio": inicio, "fin": fin}

        # 3) Catálogo (marca, filtros) y existencias sin bodegas NHE (mercadería faltante).
        #    Con miles de candidatos es más rápido leer todo una vez que consultar por lotes IN (...).
        where = [f"COALESCE(p.Marca, '') NOT IN ({', '.join(['%s'] * len(MARCAS_EXCLUIDAS))})", "p.`Codigo Producto` IS NOT NULL"]
        params = list(MARCAS_EXCLUIDAS)
        if f["marca"]:
            where.append("p.Marca = %s"); params.append(f["marca"])
        for palabra in f["q"].split():
            where.append("(p.`Codigo Producto` LIKE %s OR p.Nombre LIKE %s OR p.`Numero Parte` LIKE %s)")
            params += [f"%{palabra}%"] * 3
        cursor.execute(f"""
            SELECT /*+ NO_INDEX(c) */ p.`Codigo Producto`, p.Nombre, p.`Numero Parte`, p.Marca, p.`Ultimo Costo`, c.segmento, c.familia, c.grupo_equival,
                   c.subfamilia
            FROM catalogo_productos p LEFT JOIN clase_prod c ON c.codigo = p.`Codigo Producto`
            WHERE {' AND '.join(where)}
        """, params)
        catalogo = {r[0]: r for r in cursor.fetchall() if r[0] in candidatos}
        codigos = sorted(catalogo)
        cursor.execute("SELECT Codigo_producto, SUM(Cant_Final), SUM(CASE WHEN UPPER(Bodega) LIKE '%NHE%' THEN Cant_Final ELSE 0 END) "
                       "FROM existencias_bodega GROUP BY Codigo_producto")
        # Todas: también sirven para la existencia de los productos equivalentes
        existencias = {r[0]: (float(r[1] or 0), float(r[2] or 0)) for r in cursor.fetchall()}

        # 4) Cobertura y sugerido. Al promedio vendido se le suman (opcional) las ventas perdidas del período.
        perdidas = perdidas_por_codigo(inicio, fin)
        filas = []
        for cod in codigos:
            c = candidatos[cod]
            total, nhe = existencias.get(cod, (0.0, 0.0))
            disponible = total - nhe
            cobertura = max(disponible, 0) / c["promedio"]
            if cobertura >= f["cobertura"]:
                continue
            _, nombre, parte, marca, costo, segmento, familia, grupo, subfamilia = catalogo[cod]
            perdido = perdidas[cod]["unidades"] if cod in perdidas else 0.0
            def sugerir(promedio):
                return max(0, math.ceil(promedio * f["comprar_para"] - max(disponible, 0) - 1e-9))
            sin_perdidas = sugerir(c["promedio"])
            sugerido = sugerir(c["promedio"] + perdido / f["historial"]) if f["con_perdidas"] else sin_perdidas
            filas.append({"codigo": cod, "nombre": nombre, "numero_parte": parte, "marca": marca, "segmento": segmento,
                          "familia": familia, "subfamilia": subfamilia, "grupo": grupo, "costo": float(costo or 0),
                          "existencia": total, "nhe": nhe, "perdidas": perdido, "sugerido_sin_perdidas": sin_perdidas,
                          "disponible": disponible, "promedio": c["promedio"], "meses": c["meses"],
                          "meses_con_venta": c["meses_con_venta"], "cobertura": cobertura, "sugerido": sugerido,
                          "ultima_compra": None, "dias_compra": None, "ultima_venta": None})

        # 5) Última compra (índice por producto en transacciones_inventario)
        resultado = [x["codigo"] for x in filas]
        compras = dict(en_lotes(cursor,
            "SELECT codigo_producto, MAX(fecha) FROM transacciones_inventario WHERE codigo_producto IN ({IN}) "
            "AND tipo_transaccion = 'ENTRADA' AND tipo_documento_origen = 'ORDEN COMPRA' GROUP BY codigo_producto", resultado))

        # 6) Equivalentes (mismo grupo en Clasificación de productos): ¿hay del mismo producto en otra marca?
        vendido = {cod: sum(meses) for cod, meses in ventas.items()}
        equivalentes = buscar_equivalentes(cursor, filas, existencias, vendido, f["historial"])
    finally:
        cursor.close()
        conn.close()

    for x in filas:
        x["ultima_compra"] = a_fecha(compras.get(x["codigo"])) if compras.get(x["codigo"]) else None
        x["dias_compra"] = (fin - x["ultima_compra"]).days if x["ultima_compra"] else None
        x["ultima_venta"] = ultima_venta.get(x["codigo"])  # de las ventas del período (siempre venden en él)
        x["equivalentes"] = equivalentes.get(x["codigo"], [])
        x["eq_disponible"] = sum(max(e["disponible"], 0) for e in x["equivalentes"])
        x["eq_marcas"] = len({e["marca"] for e in x["equivalentes"] if e["marca"] != x["marca"]})
    # Primero lo que se agota antes; a igual cobertura, lo que más se vende
    filas.sort(key=lambda x: (x["cobertura"], -x["promedio"]))
    return filas, {"inicio": inicio, "fin": fin}


def lista_marcas():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT DISTINCT Marca FROM catalogo_productos WHERE Marca IS NOT NULL AND Marca <> '' ORDER BY Marca")
    marcas = [r[0] for r in cursor.fetchall()]
    cursor.close()
    conn.close()
    return marcas


@productos_agotarse_bp.route("/productos_agotarse")
@requiere_acceso("productos_agotarse")
def reporte_productos_agotarse():
    f = leer_filtros(request.args)
    filas, periodo = calcular(f)
    return render_template(
        "reporte_productos_agotarse.html", f=f, filas=filas, periodo=periodo, marcas=lista_marcas(),
        historiales=HISTORIAL, coberturas=COBERTURAS, comprar_para=COMPRAR_PARA,
        agotados=sum(1 for x in filas if x["disponible"] <= 0),
        menos_un_mes=sum(1 for x in filas if 0 < x["disponible"] and x["cobertura"] < 1),
        unidades=sum(x["sugerido"] for x in filas), costo=sum(x["sugerido"] * x["costo"] for x in filas),
        con_equivalente=sum(1 for x in filas if x["eq_disponible"] > 0),
        por_perdidas=sum(x["sugerido"] - x["sugerido_sin_perdidas"] for x in filas),
    )


@productos_agotarse_bp.route("/productos_agotarse/excel")
@requiere_acceso("productos_agotarse")
def descargar_excel():
    f = leer_filtros(request.args)
    filas, periodo = calcular(f)
    wb = Workbook()
    ws = wb.active
    ws.title = "Por agotarse"
    ws["A1"] = "Productos próximos a agotarse"
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    ws["A2"] = (f"Buen movimiento: vendió en al menos {f['meses_venta']} de los últimos {f['historial']} meses "
                f"({periodo['inicio']:%d/%m/%Y} al {periodo['fin']:%d/%m/%Y}) y promedio ≥ {f['promedio_min']:g} unid./mes · "
                f"Cobertura menor a {f['cobertura']} meses · Sugerido para {f['comprar_para']} meses"
                f"{' (incluye ventas perdidas)' if f['con_perdidas'] else ''} · "
                f"Disponible = existencia − bodegas NHE · Generado {datetime.now():%d/%m/%Y %H:%M}")
    ws["A2"].font = Font(size=9, color="52514E")

    meses_cols = [f"Mes -{f['historial'] - i - 1}" if i < f["historial"] - 1 else "Último mes" for i in range(f["historial"])]
    columnas = ([("Código", 11), ("Nº parte", 20), ("Descripción", 36), ("Marca", 14), ("Familia", 20), ("Subfamilia", 22)]
                + [(m, 9) for m in meses_cols]
                + [("Prom. mensual", 11), ("Meses con venta", 9), ("Ventas perdidas (unid.)", 10), ("Existencia", 10),
                   ("En NHE", 8), ("Disponible", 10), ("Cobertura (meses)", 10), ("Sugerido sin ventas perdidas", 12),
                   ("Sugerido", 10), ("Último costo C$", 12), ("Costo sugerido C$", 14), ("Última compra", 12), ("Última venta", 12),
                   ("Grupo equiv.", 11), ("Equivalentes", 11), ("Disponible en equivalentes", 13), ("Detalle equivalentes", 60)])
    col = {nombre: i for i, (nombre, _) in enumerate(columnas, start=1)}
    letra = {nombre: get_column_letter(i) for nombre, i in col.items()}
    ft = 4
    for i, (nombre, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=ft, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NARANJA if nombre in ("Sugerido", "Cobertura (meses)", "Ventas perdidas (unid.)") else VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[ft].height = 42
    resaltado = PatternFill("solid", fgColor="FFF1E8")
    for n, x in enumerate(filas, start=ft + 1):
        valores = ([x["codigo"], x["numero_parte"], x["nombre"], x["marca"], x["familia"], x["subfamilia"]]
                   + [round(m, 2) for m in x["meses"]]
                   + [round(x["promedio"], 2), x["meses_con_venta"], x["perdidas"] or None, x["existencia"], x["nhe"],
                      x["disponible"], round(x["cobertura"], 2), x["sugerido_sin_perdidas"], x["sugerido"], round(x["costo"], 2),
                      f"={letra['Sugerido']}{n}*{letra['Último costo C$']}{n}", x["ultima_compra"], x["ultima_venta"],
                      x["grupo"], len(x["equivalentes"]) or None, x["eq_disponible"] if x["equivalentes"] else None,
                      "; ".join(f"{e['marca']} {e['codigo']} ({e['numero_parte'] or 's/n'}): disp. {e['disponible']:,.0f}"
                                for e in x["equivalentes"]) or None])
        for i, v in enumerate(valores, start=1):
            ws.cell(row=n, column=i, value=v)
        for nombre in ("Último costo C$", "Costo sugerido C$"):
            ws.cell(row=n, column=col[nombre]).number_format = "#,##0.00"
        for nombre in ("Última compra", "Última venta"):
            ws.cell(row=n, column=col[nombre]).number_format = "dd/mm/yyyy"
        ws.cell(row=n, column=col["Sugerido"]).font = Font(bold=True)
        if x["perdidas"]:
            ws.cell(row=n, column=col["Ventas perdidas (unid.)"]).fill = resaltado
    if filas:
        tot = ft + len(filas) + 1
        ws.cell(row=tot, column=1, value="TOTAL").font = Font(bold=True)
        for nombre, fmt in (("Ventas perdidas (unid.)", "#,##0"), ("Sugerido sin ventas perdidas", "#,##0"),
                            ("Sugerido", "#,##0"), ("Costo sugerido C$", "#,##0.00")):
            l = letra[nombre]
            c = ws.cell(row=tot, column=col[nombre], value=f"=SUM({l}{ft + 1}:{l}{tot - 1})")
            c.font = Font(bold=True)
            c.number_format = fmt
    ws.freeze_panes = ws.cell(row=ft + 1, column=4)
    ws.auto_filter.ref = f"A{ft}:{get_column_letter(len(columnas))}{ft + max(len(filas), 1)}"

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    nombre = f"por_agotarse_{f['marca'].replace(' ', '_') + '_' if f['marca'] else ''}{datetime.now():%Y-%m-%d}.xlsx"
    return send_file(salida, as_attachment=True, download_name=nombre,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
