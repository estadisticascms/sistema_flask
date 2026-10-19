"""Ficha del catálogo: productos con existencia a los que les falta medida, aplicación (Descripción Larga) o
equivalencias, con enlaces de búsqueda en la web. La pantalla vive en Clasificación de productos (/clase_prod);
aquí quedan las funciones comunes y la descarga en Excel."""
import io
import re
from datetime import datetime
from urllib.parse import quote_plus

from flask import Blueprint, redirect, request, send_file, url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.db import get_connection
from router.accesos import requiere_acceso

productos_incompletos_bp = Blueprint("productos_incompletos", __name__)

VERDE, NARANJA = "027266", "FE6412"

# Campo de la ficha -> columna del catálogo. "Aplicación" vive en Descripción Larga.
CAMPOS = {"medida": "Medida", "aplicacion": "Descripcion Larga", "equivalencia": "Equivalencias"}
NOMBRES = {"medida": "Medida", "aplicacion": "Aplicación", "equivalencia": "Equivalencia"}
FALTAS = {"cualquiera": "Falta al menos uno", "medida": "Falta medida", "aplicacion": "Falta aplicación",
          "equivalencia": "Falta equivalencia", "todos": "Faltan los tres"}
ORDENES = {
    "valor": "valor DESC, p.Marca, p.Nombre",
    "existencia": "p.`Existencia Actual` DESC, p.Marca, p.Nombre",
    "marca": "p.Marca, p.Nombre",
}

# Sitio del fabricante por marca (solo dominios conocidos; se busca dentro del sitio con Google).
# Para agregar una marca: "MARCA EN CATÁLOGO": "dominio.com"
SITIOS_FABRICANTE = {
    "SKF": "skf.com", "TIMKEN": "timken.com", "GATES": "gates.com", "NSK": "nsk.com",
    "FAG": "schaeffler.com", "INA": "schaeffler.com", "LUK": "schaeffler.com",
    "AUTOMANN": "automann.com", "WORLD AMERICAN": "sandstruck.com", "NEWSTAR": "sandstruck.com",
    "CORTECO": "corteco.com", "FPDIESEL": "fpdiesel.com", "BANDO": "bandousa.com",
    "HALDEX": "haldex.com", "BENDIX": "bendix.com", "MEGADYNE": "megadynegroup.com",
}
# Cómo buscar la marca en la web cuando en el catálogo tiene sufijos (SKF-GENA, COG/MEGADYNE, ...)
ALIAS_MARCA = {"SKF-GENA": "SKF", "SKF-USA": "SKF", "SKF-GRASA": "SKF", "SKF-BANDAS": "SKF",
               "COG/MEGADYNE": "MEGADYNE", "COG": "MEGADYNE"}


def vacio(columna):
    """Condición SQL: el campo está vacío o con un valor que no aporta información."""
    return f"(COALESCE(TRIM(p.`{columna}`), '') IN ('', '-', '.', 'null', 'NULL', 'N/A', 'n/a'))"


def tiene_valor(texto):
    return (texto or "").strip() not in ("", "-", ".", "null", "NULL", "N/A", "n/a")


def marca_busqueda(marca):
    m = (marca or "").strip().upper()
    return ALIAS_MARCA.get(m, m)


def enlaces(numero_parte, marca):
    """Enlaces de búsqueda en la web para un producto."""
    completo = (numero_parte or "").strip()
    if not completo:
        return {}
    # "BB1-4829A/6203-2RSC3" son dos números de parte; "NU 222 ECM/C3" o "6203-2Z/C3" es uno solo con sufijo.
    segmentos = [x.strip() for x in re.split(r"[/;,]", completo) if x.strip()]
    partes = [completo]
    if len(segmentos) > 1 and all(len(x) >= 4 and re.search(r"\d", x) for x in segmentos):
        partes = segmentos
    numero = " OR ".join(f'"{x}"' for x in partes[:3])  # varios números de parte: cualquiera de ellos
    marca_q = marca_busqueda(marca)
    consulta = f"{numero} {marca_q}".strip()
    resultado = {
        "google": "https://www.google.com/search?q=" + quote_plus(consulta),
        "imagenes": "https://www.google.com/search?tbm=isch&q=" + quote_plus(consulta),
    }
    sitio = SITIOS_FABRICANTE.get(marca_q)
    if sitio:
        resultado["fabricante"] = "https://www.google.com/search?q=" + quote_plus(f"site:{sitio} {numero}")
        resultado["sitio"] = sitio
    return resultado


def leer_filtros(args):
    falta = args.get("falta") if args.get("falta") in FALTAS else "cualquiera"
    orden = args.get("orden") if args.get("orden") in ORDENES else "valor"
    try:
        pagina = max(int(args.get("page", 1)), 1)
    except ValueError:
        pagina = 1
    return {"falta": falta, "orden": orden, "pagina": pagina,
            "marca": (args.get("marca") or "").strip(), "q": (args.get("q") or "").strip()}


def condiciones(f):
    where = ["p.`Codigo Producto` IS NOT NULL", "p.`Existencia Actual` > 0"]
    params = []
    vacios = {k: vacio(col) for k, col in CAMPOS.items()}
    if f["falta"] == "cualquiera":
        where.append("(" + " OR ".join(vacios.values()) + ")")
    elif f["falta"] == "todos":
        where.append("(" + " AND ".join(vacios.values()) + ")")
    else:
        where.append(vacios[f["falta"]])
    if f["marca"]:
        where.append("p.Marca = %s")
        params.append(f["marca"])
    for palabra in f["q"].split():
        like = f"%{palabra}%"
        where.append("(p.`Codigo Producto` LIKE %s OR p.Nombre LIKE %s OR p.`Numero Parte` LIKE %s)")
        params += [like] * 3
    return " AND ".join(where), params


SELECT = """
    SELECT /*+ NO_INDEX(c) */ p.`Codigo Producto` AS codigo, p.`Numero Parte` AS numero_parte, p.Nombre AS nombre,
           p.Marca AS marca, c.familia, p.`Existencia Actual` AS existencia, p.`Costo Promedio` AS costo,
           p.`Existencia Actual` * COALESCE(p.`Costo Promedio`, 0) AS valor,
           p.Medida AS medida, p.`Descripcion Larga` AS aplicacion, p.Equivalencias AS equivalencia
    FROM catalogo_productos p
    LEFT JOIN clase_prod c ON c.codigo = p.`Codigo Producto`
"""


def completar(fila):
    fila["faltan"] = [k for k in CAMPOS if not tiene_valor(fila[k])]
    fila["enlaces"] = enlaces(fila["numero_parte"], fila["marca"])
    return fila


@productos_incompletos_bp.route("/productos_incompletos", methods=["GET"])
@requiere_acceso("clase_prod")
def productos_incompletos():
    """La pantalla se unió a Clasificación de productos: se redirige con el filtro equivalente."""
    falta = request.args.get("falta")
    return redirect(url_for("clase_prod.clase_prod", estado="todos", existencia="1",
                            ficha=falta if falta in ("medida", "aplicacion", "equivalencia") else "incompleta",
                            q=" ".join(filter(None, [request.args.get("marca"), request.args.get("q")]))))


@productos_incompletos_bp.route("/productos_incompletos/excel", methods=["GET"])
@requiere_acceso("clase_prod")
def descargar_excel():
    f = leer_filtros(request.args)
    where, params = condiciones(f)
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute(f"{SELECT} WHERE {where} ORDER BY {ORDENES[f['orden']]}", params)
        filas = [completar(r) for r in cursor.fetchall()]
    finally:
        cursor.close()
        conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Productos incompletos"
    ws["A1"] = "Productos con existencia sin ficha completa"
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    filtros = [FALTAS[f["falta"]]] + ([f"Marca: {f['marca']}"] if f["marca"] else []) + ([f"Búsqueda: {f['q']}"] if f["q"] else [])
    ws["A2"] = (" · ".join(filtros) + f" · {len(filas):,} productos · Aplicación = Descripción Larga del catálogo · "
                f"Generado {datetime.now():%d/%m/%Y %H:%M}")
    ws["A2"].font = Font(size=9, color="52514E")

    columnas = [("Código", 11), ("Nº parte", 20), ("Descripción", 36), ("Marca", 14), ("Familia", 20),
                ("Existencia", 10), ("Costo prom. C$", 12), ("Valor inventario C$", 14), ("Falta", 26),
                ("Medida", 18), ("Aplicación", 36), ("Equivalencias", 30),
                ("Buscar en Google", 14), ("Buscar imágenes", 14), ("Sitio del fabricante", 18)]
    ft = 4
    for i, (nombre, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=ft, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NARANJA if nombre == "Falta" else VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[ft].height = 30

    falta_fill = PatternFill("solid", fgColor="FFF1E8")
    enlace = Font(color="0563C1", underline="single")  # un solo objeto de estilo para todos los enlaces
    for n, r in enumerate(filas, start=ft + 1):
        valores = [r["codigo"], r["numero_parte"], r["nombre"], r["marca"], r["familia"],
                   float(r["existencia"] or 0), round(float(r["costo"] or 0), 2), round(float(r["valor"] or 0), 2),
                   ", ".join(NOMBRES[k] for k in r["faltan"]),
                   r["medida"] if tiene_valor(r["medida"]) else None,
                   r["aplicacion"] if tiene_valor(r["aplicacion"]) else None,
                   r["equivalencia"] if tiene_valor(r["equivalencia"]) else None]
        for col, v in enumerate(valores, start=1):
            ws.cell(row=n, column=col, value=v)
        for k, col in (("medida", 10), ("aplicacion", 11), ("equivalencia", 12)):
            if k in r["faltan"]:
                ws.cell(row=n, column=col).fill = falta_fill
        for col, (clave, texto) in enumerate((("google", "Buscar"), ("imagenes", "Imágenes"), ("fabricante", None)), start=13):
            url = r["enlaces"].get(clave)
            if url:
                # fórmula HYPERLINK: Excel la abre igual que un enlace y el archivo se genera mucho más rápido
                etiqueta = (texto or r["enlaces"].get("sitio")).replace('"', "")
                if len(url) <= 250:
                    c = ws.cell(row=n, column=col, value=f'=HYPERLINK("{url}","{etiqueta}")')
                else:  # Excel no acepta textos de más de 255 caracteres dentro de una fórmula
                    c = ws.cell(row=n, column=col, value=etiqueta)
                    c.hyperlink = url
                c.font = enlace
        ws.cell(row=n, column=6).number_format = "#,##0.##"
        ws.cell(row=n, column=7).number_format = "#,##0.00"
        ws.cell(row=n, column=8).number_format = "#,##0.00"
    if filas:
        tot = ft + len(filas) + 1
        ws.cell(row=tot, column=1, value="TOTAL").font = Font(bold=True)
        c = ws.cell(row=tot, column=8, value=f"=SUM(H{ft + 1}:H{tot - 1})")
        c.font = Font(bold=True)
        c.number_format = "#,##0.00"
    ws.freeze_panes = ws.cell(row=ft + 1, column=3)
    ws.auto_filter.ref = f"A{ft}:{get_column_letter(len(columnas))}{ft + max(len(filas), 1)}"

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    nombre = f"productos_sin_ficha_{f['falta']}" + (f"_{f['marca'].replace(' ', '_')}" if f["marca"] else "") + f"_{datetime.now():%Y-%m-%d}.xlsx"
    return send_file(salida, as_attachment=True, download_name=nombre,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
