"""Ventas por departamento, comparativo por mes (reemplaza el reporte de ventas por zona en /reporteZona).

Nivel 1: departamentos × meses. Nivel 2: clientes del departamento × meses con su vendedor.
Nivel 3: productos que compró un cliente (cantidad, monto, vendedor), en una ventana.
Venta neta sin IVA (Sub-Total 2 de ventas_por_producto) sin facturas anuladas; el departamento sale de la
ficha del cliente, normalizado igual que en Ventas por departamento."""
import io
from datetime import date, datetime

from flask import Blueprint, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.db import get_connection
from router.accesos import requiere_acceso
from router.dashboard import ultima_fecha_con_ventas
from router.ventas_bodegas import en_cache
from router.ventas_departamento import (DEPARTAMENTOS, FUERA_DEL_MAPA, OCASIONAL, SIN_DEPTO, cartera_por_departamento, departamento_de,
                                        departamentos_clientes)
from router.ventas_vendedor import rangos_rapidos

ventas_depto_mes_bp = Blueprint("ventas_depto_mes", __name__)

MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
VERDE, NARANJA = "027266", "FE6412"


def clave(nombre):
    return " ".join((nombre or "").upper().split())


def fecha_param(valor, defecto):
    try:
        return datetime.strptime(valor or "", "%Y-%m-%d").date()
    except ValueError:
        return defecto


def restar_meses(d, n):
    m = d.month - 1 - n
    return date(d.year + m // 12, m % 12 + 1, 1)


def leer_filtros(args):
    ultima = ultima_fecha_con_ventas() or date.today()
    if isinstance(ultima, datetime):
        ultima = ultima.date()
    hasta = fecha_param(args.get("hasta"), ultima)
    desde = min(fecha_param(args.get("desde"), restar_meses(hasta, 5)), hasta)   # por defecto: 6 meses
    depto = (args.get("depto") or "").strip().upper()
    return {"desde": desde, "hasta": hasta, "ultima": ultima, "vendedor": (args.get("vendedor") or "").strip(),
            "depto": depto if depto in DEPARTAMENTOS or depto in (SIN_DEPTO.upper(), OCASIONAL.upper()) else "",
            "q": (args.get("q") or "").strip(), "sin_compra": args.get("sin_compra") == "1"}


def meses_entre(desde, hasta):
    meses, a, m = [], desde.year, desde.month
    while (a, m) <= (hasta.year, hasta.month):
        meses.append((f"{a}-{m:02d}", f"{MESES[m - 1]} {str(a)[2:]}"))
        a, m = (a + 1, 1) if m == 12 else (a, m + 1)
    return meses


def nombre_depto(d):
    return d if d in ("RACCN", "RACCS") else d.title()


def fichas_clientes():
    """Clientes de la tabla por nombre: departamento normalizado, vendedor asignado y estado (10 min en memoria)."""
    def leer():
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT nombre, departamento, municipio, vendedor, estado FROM clientes")
            fichas = {}
            for nombre, depto, municipio, vendedor, estado in cursor.fetchall():
                k = clave(nombre)
                if not k:
                    continue
                actual = fichas.get(k)
                if actual is None or (estado == "Activo" and not actual["activo"]):
                    fichas[k] = {"nombre": " ".join(nombre.split()), "depto": departamento_de(depto, municipio) or SIN_DEPTO,
                                 "vendedor": (vendedor or "").strip(), "activo": estado == "Activo"}
            return fichas
        finally:
            cursor.close()
            conn.close()
    return en_cache("vdm_fichas", leer)


def ventas(f):
    """Filas (cliente, vendedor, mes, monto, facturas, última fecha) del período."""
    where, params = ["Fecha BETWEEN %s AND %s", "Estatus <> 'Anulado'"], [f["desde"].isoformat(), f["hasta"].isoformat()]
    if f["vendedor"]:
        where.append("TRIM(Vendedor) = %s"); params.append(f["vendedor"])
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(f"""
            SELECT TRIM(Cliente), TRIM(Vendedor), LEFT(Fecha, 7), SUM(`Sub-Total 2`), COUNT(DISTINCT `No. Factura`), MAX(Fecha)
            FROM ventas_por_producto WHERE {' AND '.join(where)}
            GROUP BY TRIM(Cliente), TRIM(Vendedor), LEFT(Fecha, 7)
        """, params)
        return cursor.fetchall()
    finally:
        cursor.close()
        conn.close()


def lista_vendedores():
    def leer():
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT DISTINCT TRIM(vendedor) FROM ventas_historicas WHERE Fecha >= %s AND vendedor <> '' ORDER BY 1",
                           (date.today().replace(day=1, year=date.today().year - 1),))
            return [r[0] for r in cursor.fetchall() if r[0]]
        finally:
            cursor.close()
            conn.close()
    return en_cache("vdm_vendedores", leer)


def calcular(f):
    meses = meses_entre(f["desde"], f["hasta"])
    deptos = departamentos_clientes()
    cartera = cartera_por_departamento()
    por_depto, por_cliente = {}, {}
    for cliente, vendedor, mes, monto, facturas, ultima in ventas(f):
        monto = float(monto or 0)
        depto = deptos.get(clave(cliente), OCASIONAL)
        d = por_depto.setdefault(depto, {"depto": depto, "meses": {}, "total": 0.0, "clientes": set()})
        d["meses"][mes] = d["meses"].get(mes, 0.0) + monto
        d["total"] += monto
        d["clientes"].add(clave(cliente))
        if f["depto"] and depto.upper() == f["depto"]:
            c = por_cliente.setdefault(clave(cliente), {"cliente": cliente, "meses": {}, "total": 0.0, "facturas": 0,
                                                        "vendedores": {}, "ultima": None})
            c["meses"][mes] = c["meses"].get(mes, 0.0) + monto
            c["total"] += monto
            c["facturas"] += int(facturas or 0)
            if vendedor:
                c["vendedores"][vendedor] = c["vendedores"].get(vendedor, 0.0) + monto
            u = str(ultima)[:10] if ultima else None
            if u and (c["ultima"] is None or u > c["ultima"]):
                c["ultima"] = u

    total = sum(d["total"] for d in por_depto.values())
    fichas = fichas_clientes()
    filas = []
    for nombre in DEPARTAMENTOS + [SIN_DEPTO, OCASIONAL]:
        d = por_depto.get(nombre, {"depto": nombre, "meses": {}, "total": 0.0, "clientes": set()})
        c = cartera.get(nombre, {"total": 0, "activos": 0})
        filas.append({"depto": nombre, "nombre": nombre if nombre in FUERA_DEL_MAPA else nombre_depto(nombre),
                      "meses": d["meses"], "total": d["total"],
                      "compraron": len(d["clientes"]),
                      "cartera": c["total"], "activos": c["activos"], "pct": d["total"] / total * 100 if total else 0})
    filas.sort(key=lambda x: (x["depto"] in FUERA_DEL_MAPA, x["depto"] == OCASIONAL, -x["total"]))
    claves = [m for m, _ in meses]
    for x in filas:   # comparativo: último mes contra el anterior
        ult, ant = (x["meses"].get(claves[-1], 0), x["meses"].get(claves[-2], 0)) if len(claves) > 1 else (0, 0)
        x["variacion"] = (ult - ant) / ant * 100 if ant else None

    clientes = []
    if f["depto"]:
        for k, c in por_cliente.items():
            vend = sorted(c["vendedores"].items(), key=lambda x: -x[1])
            ficha = fichas.get(k, {})
            clientes.append({**c, "vendedor": vend[0][0] if vend else ficha.get("vendedor", ""),
                             "otros_vendedores": [v for v, _ in vend[1:]], "asignado": ficha.get("vendedor", ""),
                             "activo": ficha.get("activo")})
        if f["sin_compra"]:   # clientes del departamento que no compraron en el período
            for k, ficha in fichas.items():
                if ficha["depto"].upper() == f["depto"] and k not in por_cliente and ficha["activo"]:
                    if f["vendedor"] and ficha["vendedor"] != f["vendedor"]:
                        continue
                    clientes.append({"cliente": ficha["nombre"], "meses": {}, "total": 0.0, "facturas": 0, "ultima": None,
                                     "vendedor": ficha["vendedor"], "otros_vendedores": [], "asignado": ficha["vendedor"],
                                     "activo": True, "sin_compra": True})
        for palabra in f["q"].upper().split():
            clientes = [c for c in clientes if palabra in (c["cliente"] or "").upper() or palabra in (c["vendedor"] or "").upper()]
        clientes.sort(key=lambda c: (-c["total"], c["cliente"] or ""))
        total_depto = next((x["total"] for x in filas if x["depto"].upper() == f["depto"]), 0)
        for i, c in enumerate(clientes, start=1):
            c["posicion"] = i
            c["pct"] = c["total"] / total_depto * 100 if total_depto else 0

    por_mes = {m: sum(x["meses"].get(m, 0) for x in filas) for m in claves}
    siguiente = date(f["hasta"].year + f["hasta"].month // 12, f["hasta"].month % 12 + 1, 1)
    parcial = (siguiente - f["hasta"]).days > 1   # el último mes no está completo
    return {"meses": meses, "filas": filas, "total": total, "por_mes": por_mes, "clientes": clientes,
            "parcial": parcial,
            "maximo": max((v for x in filas if x["depto"] not in FUERA_DEL_MAPA for v in x["meses"].values()), default=0),
            "maximo_cliente": max((v for c in clientes for v in c["meses"].values()), default=0),
            "depto_fila": next((x for x in filas if x["depto"].upper() == f["depto"]), None)}


@ventas_depto_mes_bp.route("/reporteZona")
@requiere_acceso("reporteZona")
def reporte():
    f = leer_filtros(request.args)
    d = calcular(f)
    return render_template("ventas_departamento_mes.html", f=f, d=d, vendedores=lista_vendedores(),
                           departamentos=DEPARTAMENTOS, sin_depto=SIN_DEPTO, ocasional=OCASIONAL, fuera_del_mapa=FUERA_DEL_MAPA,
                           rangos=rangos_rapidos(f["ultima"]),
                           nombre_depto=nombre_depto)


@ventas_depto_mes_bp.route("/reporteZona/productos")
@requiere_acceso("reporteZona")
def productos_cliente():
    """Productos que compró un cliente en el período (para la ventana de detalle)."""
    f = leer_filtros(request.args)
    cliente = (request.args.get("cliente") or "").strip()
    if not cliente:
        return jsonify([]), 400
    where = ["Fecha BETWEEN %s AND %s", "Estatus <> 'Anulado'", "TRIM(Cliente) = %s"]
    params = [f["desde"].isoformat(), f["hasta"].isoformat(), cliente]
    if f["vendedor"]:
        where.append("TRIM(Vendedor) = %s"); params.append(f["vendedor"])
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute(f"""
            SELECT Codigo AS codigo, MAX(Numero_parte) AS parte, MAX(`Descrip Prod.`) AS descripcion, MAX(Marca) AS marca,
                   TRIM(Vendedor) AS vendedor, SUM(`Cantidad Vendida`) AS cantidad, SUM(`Sub-Total 2`) AS monto,
                   COUNT(DISTINCT `No. Factura`) AS facturas, MAX(Fecha) AS ultima
            FROM ventas_por_producto WHERE {' AND '.join(where)}
            GROUP BY Codigo, TRIM(Vendedor) ORDER BY monto DESC
        """, params)
        filas = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()
    return jsonify([{**r, "cantidad": float(r["cantidad"] or 0), "monto": float(r["monto"] or 0),
                     "ultima": str(r["ultima"])[:10] if r["ultima"] else ""} for r in filas])


@ventas_depto_mes_bp.route("/reporteZona/excel")
@requiere_acceso("reporteZona")
def descargar_excel():
    f = leer_filtros(request.args)
    d = calcular(f)
    wb = Workbook()
    nota = (f"Del {f['desde']:%d/%m/%Y} al {f['hasta']:%d/%m/%Y} · venta neta sin IVA, sin anuladas"
            f"{' · Vendedor: ' + f['vendedor'] if f['vendedor'] else ''} · Generado {datetime.now():%d/%m/%Y %H:%M}")

    def hoja(ws, titulo, columnas):
        ws["A1"], ws["A2"] = titulo, nota
        ws["A1"].font = Font(bold=True, size=14, color=VERDE)
        ws["A2"].font = Font(size=9, color="52514E")
        for i, (nombre, ancho) in enumerate(columnas, start=1):
            c = ws.cell(row=4, column=i, value=nombre)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=NARANJA if nombre == "Total C$" else VERDE)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = ancho
        ws.freeze_panes = "B5"

    def totales(ws, fila, desde_col, hasta_col):
        ws.cell(row=fila, column=1, value="TOTAL").font = Font(bold=True)
        for i in range(desde_col, hasta_col + 1):
            l = get_column_letter(i)
            c = ws.cell(row=fila, column=i, value=f"=SUM({l}5:{l}{fila - 1})")
            c.font = Font(bold=True)
            c.number_format = "#,##0.00"

    meses = d["meses"]
    ws = wb.active
    ws.title = "Departamentos"
    hoja(ws, "Ventas por departamento y mes", [("Departamento", 20)] + [(e, 13) for _, e in meses]
         + [("Total C$", 15), ("% del total", 9), ("Compraron", 10), ("Clientes en cartera", 10), ("Var. último mes", 10)])
    for n, x in enumerate(d["filas"], start=5):
        valores = ([x["nombre"]] + [round(x["meses"].get(m, 0), 2) or None for m, _ in meses]
                   + [round(x["total"], 2), x["pct"] / 100, x["compraron"], x["cartera"],
                      x["variacion"] / 100 if x["variacion"] is not None else None])
        for i, v in enumerate(valores, start=1):
            c = ws.cell(row=n, column=i, value=v)
            if 1 < i <= len(meses) + 1 or i == len(meses) + 2:
                c.number_format = "#,##0.00"
        ws.cell(row=n, column=len(meses) + 3).number_format = "0.0%"
        ws.cell(row=n, column=len(meses) + 6).number_format = "+0.0%;-0.0%"
    totales(ws, 5 + len(d["filas"]), 2, len(meses) + 2)

    if f["depto"]:
        nombre = d["depto_fila"]["nombre"] if d["depto_fila"] else f["depto"]
        ws2 = wb.create_sheet("Clientes")
        hoja(ws2, f"Clientes de {nombre} por mes", [("Cliente", 40), ("Vendedor", 28), ("Otros vendedores", 26)]
             + [(e, 13) for _, e in meses] + [("Total C$", 15), ("% del depto.", 9), ("Facturas", 9), ("Última compra", 12)])
        for n, c in enumerate(d["clientes"], start=5):
            valores = ([c["cliente"], c["vendedor"], ", ".join(c["otros_vendedores"]) or None]
                       + [round(c["meses"].get(m, 0), 2) or None for m, _ in meses]
                       + [round(c["total"], 2), c["pct"] / 100, c["facturas"],
                          datetime.strptime(c["ultima"], "%Y-%m-%d").date() if c["ultima"] else None])
            for i, v in enumerate(valores, start=1):
                cel = ws2.cell(row=n, column=i, value=v)
                if 3 < i <= len(meses) + 4:
                    cel.number_format = "#,##0.00"
            ws2.cell(row=n, column=len(meses) + 5).number_format = "0.0%"
            ws2.cell(row=n, column=len(meses) + 7).number_format = "dd/mm/yyyy"
        totales(ws2, 5 + len(d["clientes"]), 4, len(meses) + 4)

        # Productos de los clientes del departamento
        deptos = departamentos_clientes()
        where, params = ["Fecha BETWEEN %s AND %s", "Estatus <> 'Anulado'"], [f["desde"].isoformat(), f["hasta"].isoformat()]
        if f["vendedor"]:
            where.append("TRIM(Vendedor) = %s"); params.append(f["vendedor"])
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute(f"""
                SELECT TRIM(Cliente), TRIM(Vendedor), Codigo, MAX(Numero_parte), MAX(`Descrip Prod.`), MAX(Marca),
                       SUM(`Cantidad Vendida`), SUM(`Sub-Total 2`), MAX(Fecha)
                FROM ventas_por_producto WHERE {' AND '.join(where)}
                GROUP BY TRIM(Cliente), TRIM(Vendedor), Codigo
            """, params)
            filas = [r for r in cursor.fetchall() if deptos.get(clave(r[0]), OCASIONAL).upper() == f["depto"]]
        finally:
            cursor.close()
            conn.close()
        filas.sort(key=lambda r: (r[0] or "", -float(r[7] or 0)))
        ws3 = wb.create_sheet("Productos")
        hoja(ws3, f"Productos comprados por los clientes de {nombre}",
             [("Cliente", 38), ("Vendedor", 26), ("Código", 11), ("Nº parte", 18), ("Descripción", 38), ("Marca", 14),
              ("Cantidad", 10), ("Total C$", 14), ("Última compra", 12)])
        for n, r in enumerate(filas, start=5):
            valores = list(r[:6]) + [float(r[6] or 0), round(float(r[7] or 0), 2),
                                     datetime.strptime(str(r[8])[:10], "%Y-%m-%d").date() if r[8] else None]
            for i, v in enumerate(valores, start=1):
                ws3.cell(row=n, column=i, value=v)
            ws3.cell(row=n, column=8).number_format = "#,##0.00"
            ws3.cell(row=n, column=9).number_format = "dd/mm/yyyy"
        ws3.auto_filter.ref = f"A4:I{4 + max(len(filas), 1)}"

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    sufijo = f"_{f['depto'].replace(' ', '_').lower()}" if f["depto"] else ""
    return send_file(salida, as_attachment=True, download_name=f"ventas_departamento_mes{sufijo}_{f['desde']}_al_{f['hasta']}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
