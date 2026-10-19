"""Productos sin movimiento: tienen existencia pero no se venden desde hace N meses o más.
Muestra el último ingreso (compra), la última venta, costos, precios de venta y el margen sobre la venta."""
import io
from datetime import datetime, timedelta

from flask import Blueprint, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.db import get_connection
from router.accesos import requiere_acceso
from router.sugerido_compras import a_fecha
from router.productos_agotarse import lista_marcas
from router.productos_incompletos import tiene_valor

productos_sin_movimiento_bp = Blueprint("productos_sin_movimiento", __name__)

MESES_SIN_VENTA = [6, 9, 12, 18, 24, 36]
DIAS_MES = 30.44
TC_DEFECTO = 36.6243            # los precios del catálogo están en U$; los costos en C$
PRECIOS = ["Precio 1", "Precio 2", "Precio 3", "Precio 4"]
INFO = {"": "Todos", "falta": "Les falta aplicación o equivalencia", "completa": "Con aplicación y equivalencia"}
ORDENES = {"valor": "Mayor valor en inventario", "antiguedad": "Más tiempo sin vender", "existencia": "Mayor existencia"}
VERDE, NARANJA = "027266", "FE6412"
POR_PAGINA = 200   # filas en pantalla; el Excel lleva todas


def leer_filtros(args):
    try:
        meses = int(args.get("meses", 6))
    except ValueError:
        meses = 6
    try:
        tc = float(args.get("tc") or TC_DEFECTO)
        tc = tc if tc > 0 else TC_DEFECTO
    except ValueError:
        tc = TC_DEFECTO
    orden = args.get("orden")
    return {"meses": meses if meses in MESES_SIN_VENTA else 6, "tc": tc,
            "orden": orden if orden in ORDENES else "valor",
            "info": args.get("info") if args.get("info") in INFO else "",
            "marca": (args.get("marca") or "").strip(), "q": (args.get("q") or "").strip()}


def margen(precio, costo):
    """Margen sobre la venta: (precio − costo) / precio."""
    return (precio - costo) / precio if precio > 0 else None


def calcular(f):
    conn = get_connection()
    cursor = conn.cursor()
    try:
        # 1) Existencia por producto (todas las bodegas; aparte lo que está en bodegas NHE)
        cursor.execute("SELECT Codigo_producto, SUM(Cant_Final), SUM(CASE WHEN UPPER(Bodega) LIKE '%NHE%' THEN Cant_Final ELSE 0 END) "
                       "FROM existencias_bodega GROUP BY Codigo_producto HAVING SUM(Cant_Final) > 0")
        existencias = {r[0]: (float(r[1] or 0), float(r[2] or 0)) for r in cursor.fetchall()}

        # 2) Última venta = última salida por factura en transacciones de inventario (desde 2021, más historia
        #    que ventas_por_producto). Solo con el filtro de tipo la consulta usa el índice producto/fecha/tipo
        #    y tarda ~2 s; filtrar además por documento obliga a leer cada fila (~12 s).
        cursor.execute("SELECT /*+ INDEX(transacciones_inventario idx_transacciones_producto_fecha) */ "
                       "codigo_producto, MAX(fecha) FROM transacciones_inventario "
                       "WHERE tipo_transaccion = 'SALIDA' GROUP BY codigo_producto")
        salidas = [(cod, a_fecha(fecha)) for cod, fecha in cursor.fetchall() if fecha]
        ultima_venta = {cod: fecha for cod, fecha in salidas if cod in existencias}

        # Se mide contra el último día cargado (no contra hoy, por si la carga va atrasada)
        fin = max((fecha for _, fecha in salidas), default=None) or datetime.today().date()
        corte = fin - timedelta(days=round(f["meses"] * DIAS_MES))
        candidatos = {c for c in existencias if c not in ultima_venta or ultima_venta[c] <= corte}

        # 3) Catálogo: descripción, marca, costos y precios
        where, params = ["p.`Codigo Producto` IS NOT NULL"], []
        if f["marca"]:
            where.append("p.Marca = %s"); params.append(f["marca"])
        for palabra in f["q"].split():
            where.append("(p.`Codigo Producto` LIKE %s OR p.Nombre LIKE %s OR p.`Numero Parte` LIKE %s)")
            params += [f"%{palabra}%"] * 3
        cursor.execute(f"""
            SELECT /*+ NO_INDEX(c) */ p.`Codigo Producto`, p.Nombre, p.`Numero Parte`, p.Marca, p.Moneda,
                   p.`Costo Promedio`, p.`Ultimo Costo`, p.`Precio 1`, p.`Precio 2`, p.`Precio 3`, p.`Precio 4`, c.familia,
                   p.`Descripcion Larga`, p.Equivalencias
            FROM catalogo_productos p LEFT JOIN clase_prod c ON c.codigo = p.`Codigo Producto`
            WHERE {' AND '.join(where)}
        """, params)
        catalogo, con_venta = {}, {"productos": 0, "faltan": 0}
        for r in cursor.fetchall():
            if r[0] in candidatos:
                catalogo[r[0]] = r
            elif r[0] in existencias:  # con existencia y con ventas recientes: sirve de comparación
                con_venta["productos"] += 1
                con_venta["faltan"] += not (tiene_valor(r[12]) and tiene_valor(r[13]))

        # 4) Último ingreso por compra: fecha y cantidad (suma de ese día)
        cursor.execute("SELECT codigo_producto, fecha, cantidad FROM transacciones_inventario "
                       "WHERE tipo_transaccion = 'ENTRADA' AND tipo_documento_origen = 'ORDEN COMPRA'")
        ingresos = {}
        for cod, fecha, cant in cursor.fetchall():
            if cod not in catalogo or not fecha:
                continue
            fecha = a_fecha(fecha)
            actual = ingresos.get(cod)
            if actual is None or fecha > actual[0]:
                ingresos[cod] = (fecha, float(cant or 0))
            elif fecha == actual[0]:
                ingresos[cod] = (fecha, actual[1] + float(cant or 0))
    finally:
        cursor.close()
        conn.close()

    filas = []
    for cod, r in catalogo.items():
        _, nombre, parte, marca, moneda, costo_prom, ultimo_costo, p1, p2, p3, p4, familia, aplicacion, equivalencias = r
        precios = [p1, p2, p3, p4]
        aplicacion = aplicacion.strip() if tiene_valor(aplicacion) else ""
        equivalencias = equivalencias.strip() if tiene_valor(equivalencias) else ""
        if (f["info"] == "falta" and aplicacion and equivalencias) or (f["info"] == "completa" and not (aplicacion and equivalencias)):
            continue
        total, nhe = existencias[cod]
        costo_prom, ultimo_costo = float(costo_prom or 0), float(ultimo_costo or 0)
        costo = ultimo_costo or costo_prom   # el margen se calcula con el último costo
        factor = f["tc"] if (moneda or "U$") != "C$" else 1.0
        lista = []
        for nombre_precio, p in zip(PRECIOS, precios):
            p_cs = float(p or 0) * factor
            lista.append({"nombre": nombre_precio, "usd": float(p or 0) if factor != 1 else None,
                          "cs": p_cs, "margen": margen(p_cs, costo) if p_cs > 0 else None})
        uv = ultima_venta.get(cod)
        ingreso = ingresos.get(cod)
        filas.append({
            "codigo": cod, "nombre": nombre, "numero_parte": parte, "marca": marca, "familia": familia,
            "aplicacion": aplicacion, "equivalencias": equivalencias,
            "existencia": total, "nhe": nhe, "costo_promedio": costo_prom, "ultimo_costo": ultimo_costo,
            "valor": total * (costo_prom or ultimo_costo), "precios": lista,
            "ultima_venta": uv, "meses_sin_venta": (fin - uv).days / DIAS_MES if uv else None,
            "ultimo_ingreso": ingreso[0] if ingreso else None, "cantidad_ingreso": ingreso[1] if ingreso else None,
            "meses_ingreso": (fin - ingreso[0]).days / DIAS_MES if ingreso else None,
        })

    if f["orden"] == "antiguedad":   # sin venta registrada primero, luego la venta más vieja
        filas.sort(key=lambda x: (x["ultima_venta"] is not None, x["ultima_venta"] or fin, -x["valor"]))
    elif f["orden"] == "existencia":
        filas.sort(key=lambda x: -x["existencia"])
    else:
        filas.sort(key=lambda x: -x["valor"])
    return filas, {"fin": fin, "corte": corte, "con_venta": con_venta}


def estado_info(x):
    if x["aplicacion"] and x["equivalencias"]:
        return "Completa"
    if not x["aplicacion"] and not x["equivalencias"]:
        return "Falta ambas"
    return "Falta aplicación" if not x["aplicacion"] else "Falta equivalencia"


def resumen(filas, meses):
    rangos = [(f"{meses} a 12 meses", 0, 12), (f"{max(meses, 12)} a 24 meses", 12, 24), (f"más de {max(meses, 24)} meses", 24, 10 ** 6)]
    grupos = []
    for etiqueta, desde, hasta in rangos:
        if hasta <= meses:   # rango que el filtro ya dejó fuera
            continue
        sel = [x for x in filas if x["meses_sin_venta"] is not None and desde <= x["meses_sin_venta"] < hasta]
        grupos.append({"etiqueta": etiqueta, "productos": len(sel), "valor": sum(x["valor"] for x in sel)})
    sin = [x for x in filas if x["ultima_venta"] is None]
    grupos.append({"etiqueta": "sin venta registrada", "productos": len(sin), "valor": sum(x["valor"] for x in sin)})
    falta_apl = sum(1 for x in filas if not x["aplicacion"])
    falta_eq = sum(1 for x in filas if not x["equivalencias"])
    faltan = sum(1 for x in filas if not (x["aplicacion"] and x["equivalencias"]))
    return {"falta_aplicacion": falta_apl, "falta_equivalencias": falta_eq, "faltan": faltan,
            "productos": len(filas), "unidades": sum(x["existencia"] for x in filas),
            "valor": sum(x["valor"] for x in filas), "grupos": grupos}


@productos_sin_movimiento_bp.route("/productos_sin_movimiento")
@requiere_acceso("productos_sin_movimiento")
def reporte():
    f = leer_filtros(request.args)
    filas, periodo = calcular(f)
    paginas = max(1, -(-len(filas) // POR_PAGINA))
    try:
        pagina = min(max(1, int(request.args.get("pagina", 1))), paginas)
    except ValueError:
        pagina = 1
    return render_template("productos_sin_movimiento.html", f=f, periodo=periodo, r=resumen(filas, f["meses"]),
                           filas=filas[(pagina - 1) * POR_PAGINA: pagina * POR_PAGINA], pagina=pagina, paginas=paginas,
                           desde=(pagina - 1) * POR_PAGINA, marcas=lista_marcas(), opciones_meses=MESES_SIN_VENTA,
                           ordenes=ORDENES, precios=PRECIOS, info=INFO)


@productos_sin_movimiento_bp.route("/productos_sin_movimiento/excel")
@requiere_acceso("productos_sin_movimiento")
def descargar_excel():
    f = leer_filtros(request.args)
    filas, periodo = calcular(f)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sin movimiento"
    ws["A1"] = f"Productos con existencia sin ventas en {f['meses']} meses o más"
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    ws["A2"] = (f"Última venta anterior al {periodo['corte']:%d/%m/%Y} (datos de ventas hasta {periodo['fin']:%d/%m/%Y}) · "
                f"Precios del catálogo en U$ convertidos a C$ con TC {f['tc']:g} · Margen = (precio − último costo) / precio · "
                f"Valor = existencia × costo promedio · Generado {datetime.now():%d/%m/%Y %H:%M}")
    ws["A2"].font = Font(size=9, color="52514E")

    columnas = [("Código", 11), ("Nº parte", 20), ("Descripción", 36), ("Marca", 14), ("Familia", 20),
                ("Información", 16), ("Aplicación", 40), ("Equivalencias", 30),
                ("Existencia", 10), ("En NHE", 8), ("Costo promedio C$", 13), ("Último costo C$", 13), ("Valor inventario C$", 15),
                ("Último ingreso", 12), ("Cant. último ingreso", 11), ("Última venta", 12), ("Meses sin venta", 10)]
    for p in PRECIOS:
        columnas += [(f"{p} U$", 10), (f"{p} C$", 11), (f"Margen {p}", 10)]
    ft = 4
    for i, (nombre, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=ft, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NARANJA if nombre.startswith("Margen") or nombre == "Meses sin venta" else VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[ft].height = 32

    tc_celda = "$B$3"
    ws["A3"], ws["B3"] = "Tipo de cambio", f["tc"]
    ws["A3"].font = Font(bold=True)
    for n, x in enumerate(filas, start=ft + 1):
        valores = [x["codigo"], x["numero_parte"], x["nombre"], x["marca"], x["familia"], estado_info(x),
                   x["aplicacion"][:500], x["equivalencias"][:500], x["existencia"], x["nhe"],
                   round(x["costo_promedio"], 4), round(x["ultimo_costo"], 4), round(x["valor"], 2),
                   x["ultimo_ingreso"], x["cantidad_ingreso"], x["ultima_venta"],
                   round(x["meses_sin_venta"], 1) if x["meses_sin_venta"] is not None else "Sin venta"]
        for col, v in enumerate(valores, start=1):
            ws.cell(row=n, column=col, value=v)
        for col in (11, 12, 13):
            ws.cell(row=n, column=col).number_format = "#,##0.00"
        for col in (14, 16):
            ws.cell(row=n, column=col).number_format = "dd/mm/yyyy"
        if x["aplicacion"] == "" or x["equivalencias"] == "":
            ws.cell(row=n, column=6).font = Font(bold=True, color="B42318")
        costo = f"IF(L{n}>0,L{n},K{n})"
        col = len(valores) + 1
        for p in x["precios"]:
            usd, cs, mg = (get_column_letter(col + k) for k in range(3))
            if p["usd"] is not None:   # precio en U$: C$ y margen como fórmulas (cambian si se edita el TC)
                ws.cell(row=n, column=col, value=round(p["usd"], 4)).number_format = "#,##0.00"
                ws.cell(row=n, column=col + 1, value=f"={usd}{n}*{tc_celda}")
            else:
                ws.cell(row=n, column=col + 1, value=round(p["cs"], 2))
            ws.cell(row=n, column=col + 1).number_format = "#,##0.00"
            ws.cell(row=n, column=col + 2, value=f'=IF({cs}{n}>0,({cs}{n}-{costo})/{cs}{n},"")').number_format = "0.0%"
            col += 3
    if filas:
        tot = ft + len(filas) + 1
        ws.cell(row=tot, column=1, value="TOTAL").font = Font(bold=True)
        for col, fmt in ((9, "#,##0"), (13, "#,##0.00")):
            letra = get_column_letter(col)
            c = ws.cell(row=tot, column=col, value=f"=SUM({letra}{ft + 1}:{letra}{tot - 1})")
            c.font = Font(bold=True)
            c.number_format = fmt
    ws.freeze_panes = ws.cell(row=ft + 1, column=4)
    ws.auto_filter.ref = f"A{ft}:{get_column_letter(len(columnas))}{ft + max(len(filas), 1)}"

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    nombre = f"sin_movimiento_{f['meses']}m_{f['marca'].replace(' ', '_') + '_' if f['marca'] else ''}{datetime.now():%Y-%m-%d}.xlsx"
    return send_file(salida, as_attachment=True, download_name=nombre,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
