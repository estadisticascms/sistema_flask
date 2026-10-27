"""Productos con existencia y poco movimiento, para armar estrategias de marketing, y seguimiento de campañas.

- Poco movimiento: cobertura = existencia / venta mensual promedio de los últimos 12 meses (salidas por factura menos
  devoluciones). Se listan los que tienen N meses o más de cobertura y los que no se vendieron en 12 meses.
- Revisión de precio: margen del precio de referencia sobre el último costo y el descuento máximo que deja el margen objetivo.
- Sugerencia de estrategia según margen, cobertura y si la ficha (aplicación / equivalencias) está completa.
- Campañas (mkt_campanas / mkt_campana_productos): al agregar productos se guarda la foto inicial (existencia, costo, precio y
  venta mensual de antes); el seguimiento compara lo vendido desde el inicio contra ese ritmo anterior."""
import io
import re
from datetime import date, datetime, timedelta

from flask import Blueprint, flash, jsonify, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.db import get_connection
from router.accesos import requiere_acceso
from router.productos_agotarse import lista_marcas
from router.productos_incompletos import tiene_valor
from router.productos_sin_movimiento import DIAS_MES, PRECIOS, TC_DEFECTO
from router.sugerido_compras import a_fecha
from router.ventas_bodegas import en_cache

estrategia_bp = Blueprint("estrategia_productos", __name__)

HINT = "/*+ INDEX(transacciones_inventario idx_transacciones_producto_fecha) */"
COBERTURAS = [6, 12, 18, 24, 36]
CLASES = {"": "Todos", "sin_venta": "Sin venta en 12 meses", "muy_lento": "Muy lento (24+ meses de existencia)", "lento": "Lento"}
ORDENES = {"valor": "Mayor valor en inventario", "cobertura": "Más meses de existencia", "margen": "Mayor margen", "descuento": "Mayor descuento posible"}
ESTRATEGIAS = ["Descuento / promoción", "Combo con producto de alta rotación", "Liquidación", "Completar ficha y difundir",
               "Ofrecer a clientes que lo compraron", "Exhibición / catálogo / redes", "Otra"]
VERDE, NARANJA = "027266", "FE6412"
POR_PAGINA = 200
DUDOSO_MIN, DUDOSO_MAX = -1.0, 0.95   # márgenes fuera de este rango suelen ser error de costo o precio


def usuario():
    return session.get("usuario") or ""


def numero(valor, defecto, minimo, maximo):
    try:
        v = float(valor)
    except (TypeError, ValueError):
        return defecto
    return v if minimo <= v <= maximo else defecto


def leer_filtros(args):
    cobertura = int(numero(args.get("cobertura"), 12, 1, 120))
    return {"cobertura": cobertura if cobertura in COBERTURAS else 12,
            "clase": args.get("clase") if args.get("clase") in CLASES else "",
            "orden": args.get("orden") if args.get("orden") in ORDENES else "valor",
            "marca": (args.get("marca") or "").strip(), "familia": (args.get("familia") or "").strip(),
            "q": (args.get("q") or "").strip(), "info": args.get("info") if args.get("info") in ("", "falta", "completa") else "",
            "valor_min": numero(args.get("valor_min"), 0, 0, 10 ** 9),
            "precio": int(numero(args.get("precio"), 1, 1, 4)), "tc": numero(args.get("tc"), TC_DEFECTO, 1, 1000),
            "margen_obj": numero(args.get("margen_obj"), 25, 0, 90),
            "sin_recientes": args.get("sin_recientes") == "1"}


# ---------------------------------------------------------------- datos base (con caché)
def datos_base():
    """Existencia, ventas netas de 12 y 3 meses, clientes, última venta y último ingreso por producto."""
    def leer():
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT Codigo_producto, SUM(Cant_Final) FROM existencias_bodega GROUP BY Codigo_producto HAVING SUM(Cant_Final) > 0")
            existencias = {r[0]: float(r[1] or 0) for r in cursor.fetchall()}
            cursor.execute(f"SELECT {HINT} codigo_producto, MAX(fecha) FROM transacciones_inventario WHERE tipo_transaccion = 'SALIDA' GROUP BY codigo_producto")
            ultima_venta = {c: a_fecha(f) for c, f in cursor.fetchall() if f}
            fin = max(ultima_venta.values(), default=None) or date.today()
            desde12, desde3 = fin - timedelta(days=365), fin - timedelta(days=round(3 * DIAS_MES))
            cursor.execute(f"""
                SELECT {HINT} codigo_producto,
                       SUM(CASE WHEN tipo_transaccion = 'SALIDA' THEN cantidad ELSE -cantidad END),
                       SUM(CASE WHEN fecha >= %s THEN (CASE WHEN tipo_transaccion = 'SALIDA' THEN cantidad ELSE -cantidad END) ELSE 0 END),
                       COUNT(DISTINCT CASE WHEN tipo_transaccion = 'SALIDA' THEN cliente_proveedor END)
                FROM transacciones_inventario
                WHERE fecha > %s AND (tipo_transaccion = 'SALIDA' OR (tipo_transaccion = 'ENTRADA' AND tipo_documento_origen = 'DEVOLUCION'))
                GROUP BY codigo_producto
            """, (desde3, desde12))
            ventas = {r[0]: (max(float(r[1] or 0), 0.0), max(float(r[2] or 0), 0.0), int(r[3] or 0)) for r in cursor.fetchall()}
            cursor.execute("SELECT codigo_producto, MAX(fecha) FROM transacciones_inventario "
                           "WHERE tipo_transaccion = 'ENTRADA' AND tipo_documento_origen = 'ORDEN COMPRA' GROUP BY codigo_producto")
            ingreso = {c: a_fecha(f) for c, f in cursor.fetchall() if f}
            return {"existencias": existencias, "ultima_venta": ultima_venta, "ventas": ventas, "ingreso": ingreso, "fin": fin}
        finally:
            cursor.close()
            conn.close()
    return en_cache("estrategia_datos_base", leer, segundos=600)


def catalogo(codigos=None):
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("""
            SELECT /*+ NO_INDEX(c) */ p.`Codigo Producto` AS codigo, p.Nombre AS nombre, p.`Numero Parte` AS parte, p.Marca AS marca,
                   p.Moneda AS moneda, p.`Costo Promedio` AS costo_prom, p.`Ultimo Costo` AS ultimo_costo,
                   p.`Precio 1` AS p1, p.`Precio 2` AS p2, p.`Precio 3` AS p3, p.`Precio 4` AS p4,
                   p.`Descripcion Larga` AS aplicacion, p.Equivalencias AS equivalencias, c.segmento, c.familia, c.subfamilia
            FROM catalogo_productos p LEFT JOIN clase_prod c ON c.codigo = p.`Codigo Producto`
        """ + (f" WHERE p.`Codigo Producto` IN ({', '.join(['%s'] * len(codigos))})" if codigos else ""), tuple(codigos or ()))
        return {r["codigo"]: r for r in cursor.fetchall()}
    finally:
        cursor.close()
        conn.close()


def en_campanas():
    """Código -> nombre de la campaña activa en la que está."""
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT p.codigo, c.nombre FROM mkt_campana_productos p JOIN mkt_campanas c ON c.id = p.campana_id WHERE c.estado = 'Activa'")
        return {cod: nombre for cod, nombre in cursor.fetchall()}
    finally:
        cursor.close()
        conn.close()


def analizar(cod, cat, base, f):
    """Arma la fila de un producto: rotación, precio, margen, descuento posible y sugerencias."""
    existencia = base["existencias"].get(cod, 0.0)
    v12, v3, clientes = base["ventas"].get(cod, (0.0, 0.0, 0))
    mensual = v12 / 12
    cobertura = existencia / mensual if mensual > 0 else None
    clase = "sin_venta" if mensual <= 0 else ("muy_lento" if cobertura >= 24 else "lento")
    costo_prom, ultimo_costo = float(cat["costo_prom"] or 0), float(cat["ultimo_costo"] or 0)
    costo = ultimo_costo or costo_prom
    factor = f["tc"] if (cat["moneda"] or "U$") != "C$" else 1.0
    precio_usd = float(cat[f"p{f['precio']}"] or 0)
    precio = precio_usd * factor
    margen = (precio - costo) / precio if precio > 0 else None
    dudoso = margen is not None and not (DUDOSO_MIN <= margen <= DUDOSO_MAX)
    m_obj = f["margen_obj"] / 100
    # Descuento máximo que deja el margen objetivo: precio nuevo = costo / (1 − margen objetivo)
    descuento = max(0.0, 1 - (costo / (1 - m_obj)) / precio) if precio > 0 and costo > 0 and not dudoso else 0.0
    aplicacion = cat["aplicacion"].strip() if tiene_valor(cat["aplicacion"]) else ""
    equivalencias = cat["equivalencias"].strip() if tiene_valor(cat["equivalencias"]) else ""
    uv = base["ultima_venta"].get(cod)

    sug = []
    if dudoso:
        sug.append(("revisar", "Revisar costo o precio en el catálogo"))
    elif margen is not None and margen < m_obj:
        sug.append(("margen", "Margen bajo: no admite descuento; combo o revisar costo"))
    elif descuento >= 0.05:
        sug.append(("descuento", f"Promoción con hasta {descuento * 100:.0f}% de descuento"))
    if not (aplicacion and equivalencias):
        sug.append(("ficha", "Completar aplicación / equivalencias y difundir"))
    if clase == "sin_venta" and uv:
        sug.append(("clientes", "Ofrecer a clientes que lo compraron antes"))
    elif clase == "sin_venta":
        sug.append(("liquidar", "Nunca se ha vendido: exhibir, catálogo o liquidar"))
    if clase != "sin_venta" and cobertura and cobertura >= 36 and descuento >= 0.05:
        sug.append(("liquidar", "Más de 3 años de existencia: considerar liquidación"))
    return {
        "codigo": cod, "nombre": cat["nombre"] or "", "parte": cat["parte"] or "", "marca": (cat["marca"] or "").strip(),
        "segmento": cat["segmento"] or "", "familia": cat["familia"] or "", "subfamilia": cat["subfamilia"] or "",
        "aplicacion": aplicacion, "equivalencias": equivalencias, "existencia": existencia,
        "costo": costo, "costo_prom": costo_prom, "valor": existencia * (costo_prom or ultimo_costo),
        "precio": precio, "precio_usd": precio_usd if factor != 1 else None, "margen": margen, "dudoso": dudoso, "descuento": descuento,
        "precio_desc": costo / (1 - m_obj) if descuento > 0 else None,
        "v12": v12, "v3": v3, "mensual": mensual, "clientes": clientes, "cobertura": cobertura, "clase": clase,
        "ultima_venta": uv, "meses_sin_venta": (base["fin"] - uv).days / DIAS_MES if uv else None,
        "ultimo_ingreso": base["ingreso"].get(cod), "sugerencias": sug,
        # Ingresó hace menos de 3 meses: puede ser lento solo porque es compra nueva
        "reciente": bool(base["ingreso"].get(cod) and (base["fin"] - base["ingreso"][cod]).days <= 92),
    }


def calcular(f):
    base = datos_base()
    cat = catalogo()
    activos = en_campanas()
    filas = []
    marca = f["marca"].upper()
    for cod, existencia in base["existencias"].items():
        c = cat.get(cod)
        if not c:
            continue
        v12 = base["ventas"].get(cod, (0, 0, 0))[0]
        if v12 > 0 and existencia / (v12 / 12) < f["cobertura"]:
            continue   # rota lo suficiente
        m = (c["marca"] or "").strip().upper()
        if marca and m != marca and not m.startswith(marca + "-"):
            continue
        if f["familia"] and f["familia"].upper() not in f"{c['segmento'] or ''} {c['familia'] or ''} {c['subfamilia'] or ''}".upper():
            continue
        if f["q"]:
            texto = f"{cod} {c['nombre'] or ''} {c['parte'] or ''} {c['equivalencias'] or ''}".upper()
            if not all(p in texto for p in f["q"].upper().split()):
                continue
        x = analizar(cod, c, base, f)
        if x["valor"] < f["valor_min"] or (f["clase"] and x["clase"] != f["clase"]) or (f["sin_recientes"] and x["reciente"]):
            continue
        completa = bool(x["aplicacion"] and x["equivalencias"])
        if (f["info"] == "falta" and completa) or (f["info"] == "completa" and not completa):
            continue
        x["campana"] = activos.get(cod)
        filas.append(x)
    clave = {"cobertura": lambda x: -(x["cobertura"] if x["cobertura"] is not None else 10 ** 6),
             "margen": lambda x: -(x["margen"] if x["margen"] is not None and not x["dudoso"] else -9),
             "descuento": lambda x: (-x["descuento"], -x["valor"])}.get(f["orden"], lambda x: -x["valor"])
    filas.sort(key=clave)
    return filas, base["fin"]


def resumen(filas):
    r = {"productos": len(filas), "valor": sum(x["valor"] for x in filas), "unidades": sum(x["existencia"] for x in filas),
         "clases": {k: {"n": 0, "valor": 0.0} for k in ("sin_venta", "muy_lento", "lento")},
         "falta_ficha": sum(1 for x in filas if not (x["aplicacion"] and x["equivalencias"])),
         "con_descuento": sum(1 for x in filas if x["descuento"] >= 0.05),
         "valor_descuento": sum(x["valor"] for x in filas if x["descuento"] >= 0.05),
         "en_campana": sum(1 for x in filas if x["campana"])}
    for x in filas:
        r["clases"][x["clase"]]["n"] += 1
        r["clases"][x["clase"]]["valor"] += x["valor"]
    return r


def lista_familias():
    def leer():
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT DISTINCT familia FROM clase_prod WHERE familia IS NOT NULL AND familia <> '' ORDER BY familia")
            return [r[0] for r in cursor.fetchall()]
        finally:
            cursor.close()
            conn.close()
    return en_cache("estrategia_familias", leer)


def campanas_activas():
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("SELECT id, nombre, estrategia, fecha_inicio FROM mkt_campanas WHERE estado = 'Activa' ORDER BY creado_en DESC")
        return cursor.fetchall()
    finally:
        cursor.close()
        conn.close()


# ---------------------------------------------------------------- reporte de poco movimiento
@estrategia_bp.route("/productos_poco_movimiento")
@requiere_acceso("estrategia_productos")
def reporte():
    f = leer_filtros(request.args)
    filas, fin = calcular(f)
    paginas = max(1, -(-len(filas) // POR_PAGINA))
    try:
        pagina = min(max(1, int(request.args.get("pagina", 1))), paginas)
    except ValueError:
        pagina = 1
    return render_template("productos_poco_movimiento.html", f=f, fin=fin, r=resumen(filas),
                           filas=filas[(pagina - 1) * POR_PAGINA: pagina * POR_PAGINA], pagina=pagina, paginas=paginas,
                           marcas=lista_marcas(), familias=lista_familias(), coberturas=COBERTURAS, clases=CLASES, ordenes=ORDENES,
                           estrategias=ESTRATEGIAS, campanas=campanas_activas(), hoy=date.today())


@estrategia_bp.route("/productos_poco_movimiento/historial/<path:codigo>")
@requiere_acceso("estrategia_productos")
def historial(codigo):
    """Quiénes lo han comprado (para ofrecérselo) y ventas por año, desde ventas_por_producto."""
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("""
            SELECT TRIM(Cliente) AS cliente, SUBSTRING_INDEX(GROUP_CONCAT(Vendedor ORDER BY Fecha DESC SEPARATOR '|'), '|', 1) AS vendedor,
                   SUM(`Cantidad Vendida`) AS unidades, SUM(`Sub-Total 2`) AS monto, MAX(Fecha) AS ultima, COUNT(DISTINCT `No. Factura`) AS facturas
            FROM ventas_por_producto WHERE Codigo = %s AND Estatus <> 'Anulado'
            GROUP BY TRIM(Cliente) ORDER BY ultima DESC LIMIT 60
        """, (codigo,))
        clientes = cursor.fetchall()
        cursor.execute("""
            SELECT LEFT(Fecha, 4) AS anio, SUM(`Cantidad Vendida`) AS unidades, SUM(`Sub-Total 2`) AS monto
            FROM ventas_por_producto WHERE Codigo = %s AND Estatus <> 'Anulado' GROUP BY LEFT(Fecha, 4) ORDER BY anio
        """, (codigo,))
        anios = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()
    a_num = lambda v: float(v or 0)
    return jsonify(clientes=[{"cliente": c["cliente"] or "—", "vendedor": c["vendedor"] or "—", "unidades": a_num(c["unidades"]),
                              "monto": a_num(c["monto"]), "ultima": c["ultima"], "facturas": c["facturas"]} for c in clientes],
                   anios=[{"anio": a["anio"], "unidades": a_num(a["unidades"]), "monto": a_num(a["monto"])} for a in anios])


def guardar_productos(cursor, campana_id, codigos, cat, base, f):
    """Inserta los productos en la campaña con su foto inicial; devuelve cuántos eran nuevos (los repetidos se ignoran)."""
    nuevos = 0
    for cod in codigos:
        if cod not in cat:
            continue
        x = analizar(cod, cat[cod], base, f)
        cursor.execute("""INSERT IGNORE INTO mkt_campana_productos (campana_id, codigo, existencia_inicial, costo, precio_ref, venta_mensual_base,
                          ultima_venta_previa, agregado_por) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)""",
                       (campana_id, cod, x["existencia"], round(x["costo"], 4), round(x["precio"], 4), round(x["mensual"], 4),
                        x["ultima_venta"], usuario()))
        nuevos += cursor.rowcount
    return nuevos


def resolver_codigos(texto):
    """Texto pegado (códigos o números de parte, uno por línea o separados por coma, punto y coma o tabulador) ->
    (códigos del catálogo, no encontrados, números de parte que corresponden a varios códigos)."""
    tokens = list(dict.fromkeys(t.strip().upper() for t in re.split(r"[\n\r,;\t]+", texto or "") if t.strip()))
    if not tokens:
        return [], [], []
    conn = get_connection()
    cursor = conn.cursor()
    try:
        marcas = ", ".join(["%s"] * len(tokens))
        cursor.execute(f"SELECT `Codigo Producto`, UPPER(TRIM(`Numero Parte`)) FROM catalogo_productos "
                       f"WHERE UPPER(`Codigo Producto`) IN ({marcas}) OR UPPER(TRIM(`Numero Parte`)) IN ({marcas})", tokens + tokens)
        filas = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()
    por_codigo = {str(c).upper(): c for c, _ in filas}
    por_parte = {}
    for c, p in filas:
        por_parte.setdefault(p, []).append(c)
    codigos, no_encontrados, varios = [], [], []
    for t in tokens:
        if t in por_codigo:                     # el código manda sobre el número de parte
            codigos.append(por_codigo[t])
        elif len(por_parte.get(t, [])) == 1:
            codigos.append(por_parte[t][0])
        elif por_parte.get(t):
            varios.append(f"{t} ({', '.join(por_parte[t][:5])})")
        else:
            no_encontrados.append(t)
    return list(dict.fromkeys(codigos)), no_encontrados, varios


@estrategia_bp.route("/campanas_mkt/buscar")
@requiere_acceso("estrategia_productos")
def buscar_productos():
    """Busca en todo el catálogo (no solo en poco movimiento) por código, número de parte, descripción o equivalencias."""
    palabras = (request.args.get("q") or "").upper().split()[:6]
    if not palabras or len("".join(palabras)) < 2:
        return jsonify(productos=[])
    where = " AND ".join(["UPPER(CONCAT_WS(' ', `Codigo Producto`, `Numero Parte`, Nombre, Marca, Equivalencias)) LIKE %s"] * len(palabras))
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute(f"SELECT `Codigo Producto` AS codigo, Nombre AS nombre, `Numero Parte` AS parte, Marca AS marca "
                       f"FROM catalogo_productos WHERE {where} ORDER BY Nombre LIMIT 30", [f"%{p}%" for p in palabras])
        productos = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()
    base = datos_base()
    return jsonify(productos=[{**p, "marca": (p["marca"] or "").strip(), "parte": p["parte"] or "",
                               "existencia": base["existencias"].get(p["codigo"], 0.0),
                               "mensual": base["ventas"].get(p["codigo"], (0.0, 0, 0))[0] / 12} for p in productos])


@estrategia_bp.route("/productos_poco_movimiento/agregar", methods=["POST"])
@requiere_acceso("estrategia_productos")
def agregar_a_campana():
    """Agrega los productos marcados a una campaña (nueva o existente) guardando su foto inicial."""
    codigos = list(dict.fromkeys(c.strip() for c in request.form.getlist("codigo") if c.strip()))
    if not codigos:
        flash("Marque al menos un producto.", "warning")
        return redirect(request.referrer or url_for("estrategia_productos.reporte"))
    f = leer_filtros(request.form)
    base = datos_base()
    cat = catalogo(codigos)
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        conn.start_transaction()
        campana_id = request.form.get("campana_id")
        if campana_id == "nueva" or not (campana_id or "").isdigit():
            nombre = " ".join((request.form.get("nombre") or "").split())[:120]
            if not nombre:
                raise ValueError("Escriba el nombre de la campaña.")
            estrategia = request.form.get("estrategia") if request.form.get("estrategia") in ESTRATEGIAS else "Otra"
            try:
                inicio = datetime.strptime(request.form.get("fecha_inicio") or "", "%Y-%m-%d").date()
            except ValueError:
                inicio = date.today()
            descuento = numero(request.form.get("descuento"), None, 0, 100)
            cursor.execute("INSERT INTO mkt_campanas (nombre, estrategia, descripcion, fecha_inicio, descuento, creado_por) VALUES (%s,%s,%s,%s,%s,%s)",
                           (nombre, estrategia, (request.form.get("descripcion") or "").strip() or None, inicio, descuento, usuario()))
            campana_id = cursor.lastrowid
        else:
            cursor.execute("SELECT id FROM mkt_campanas WHERE id = %s AND estado = 'Activa'", (int(campana_id),))
            if not cursor.fetchone():
                raise ValueError("La campaña no existe o está cerrada.")
            campana_id = int(campana_id)
        nuevos = guardar_productos(cursor, campana_id, codigos, cat, base, f)
        conn.commit()
    except ValueError as e:
        conn.rollback()
        flash(str(e), "warning")
        return redirect(request.referrer or url_for("estrategia_productos.reporte"))
    finally:
        cursor.close()
        conn.close()
    flash(f"✅ {nuevos} producto(s) agregados a la campaña" + (f" ({len(codigos) - nuevos} ya estaban)." if nuevos < len(codigos) else "."), "success")
    return redirect(url_for("estrategia_productos.campana", id_=campana_id))


@estrategia_bp.route("/productos_poco_movimiento/excel")
@requiere_acceso("estrategia_productos")
def descargar_excel():
    f = leer_filtros(request.args)
    filas, fin = calcular(f)
    wb = Workbook()
    ws = wb.active
    ws.title = "Poco movimiento"
    ws["A1"] = f"Productos con existencia y poco movimiento ({f['cobertura']}+ meses de existencia o sin venta en 12 meses)"
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    ws["A2"] = (f"Ventas de los 12 meses al {fin:%d/%m/%Y} · Precio {f['precio']} (U$ × TC {f['tc']:g}) · Margen objetivo {f['margen_obj']:g}% · "
                f"Generado {datetime.now():%d/%m/%Y %H:%M}")
    ws["A2"].font = Font(size=9, color="52514E")
    columnas = [("Código", 11), ("Nº parte", 18), ("Descripción", 34), ("Marca", 13), ("Segmento", 14), ("Familia", 20), ("Ficha", 12),
                ("Existencia", 10), ("Valor inventario C$", 15), ("Vendido 12 m", 10), ("Vendido 3 m", 10), ("Venta mensual", 10),
                ("Meses de existencia", 11), ("Clientes 12 m", 9), ("Última venta", 12), ("Último ingreso", 12),
                ("Costo C$", 12), (f"Precio {f['precio']} C$", 12), ("Margen", 9), ("Descuento máx.", 10), ("Precio con descuento C$", 13),
                ("Estrategia sugerida", 50), ("Campaña", 22)]
    ft = 4
    for i, (nombre, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=ft, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NARANJA if nombre in ("Margen", "Descuento máx.", "Estrategia sugerida") else VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[ft].height = 32
    for n, x in enumerate(filas, start=ft + 1):
        ficha = "Completa" if x["aplicacion"] and x["equivalencias"] else ("Falta ambas" if not x["aplicacion"] and not x["equivalencias"]
                                                                         else ("Falta aplicación" if not x["aplicacion"] else "Falta equivalencia"))
        valores = [x["codigo"], x["parte"], x["nombre"], x["marca"], x["segmento"], x["familia"], ficha, x["existencia"], round(x["valor"], 2),
                   x["v12"], x["v3"], round(x["mensual"], 2), round(x["cobertura"], 1) if x["cobertura"] is not None else "Sin venta",
                   x["clientes"], x["ultima_venta"], x["ultimo_ingreso"], round(x["costo"], 2), round(x["precio"], 2),
                   x["margen"], x["descuento"] or None, round(x["precio_desc"], 2) if x["precio_desc"] else None,
                   " · ".join(t for _, t in x["sugerencias"]), x["campana"] or ""]
        for col, v in enumerate(valores, start=1):
            ws.cell(row=n, column=col, value=v)
        for col in (15, 16):
            ws.cell(row=n, column=col).number_format = "dd/mm/yyyy"
        for col in (9, 17, 18, 21):
            ws.cell(row=n, column=col).number_format = "#,##0.00"
        for col in (19, 20):
            ws.cell(row=n, column=col).number_format = "0.0%"
    ws.freeze_panes = ws.cell(row=ft + 1, column=4)
    ws.auto_filter.ref = f"A{ft}:{get_column_letter(len(columnas))}{ft + max(len(filas), 1)}"
    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=f"poco_movimiento_{datetime.now():%Y-%m-%d}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------------------------------------------------------------- seguimiento de campañas
def ventas_desde(codigos, desde):
    """Unidades netas vendidas desde 'desde' por producto (transacciones: salidas − devoluciones) y clientes."""
    if not codigos:
        return {}
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(f"""
            SELECT {HINT} codigo_producto, SUM(CASE WHEN tipo_transaccion = 'SALIDA' THEN cantidad ELSE -cantidad END),
                   COUNT(DISTINCT CASE WHEN tipo_transaccion = 'SALIDA' THEN cliente_proveedor END),
                   MAX(CASE WHEN tipo_transaccion = 'SALIDA' THEN fecha END)
            FROM transacciones_inventario
            WHERE codigo_producto IN ({', '.join(['%s'] * len(codigos))}) AND fecha >= %s
              AND (tipo_transaccion = 'SALIDA' OR (tipo_transaccion = 'ENTRADA' AND tipo_documento_origen = 'DEVOLUCION'))
            GROUP BY codigo_producto
        """, (*codigos, desde))
        return {r[0]: (max(float(r[1] or 0), 0.0), int(r[2] or 0), a_fecha(r[3]) if r[3] else None) for r in cursor.fetchall()}
    finally:
        cursor.close()
        conn.close()


def montos_desde(codigos, desde):
    """Venta sin IVA por producto desde 'desde' (ventas_por_producto)."""
    if not codigos:
        return {}
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(f"SELECT Codigo, SUM(`Sub-Total 2`) FROM ventas_por_producto WHERE Codigo IN ({', '.join(['%s'] * len(codigos))}) "
                       "AND Fecha >= %s AND Estatus <> 'Anulado' GROUP BY Codigo", (*codigos, str(desde)))
        return {c: float(m or 0) for c, m in cursor.fetchall()}
    finally:
        cursor.close()
        conn.close()


def estado_campana(c, productos):
    """Calcula el avance de cada producto de la campaña y los totales."""
    base = datos_base()
    fin = base["fin"]
    corte = min(c["fecha_fin"], fin) if c["fecha_fin"] else fin
    meses = max((corte - c["fecha_inicio"]).days + 1, 1) / DIAS_MES
    codigos = [p["codigo"] for p in productos]
    vendidos = ventas_desde(codigos, c["fecha_inicio"])
    montos = en_cache(f"mkt_montos_{c['id']}_{len(codigos)}_{fin}", lambda: montos_desde(codigos, c["fecha_inicio"]), segundos=300)
    cat = catalogo(codigos) if codigos else {}
    filas = []
    for p in productos:
        unidades, clientes, ultima = vendidos.get(p["codigo"], (0.0, 0, None))
        base_mensual = float(p["venta_mensual_base"] or 0)
        ritmo = unidades / meses
        inicial = float(p["existencia_inicial"] or 0)
        x = cat.get(p["codigo"], {})
        filas.append({**p, "nombre": x.get("nombre") or "", "parte": x.get("parte") or "", "marca": (x.get("marca") or "").strip(),
                      "existencia_inicial": inicial, "existencia": base["existencias"].get(p["codigo"], 0.0),
                      "costo": float(p["costo"] or 0), "precio_ref": float(p["precio_ref"] or 0), "base": base_mensual,
                      "vendido": unidades, "monto": montos.get(p["codigo"], 0.0), "clientes": clientes, "ultima": ultima,
                      "ritmo": ritmo, "variacion": (ritmo / base_mensual - 1) * 100 if base_mensual else None,
                      "avance": min(unidades / inicial, 1.0) if inicial else 0.0, "valor_inicial": inicial * float(p["costo"] or 0)})
    tot = {"productos": len(filas), "con_venta": sum(1 for x in filas if x["vendido"] > 0),
           "valor_inicial": sum(x["valor_inicial"] for x in filas), "vendido_costo": sum(min(x["vendido"], x["existencia_inicial"]) * x["costo"] for x in filas),
           "monto": sum(x["monto"] for x in filas), "unidades": sum(x["vendido"] for x in filas),
           "unidades_inicial": sum(x["existencia_inicial"] for x in filas),
           "base": sum(x["base"] for x in filas), "ritmo": sum(x["ritmo"] for x in filas), "meses": meses, "corte": corte}
    tot["avance"] = tot["vendido_costo"] / tot["valor_inicial"] if tot["valor_inicial"] else 0
    tot["variacion"] = (tot["ritmo"] / tot["base"] - 1) * 100 if tot["base"] else None
    return filas, tot


def cargar_campana(id_):
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("SELECT * FROM mkt_campanas WHERE id = %s", (id_,))
        c = cursor.fetchone()
        productos = []
        if c:
            cursor.execute("SELECT * FROM mkt_campana_productos WHERE campana_id = %s ORDER BY agregado_en, id", (id_,))
            productos = cursor.fetchall()
        return c, productos
    finally:
        cursor.close()
        conn.close()


@estrategia_bp.route("/campanas_mkt")
@requiere_acceso("estrategia_productos")
def campanas():
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("""SELECT c.*, COUNT(p.id) AS productos FROM mkt_campanas c LEFT JOIN mkt_campana_productos p ON p.campana_id = c.id
                          GROUP BY c.id ORDER BY c.estado = 'Cerrada', c.fecha_inicio DESC, c.id DESC""")
        lista = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()
    for c in lista:
        _, prods = cargar_campana(c["id"])
        _, c["tot"] = estado_campana(c, prods)
    return render_template("campanas_mkt.html", campanas=lista, estrategias=ESTRATEGIAS, hoy=date.today())


@estrategia_bp.route("/campanas_mkt/nueva", methods=["POST"])
@requiere_acceso("estrategia_productos")
def nueva_campana():
    nombre = " ".join((request.form.get("nombre") or "").split())[:120]
    if not nombre:
        flash("Escriba el nombre de la campaña.", "warning")
        return redirect(url_for("estrategia_productos.campanas"))
    try:
        inicio = datetime.strptime(request.form.get("fecha_inicio") or "", "%Y-%m-%d").date()
    except ValueError:
        inicio = date.today()
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("INSERT INTO mkt_campanas (nombre, estrategia, descripcion, fecha_inicio, descuento, creado_por) VALUES (%s,%s,%s,%s,%s,%s)",
                       (nombre, request.form.get("estrategia") if request.form.get("estrategia") in ESTRATEGIAS else "Otra",
                        (request.form.get("descripcion") or "").strip() or None, inicio,
                        numero(request.form.get("descuento"), None, 0, 100), usuario()))
        conn.commit()
        nuevo = cursor.lastrowid
    finally:
        cursor.close()
        conn.close()
    flash("✅ Campaña creada. Agregue productos desde el reporte de poco movimiento.", "success")
    return redirect(url_for("estrategia_productos.campana", id_=nuevo))


@estrategia_bp.route("/campanas_mkt/<int:id_>", methods=["GET", "POST"])
@requiere_acceso("estrategia_productos")
def campana(id_):
    if request.method == "POST" and request.form.get("accion") == "agregar":
        return agregar_productos(id_)
    if request.method == "POST":
        accion = request.form.get("accion")
        conn = get_connection()
        cursor = conn.cursor()
        try:
            if accion == "editar":
                try:
                    inicio = datetime.strptime(request.form.get("fecha_inicio") or "", "%Y-%m-%d").date()
                except ValueError:
                    inicio = None
                try:
                    fin = datetime.strptime(request.form.get("fecha_fin") or "", "%Y-%m-%d").date()
                except ValueError:
                    fin = None
                nombre = " ".join((request.form.get("nombre") or "").split())[:120]
                if nombre and inicio:
                    cursor.execute("""UPDATE mkt_campanas SET nombre=%s, estrategia=%s, descripcion=%s, fecha_inicio=%s, fecha_fin=%s, descuento=%s,
                                      estado=%s WHERE id=%s""",
                                   (nombre, request.form.get("estrategia") if request.form.get("estrategia") in ESTRATEGIAS else "Otra",
                                    (request.form.get("descripcion") or "").strip() or None, inicio, fin,
                                    numero(request.form.get("descuento"), None, 0, 100),
                                    "Cerrada" if request.form.get("estado") == "Cerrada" else "Activa", id_))
                    flash("✅ Campaña actualizada.", "success")
                else:
                    flash("Nombre y fecha de inicio son obligatorios.", "warning")
            elif accion == "quitar":
                cursor.execute("DELETE FROM mkt_campana_productos WHERE campana_id = %s AND codigo = %s", (id_, request.form.get("codigo")))
                flash(f"Producto {request.form.get('codigo')} quitado de la campaña.", "info")
            elif accion == "nota":
                cursor.execute("UPDATE mkt_campana_productos SET nota = %s WHERE campana_id = %s AND codigo = %s",
                               ((request.form.get("nota") or "").strip()[:255] or None, id_, request.form.get("codigo")))
            elif accion == "eliminar":
                cursor.execute("DELETE FROM mkt_campanas WHERE id = %s", (id_,))
                conn.commit()
                flash("Campaña eliminada.", "info")
                return redirect(url_for("estrategia_productos.campanas"))
            conn.commit()
        finally:
            cursor.close()
            conn.close()
        return redirect(url_for("estrategia_productos.campana", id_=id_))

    c, productos = cargar_campana(id_)
    if not c:
        flash("La campaña no existe.", "warning")
        return redirect(url_for("estrategia_productos.campanas"))
    filas, tot = estado_campana(c, productos)
    orden = request.args.get("orden", "vendido")
    filas.sort(key={"avance": lambda x: -x["avance"], "valor": lambda x: -x["valor_inicial"],
                    "sin_venta": lambda x: (x["vendido"] > 0, -x["valor_inicial"])}.get(orden, lambda x: (-x["monto"], -x["vendido"])))
    return render_template("campana_mkt.html", c=c, filas=filas, tot=tot, estrategias=ESTRATEGIAS, orden=orden)


def agregar_productos(id_):
    """Agrega a la campaña cualquier producto del catálogo, pegando o buscando códigos / números de parte."""
    volver = redirect(url_for("estrategia_productos.campana", id_=id_))
    c, _ = cargar_campana(id_)
    if not c or c["estado"] != "Activa":
        flash("Solo se pueden agregar productos a una campaña activa.", "warning")
        return volver
    codigos, no_encontrados, varios = resolver_codigos(request.form.get("codigos"))
    if codigos:
        cat = catalogo(codigos)
        conn = get_connection()
        cursor = conn.cursor()
        try:
            nuevos = guardar_productos(cursor, id_, codigos, cat, datos_base(), leer_filtros({}))
            conn.commit()
        finally:
            cursor.close()
            conn.close()
        flash(f"✅ {nuevos} producto(s) agregados a la campaña" + (f" ({len(codigos) - nuevos} ya estaban)." if nuevos < len(codigos) else "."), "success")
    if no_encontrados:
        flash(f"No están en el catálogo: {', '.join(no_encontrados[:30])}{' …' if len(no_encontrados) > 30 else ''}", "warning")
    if varios:
        flash("Estos números de parte corresponden a varios códigos; agréguelos por código: " + "; ".join(varios[:15]), "warning")
    if not (codigos or no_encontrados or varios):
        flash("Escriba o pegue al menos un código o número de parte.", "warning")
    return volver


@estrategia_bp.route("/campanas_mkt/<int:id_>/ventas/<path:codigo>")
@requiere_acceso("estrategia_productos")
def ventas_campana(id_, codigo):
    """Ventas del producto desde el inicio de la campaña: cliente, vendedor, factura, cantidad y monto."""
    c, _ = cargar_campana(id_)
    if not c:
        return jsonify(error="La campaña no existe."), 404
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("""
            SELECT Fecha AS fecha, `No. Factura` AS factura, TRIM(Cliente) AS cliente, Vendedor AS vendedor, `Contado/Credito` AS tipo,
                   SUM(`Cantidad Vendida`) AS cantidad, SUM(`Sub-Total 2`) AS monto, MAX(`Precio Facturado`) AS precio
            FROM ventas_por_producto WHERE Codigo = %s AND Fecha >= %s AND Estatus <> 'Anulado'
            GROUP BY Fecha, `No. Factura`, TRIM(Cliente), Vendedor, `Contado/Credito` ORDER BY Fecha DESC, factura DESC
        """, (codigo, str(c["fecha_inicio"])))
        filas = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()
    return jsonify(ventas=[{"fecha": r["fecha"], "factura": int(r["factura"]) if r["factura"] else "", "cliente": r["cliente"] or "—",
                            "vendedor": r["vendedor"] or "—", "tipo": r["tipo"] or "", "cantidad": float(r["cantidad"] or 0),
                            "monto": float(r["monto"] or 0), "precio": float(r["precio"] or 0)} for r in filas])


@estrategia_bp.route("/campanas_mkt/<int:id_>/excel")
@requiere_acceso("estrategia_productos")
def excel_campana(id_):
    c, productos = cargar_campana(id_)
    if not c:
        return redirect(url_for("estrategia_productos.campanas"))
    filas, tot = estado_campana(c, productos)
    wb = Workbook()
    ws = wb.active
    ws.title = "Campaña"
    ws["A1"] = f"Campaña: {c['nombre']} ({c['estrategia']})"
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    ws["A2"] = (f"Desde {c['fecha_inicio']:%d/%m/%Y} al {tot['corte']:%d/%m/%Y} ({tot['meses']:.1f} meses) · "
                f"Descuento {c['descuento'] or 0:g}% · Generado {datetime.now():%d/%m/%Y %H:%M}")
    ws["A2"].font = Font(size=9, color="52514E")
    columnas = [("Código", 11), ("Nº parte", 18), ("Descripción", 34), ("Marca", 13), ("Existencia inicial", 11), ("Existencia hoy", 11),
                ("Costo C$", 12), ("Precio ref. C$", 12), ("Venta mensual antes", 11), ("Vendido desde inicio", 11), ("Venta mensual ahora", 11),
                ("Variación", 10), ("% de la existencia vendida", 11), ("Monto vendido C$", 14), ("Clientes", 9), ("Última venta", 12), ("Nota", 30)]
    for i, (nombre, ancho) in enumerate(columnas, start=1):
        cel = ws.cell(row=4, column=i, value=nombre)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=NARANJA if i in (10, 11, 12, 13, 14) else VERDE)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[4].height = 32
    for n, x in enumerate(filas, start=5):
        valores = [x["codigo"], x["parte"], x["nombre"], x["marca"], x["existencia_inicial"], x["existencia"], x["costo"], x["precio_ref"],
                   round(x["base"], 2), x["vendido"], round(x["ritmo"], 2), x["variacion"] / 100 if x["variacion"] is not None else None,
                   x["avance"], round(x["monto"], 2), x["clientes"], x["ultima"], x["nota"] or ""]
        for col, v in enumerate(valores, start=1):
            ws.cell(row=n, column=col, value=v)
        for col in (7, 8, 14):
            ws.cell(row=n, column=col).number_format = "#,##0.00"
        for col in (12, 13):
            ws.cell(row=n, column=col).number_format = "0%"
        ws.cell(row=n, column=16).number_format = "dd/mm/yyyy"
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(columnas))}{4 + max(len(filas), 1)}"
    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=f"campana_{id_}_{datetime.now():%Y-%m-%d}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
