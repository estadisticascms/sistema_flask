"""Ventas perdidas: registro rápido para vendedores (celular o computadora), listado con filtros y
resumen por producto. Cada registro se liga al producto del catálogo (codigo_producto) para mostrarlo
en el Sugerido de compras; los registros anteriores se ligan por número de parte cuando coincide con un solo producto."""
import io
import os
import re
import time
from datetime import date, datetime, timedelta

import pandas as pd
from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, session, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from werkzeug.utils import secure_filename

from utils.db import get_connection
from router.accesos import requiere_acceso
from router.ventas_bodegas import en_cache
from router.reporte_armados import nombre_limpio

ventas_perdidas_bp = Blueprint("ventas_perdidas", __name__)

MOTIVOS = ["Sin existencia", "No lo manejamos", "Precio", "Tiempo de entrega", "Otro"]
PALABRAS_CLAVE = ["RODAMIENTO", "RETENEDOR", "CORREA", "FILTRO", "MANGUERA", "FITINERIA", "EQUIPO PESADO", "EQUIPO LIVIANO"]
LARGOS = {"nomb_vendedor": 50, "nomb_cliente": 100, "numero_parte": 100, "codigo_producto": 30,
          "descripcion": 100, "marca": 100, "motivo": 50}
MINUTOS_DESHACER = 15
DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
EMOJI_MOTIVO = {"Sin existencia": "📦", "No lo manejamos": "🚫", "Precio": "💲", "Tiempo de entrega": "⏱️", "Otro": "📝"}
VERDE, NARANJA = "027266", "FE6412"


def normal_parte(texto):
    return re.sub(r"[^A-Z0-9]", "", (texto or "").upper())


# -----------------------------
# Datos de apoyo
# -----------------------------
def lista_vendedores():
    """Nombres como los usa el sistema (los que facturaron en los últimos 6 meses)."""
    def leer():
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT vendedor, COUNT(*) FROM ventas_historicas WHERE Fecha >= %s AND vendedor <> '' "
                           "GROUP BY vendedor HAVING COUNT(*) >= 5 ORDER BY vendedor",
                           (date.today() - timedelta(days=180),))
            return [r[0].strip() for r in cursor.fetchall() if r[0] and r[0].strip()]
        finally:
            cursor.close()
            conn.close()
    return en_cache("vp_vendedores", leer)


def indice_catalogo():
    """Número de parte normalizado -> productos del catálogo (para ligar registros sin código)."""
    def leer():
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT `Codigo Producto`, `Numero Parte`, Marca FROM catalogo_productos WHERE `Codigo Producto` IS NOT NULL")
            indice = {}
            for cod, parte, marca in cursor.fetchall():
                for p in {normal_parte(parte)} | {normal_parte(x) for x in (parte or "").split("/")}:
                    if p:
                        indice.setdefault(p, {})[cod] = (marca or "").strip().upper()
            return indice
        finally:
            cursor.close()
            conn.close()
    return en_cache("vp_indice_catalogo", leer)


def ligar(registro, indice=None):
    """Código del catálogo de un registro: el guardado o, si no tiene, el único producto con ese número de parte
    (si hay varios, el de la marca anotada)."""
    if registro.get("codigo_producto"):
        return registro["codigo_producto"], "elegido"
    candidatos = (indice or indice_catalogo()).get(normal_parte(registro.get("numero_parte")), {})
    if len(candidatos) == 1:
        return next(iter(candidatos)), "por número de parte"
    pista = " ".join(filter(None, [registro.get("marca"), registro.get("nota")])).upper()
    if len(candidatos) > 1 and pista:
        por_marca = [c for c, m in candidatos.items() if m and m in pista]
        if len(por_marca) == 1:
            return por_marca[0], "por número de parte y marca"
    return None, "varios en catálogo" if candidatos else "no está en catálogo"


def perdidas_por_codigo(desde, hasta):
    """Para el Sugerido de compras: {codigo: {unidades, registros, clientes, ultima}} en el período."""
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("SELECT * FROM ventas_perdidas WHERE fecha BETWEEN %s AND %s", (desde, hasta))
        registros = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()
    indice = indice_catalogo()
    resultado = {}
    for r in registros:
        cod, _ = ligar(r, indice)
        if not cod:
            continue
        x = resultado.setdefault(cod, {"unidades": 0.0, "registros": 0, "clientes": set(), "ultima": None})
        x["unidades"] += float(r["cantidad"] or 0)
        x["registros"] += 1
        if r["nomb_cliente"]:
            x["clientes"].add(" ".join(r["nomb_cliente"].split()).upper())
        if r["fecha"] and (x["ultima"] is None or r["fecha"] > x["ultima"]):
            x["ultima"] = r["fecha"]
    return resultado


def detalle_por_codigo(codigo, desde, hasta):
    """Registros de venta perdida ligados a un producto en el período (para el Sugerido de compras)."""
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("SELECT * FROM ventas_perdidas WHERE fecha BETWEEN %s AND %s ORDER BY fecha DESC, id DESC", (desde, hasta))
        registros = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()
    indice = indice_catalogo()
    resultado = []
    for r in registros:
        cod, liga = ligar(r, indice)
        if cod == codigo:
            resultado.append({"fecha": r["fecha"].strftime("%d/%m/%Y") if r["fecha"] else "", "vendedor": r["nomb_vendedor"] or "",
                              "cliente": r["nomb_cliente"] or "", "cantidad": float(r["cantidad"] or 0),
                              "motivo": r["motivo"] or "", "nota": r["nota"] or "", "numero_parte": r["numero_parte"] or "",
                              "liga": liga})
    return resultado


def leer_formulario(datos):
    """Valida un registro (formulario o JSON). Devuelve (registro, errores)."""
    def texto(campo):
        return (str(datos.get(campo) or "")).strip()

    r = {c: texto(c) or None for c in ("nomb_vendedor", "nomb_cliente", "numero_parte", "codigo_producto",
                                       "descripcion", "marca", "motivo", "nota")}
    errores = []
    try:
        r["fecha"] = datetime.strptime(texto("fecha") or date.today().isoformat(), "%Y-%m-%d").date()
        if r["fecha"] > date.today():
            errores.append("La fecha no puede ser futura.")
    except ValueError:
        errores.append("Fecha inválida.")
    try:
        r["cantidad"] = float(texto("cantidad").replace(",", ""))
        if r["cantidad"] <= 0:
            errores.append("La cantidad debe ser mayor que cero.")
    except ValueError:
        errores.append("Indique la cantidad.")
    if not r["nomb_vendedor"]:
        errores.append("Seleccione el vendedor.")
    if not r["nomb_cliente"]:
        errores.append("Indique el cliente.")
    if not r["numero_parte"] and not r["descripcion"]:
        errores.append("Busque el producto o escriba el número de parte / descripción.")
    if r["motivo"] and r["motivo"] not in MOTIVOS:
        r["motivo"] = "Otro"
    for campo, largo in LARGOS.items():
        if r[campo] and len(r[campo]) > largo:
            r[campo] = r[campo][:largo]
    if r["codigo_producto"]:   # que el código exista y que parte/marca/descripción coincidan con el catálogo
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT `Numero Parte`, Nombre, Marca FROM catalogo_productos WHERE `Codigo Producto` = %s",
                           (r["codigo_producto"],))
            p = cursor.fetchone()
        finally:
            cursor.close()
            conn.close()
        if p:
            r["numero_parte"] = (p[0] or r["numero_parte"] or "")[:100] or None
            r["descripcion"] = (p[1] or r["descripcion"] or "")[:100] or None
            r["marca"] = (p[2] or r["marca"] or "")[:100] or None
        else:
            r["codigo_producto"] = None
    return r, errores


COLUMNAS = ["nomb_vendedor", "nomb_cliente", "fecha", "numero_parte", "codigo_producto", "descripcion", "marca",
            "cantidad", "motivo", "nota"]


def insertar(r):
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(f"INSERT INTO ventas_perdidas ({', '.join(COLUMNAS)}) VALUES ({', '.join(['%s'] * len(COLUMNAS))})",
                       [r[c] for c in COLUMNAS])
        conn.commit()
        return cursor.lastrowid
    finally:
        cursor.close()
        conn.close()


def leer_filtros(args):
    hoy = date.today()
    def fecha(nombre, defecto):
        try:
            return datetime.strptime(args.get(nombre) or "", "%Y-%m-%d").date()
        except ValueError:
            return defecto
    return {"vendedor": (args.get("vendedor") or args.get("filtro_vendedor") or "").strip(),
            "desde": fecha("desde", fecha("fecha_inicio", hoy - timedelta(days=90))),
            "hasta": fecha("hasta", fecha("fecha_fin", hoy)),
            "q": (args.get("q") or "").strip(), "motivo": (args.get("motivo") or "").strip(),
            "liga": args.get("liga") if args.get("liga") in ("si", "no") else ""}


def consultar(f):
    where, params = ["fecha BETWEEN %s AND %s"], [f["desde"], f["hasta"]]
    if f["vendedor"]:
        where.append("nomb_vendedor = %s"); params.append(f["vendedor"])
    if f["motivo"]:
        where.append("motivo = %s"); params.append(f["motivo"])
    for palabra in f["q"].split():
        where.append("(nomb_cliente LIKE %s OR numero_parte LIKE %s OR descripcion LIKE %s OR marca LIKE %s OR nota LIKE %s)")
        params += [f"%{palabra}%"] * 5
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute(f"SELECT * FROM ventas_perdidas WHERE {' AND '.join(where)} ORDER BY fecha DESC, id DESC", params)
        registros = cursor.fetchall()
        codigos = set()
        indice = indice_catalogo()
        for r in registros:
            r["codigo_ligado"], r["liga"] = ligar(r, indice)
            if r["codigo_ligado"]:
                codigos.add(r["codigo_ligado"])
        productos = {}
        lista = sorted(codigos)
        for i in range(0, len(lista), 1000):
            lote = lista[i:i + 1000]
            cursor.execute(f"SELECT `Codigo Producto` AS codigo, Nombre AS nombre, `Numero Parte` AS parte, Marca AS marca, "
                           f"`Existencia Actual` AS existencia FROM catalogo_productos "
                           f"WHERE `Codigo Producto` IN ({', '.join(['%s'] * len(lote))})", lote)
            productos.update({p["codigo"]: p for p in cursor.fetchall()})
    finally:
        cursor.close()
        conn.close()
    for r in registros:
        r["producto"] = productos.get(r["codigo_ligado"])
    if f["liga"] == "si":
        registros = [r for r in registros if r["codigo_ligado"]]
    elif f["liga"] == "no":
        registros = [r for r in registros if not r["codigo_ligado"]]
    return registros


def resumen_productos(registros):
    """Lo más pedido que se perdió: por producto del catálogo o, si no está, por número de parte."""
    grupos = {}
    for r in registros:
        clave = r["codigo_ligado"] or "P:" + (normal_parte(r["numero_parte"]) or (r["descripcion"] or "").upper())
        g = grupos.setdefault(clave, {"codigo": r["codigo_ligado"], "producto": r["producto"],
                                      "numero_parte": r["numero_parte"], "descripcion": r["descripcion"],
                                      "marca": r["marca"], "unidades": 0.0, "registros": 0, "clientes": set(),
                                      "vendedores": set(), "ultima": None})
        g["unidades"] += float(r["cantidad"] or 0)
        g["registros"] += 1
        g["clientes"].add(" ".join((r["nomb_cliente"] or "").split()).upper())
        g["vendedores"].add((r["nomb_vendedor"] or "").strip())
        if r["fecha"] and (g["ultima"] is None or r["fecha"] > g["ultima"]):
            g["ultima"] = r["fecha"]
    return sorted(grupos.values(), key=lambda g: (-len(g["clientes"]), -g["unidades"]))


def volver_seguro(url):
    """Solo se regresa a una pantalla de ventas perdidas (evita redirigir a otro sitio)."""
    return url if url and url.startswith("/ventas_perdidas") else url_for("ventas_perdidas.ventas_perdidas", vista="registros")


def periodo_txt(f):
    if f["desde"] == f["hasta"]:
        return f"{DIAS[f['desde'].weekday()].capitalize()} {f['desde']:%d/%m/%Y}"
    return f"Del {f['desde']:%d/%m/%Y} al {f['hasta']:%d/%m/%Y}"


def nombre_producto(r):
    """Número de parte · descripción · marca, del catálogo si está ligado."""
    p = r.get("producto") or {}
    parte = (p.get("parte") or r["numero_parte"] or "").strip()
    desc = (p.get("nombre") or r["descripcion"] or "").strip()
    marca = (p.get("marca") or r["marca"] or "").strip()
    return parte, desc, marca


def nota_util(r):
    """La nota solo si aporta: los registros viejos traen '-' o repiten la marca."""
    nota = (r["nota"] or "").strip()
    _, _, marca = nombre_producto(r)
    if len(re.sub(r"[^A-Za-z0-9]", "", nota)) < 2 or nota.upper() == marca.upper():
        return ""
    return nota


def reporte_dia(registros):
    """Para compartir: vendedor -> cliente -> productos, más totales por motivo."""
    vendedores = {}
    for r in sorted(registros, key=lambda x: (x["fecha"], x["id"])):
        r["nota_ver"] = nota_util(r)
        v = vendedores.setdefault(nombre_limpio((r["nomb_vendedor"] or "Sin vendedor").strip()), {"clientes": {}, "unidades": 0.0, "registros": 0})
        cliente = " ".join((r["nomb_cliente"] or "Sin cliente").split())
        v["clientes"].setdefault(cliente.upper(), {"nombre": cliente, "items": []})["items"].append(r)
        v["unidades"] += float(r["cantidad"] or 0)
        v["registros"] += 1
    motivos = {}
    for r in registros:
        if r["motivo"]:
            motivos[r["motivo"]] = motivos.get(r["motivo"], 0) + 1
    return {"vendedores": [{"nombre": k, **v, "clientes": list(v["clientes"].values())} for k, v in sorted(vendedores.items())],
            "motivos": sorted(motivos.items(), key=lambda x: -x[1]),
            "clientes": len({" ".join((r["nomb_cliente"] or "").split()).upper() for r in registros}),
            "unidades": sum(float(r["cantidad"] or 0) for r in registros), "registros": len(registros)}


def texto_whatsapp(rep_, f):
    """Mensaje listo para pegar en el grupo (WhatsApp: *negrita*, _cursiva_)."""
    def n(v):
        return f"{v:,.0f}" if float(v) == int(v) else f"{v:,.2f}"
    lineas = ["📉 *VENTAS PERDIDAS*", f"📅 {periodo_txt(f)}"]
    if f["vendedor"]:
        lineas.append(f"👤 {nombre_limpio(f['vendedor'])}")
    if not rep_["registros"]:
        return "\n".join(lineas + ["", "Sin ventas perdidas registradas. ✅"])
    for v in rep_["vendedores"]:
        lineas.append("")
        if not f["vendedor"]:
            lineas.append(f"👤 *{v['nombre']}* ({v['registros']})")
        for c in v["clientes"]:
            lineas.append(f"🏪 *{c['nombre']}*")
            for r in c["items"]:
                parte, desc, marca = nombre_producto(r)
                producto = " – ".join(x for x in (parte, desc[:45]) if x) + (f" ({marca})" if marca else "")
                fecha = f"{r['fecha']:%d/%m} " if f["desde"] != f["hasta"] else ""
                linea = f"   • {fecha}{n(float(r['cantidad'] or 0))} × {producto}"
                if r["motivo"]:
                    linea += f" · _{r['motivo']}_"
                lineas.append(linea)
                if r["nota_ver"]:
                    lineas.append(f"      ↳ {r['nota_ver']}")
    lineas += ["", f"*Total:* {rep_['registros']} producto{'s' if rep_['registros'] != 1 else ''} · {n(rep_['unidades'])} unidades · "
                   f"{rep_['clientes']} cliente{'s' if rep_['clientes'] != 1 else ''}"]
    if rep_["motivos"]:
        lineas.append(" · ".join(f"{EMOJI_MOTIVO.get(m, '•')} {m}: {c}" for m, c in rep_["motivos"]))
    return "\n".join(lineas)


# -----------------------------
# Vistas
# -----------------------------
@ventas_perdidas_bp.route("/ventas_perdidas", methods=["GET", "POST"])
@requiere_acceso("ventas_perdidas")
def ventas_perdidas():
    if request.method == "POST":
        es_fetch = request.headers.get("X-Requested-With") == "fetch"
        r, errores = leer_formulario(request.get_json(silent=True) if request.is_json else request.form)
        if errores:
            if es_fetch:
                return jsonify(ok=False, errores=errores), 400
            for e in errores:
                flash(f"⚠️ {e}", "warning")
            return redirect(url_for("ventas_perdidas.ventas_perdidas"))
        nuevo_id = insertar(r)
        # El vendedor puede deshacer lo que acaba de registrar aunque no tenga permiso de eliminar
        recientes = [x for x in session.get("vp_recientes", []) if time.time() - x[1] < MINUTOS_DESHACER * 60]
        session["vp_recientes"] = recientes + [[nuevo_id, time.time()]]
        if es_fetch:
            return jsonify(ok=True, id=nuevo_id, registro={**{k: r[k] for k in COLUMNAS if k != "fecha"},
                                                           "fecha": r["fecha"].isoformat(), "cantidad": r["cantidad"]})
        flash("✅ Venta perdida registrada", "success")
        return redirect(url_for("ventas_perdidas.ventas_perdidas"))

    vista = request.args.get("vista") if request.args.get("vista") in ("registrar", "registros", "resumen") else "registrar"
    f = leer_filtros(request.args)
    if vista == "registros" and not request.args.get("desde") and not request.args.get("hasta"):
        f["desde"] = f["hasta"] = date.today()   # el reporte se pasa a diario: por defecto, hoy
    registros, resumen, reporte, texto = [], [], None, ""
    if vista in ("registros", "resumen"):
        registros = consultar(f)
        if vista == "resumen":
            resumen = resumen_productos(registros)
        else:
            reporte = reporte_dia(registros)
            texto = texto_whatsapp(reporte, f)
    vendedores = lista_vendedores()
    # Vendedores de registros anteriores escritos a mano (solo para filtrar)
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT DISTINCT nomb_vendedor FROM ventas_perdidas WHERE nomb_vendedor IS NOT NULL ORDER BY nomb_vendedor")
        vendedores_filtro = sorted(set(vendedores) | {x[0] for x in cursor.fetchall() if x[0]}, key=str.upper)
    finally:
        cursor.close()
        conn.close()
    return render_template("ventas_perdidas.html", vista=vista, f=f, registros=registros, resumen=resumen,
                           vendedores=vendedores, vendedores_filtro=vendedores_filtro, motivos=MOTIVOS,
                           palabras_clave=PALABRAS_CLAVE, hoy=date.today(), reporte=reporte, texto=texto,
                           periodo=periodo_txt(f), emoji_motivo=EMOJI_MOTIVO, nombre_limpio=nombre_limpio, dias=DIAS, timedelta=timedelta,
                           ligados=sum(1 for r in registros if r["codigo_ligado"]),
                           unidades=sum(float(r["cantidad"] or 0) for r in registros))


@ventas_perdidas_bp.route("/ventas_perdidas/buscar_producto")
@requiere_acceso("ventas_perdidas")
def buscar_producto():
    q = (request.args.get("q") or "").strip()
    if len(q) < 2:
        return jsonify([])
    palabras = q.split()[:5]
    exacto = normal_parte(q)
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        where = " AND ".join(["(`Numero Parte` LIKE %s OR Nombre LIKE %s OR `Codigo Producto` LIKE %s)"] * len(palabras))
        params = [x for p in palabras for x in [f"%{p}%"] * 3]
        cursor.execute(f"""
            SELECT `Codigo Producto` AS codigo, `Numero Parte` AS parte, Nombre AS nombre, Marca AS marca,
                   COALESCE(`Existencia Actual`, 0) AS existencia
            FROM catalogo_productos WHERE `Codigo Producto` IS NOT NULL AND {where} LIMIT 200
        """, params)
        filas = cursor.fetchall()
        # También por número de parte escrito sin guiones/espacios (ej. 6204-2RS = 62042RS)
        extra = []
        if len(exacto) >= 4:   # empieza igual, sin importar guiones ni espacios (62042RS -> 6204-2RSC3)
            for parte, codigos in indice_catalogo().items():
                if parte.startswith(exacto):
                    extra.extend(codigos)
                    if len(extra) >= 40:
                        break
        vistos = {f["codigo"] for f in filas}
        faltan = [c for c in dict.fromkeys(extra) if c not in vistos]
        if faltan:
            cursor.execute(f"""
                SELECT `Codigo Producto` AS codigo, `Numero Parte` AS parte, Nombre AS nombre, Marca AS marca,
                       COALESCE(`Existencia Actual`, 0) AS existencia
                FROM catalogo_productos WHERE `Codigo Producto` IN ({', '.join(['%s'] * len(faltan))})
            """, faltan)
            filas += cursor.fetchall()
    finally:
        cursor.close()
        conn.close()
    # Primero coincidencia exacta de número de parte, luego los que tienen existencia
    for x in filas:
        partes = {normal_parte(x["parte"])} | {normal_parte(p) for p in (x["parte"] or "").split("/")}
        x["exacto"] = exacto in partes
        x["existencia"] = float(x["existencia"] or 0)
    filas.sort(key=lambda x: (not x["exacto"], x["existencia"] <= 0, (x["parte"] or "")))
    return jsonify(filas[:15])


@ventas_perdidas_bp.route("/ventas_perdidas/buscar_cliente")
@requiere_acceso("ventas_perdidas")
def buscar_cliente():
    q = (request.args.get("q") or "").strip()
    if len(q) < 2:
        return jsonify([])
    palabras = q.split()[:5]
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute(f"SELECT nombre, estado, municipio FROM clientes WHERE {' AND '.join(['nombre LIKE %s'] * len(palabras))} "
                       f"ORDER BY estado = 'Activo' DESC, nombre LIMIT 12", [f"%{p}%" for p in palabras])
        return jsonify([{"nombre": n, "activo": e == "Activo", "municipio": m} for n, e, m in cursor.fetchall()])
    finally:
        cursor.close()
        conn.close()


@ventas_perdidas_bp.route("/ventas_perdidas/recientes")
@requiere_acceso("ventas_perdidas")
def recientes():
    """Lo que el vendedor registró hoy (para verlo debajo del formulario)."""
    vendedor = (request.args.get("vendedor") or "").strip()
    if not vendedor:
        return jsonify([])
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("SELECT id, nomb_cliente, numero_parte, descripcion, marca, cantidad, motivo FROM ventas_perdidas "
                       "WHERE nomb_vendedor = %s AND fecha = %s ORDER BY id DESC LIMIT 30", (vendedor, date.today()))
        filas = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()
    propios = {x[0] for x in session.get("vp_recientes", []) if time.time() - x[1] < MINUTOS_DESHACER * 60}
    for x in filas:
        x["cantidad"] = float(x["cantidad"] or 0)
        x["deshacer"] = x["id"] in propios
    return jsonify(filas)


@ventas_perdidas_bp.route("/ventas_perdidas/deshacer/<int:id>", methods=["POST"])
@requiere_acceso("ventas_perdidas")
def deshacer(id):
    propios = {x[0] for x in session.get("vp_recientes", []) if time.time() - x[1] < MINUTOS_DESHACER * 60}
    if id not in propios:
        return jsonify(ok=False, error=f"Solo se puede deshacer un registro propio en los primeros {MINUTOS_DESHACER} minutos."), 403
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM ventas_perdidas WHERE id = %s", (id,))
        conn.commit()
    finally:
        cursor.close()
        conn.close()
    session["vp_recientes"] = [x for x in session.get("vp_recientes", []) if x[0] != id]
    return jsonify(ok=True)


@ventas_perdidas_bp.route("/ventas_perdidas/excel")
@requiere_acceso("ventas_perdidas")
def descargar_excel():
    f = leer_filtros(request.args)
    registros = consultar(f)
    wb = Workbook()
    ws = wb.active
    ws.title = "Registros"
    titulo = f"Ventas perdidas del {f['desde']:%d/%m/%Y} al {f['hasta']:%d/%m/%Y}" + (f" · {f['vendedor']}" if f["vendedor"] else "")
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    columnas = [("Fecha", 11), ("Vendedor", 24), ("Cliente", 34), ("Nº parte", 18), ("Descripción", 34), ("Marca", 14),
                ("Cantidad", 9), ("Motivo", 16), ("Nota", 30), ("Código catálogo", 13), ("Ligado", 22), ("Existencia hoy", 10)]
    hoja_encabezado(ws, columnas, 3)
    for n, r in enumerate(registros, start=4):
        p = r["producto"] or {}
        for col, v in enumerate([r["fecha"], r["nomb_vendedor"], r["nomb_cliente"], r["numero_parte"], r["descripcion"],
                                 r["marca"], float(r["cantidad"] or 0), r["motivo"], r["nota"], r["codigo_ligado"],
                                 r["liga"], p.get("existencia")], start=1):
            ws.cell(row=n, column=col, value=v)
        ws.cell(row=n, column=1).number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(columnas))}{3 + max(len(registros), 1)}"

    ws2 = wb.create_sheet("Por producto")
    ws2["A1"] = titulo + " · agrupado por producto"
    ws2["A1"].font = Font(bold=True, size=14, color=VERDE)
    columnas2 = [("Código catálogo", 13), ("Nº parte", 18), ("Descripción", 36), ("Marca", 14), ("Unidades", 10),
                 ("Clientes", 9), ("Registros", 9), ("Vendedores", 30), ("Última", 11), ("Existencia hoy", 10)]
    hoja_encabezado(ws2, columnas2, 3)
    for n, g in enumerate(resumen_productos(registros), start=4):
        p = g["producto"] or {}
        for col, v in enumerate([g["codigo"], p.get("parte") or g["numero_parte"], p.get("nombre") or g["descripcion"],
                                 p.get("marca") or g["marca"], g["unidades"], len(g["clientes"]), g["registros"],
                                 ", ".join(sorted(v for v in g["vendedores"] if v)), g["ultima"], p.get("existencia")], start=1):
            ws2.cell(row=n, column=col, value=v)
        ws2.cell(row=n, column=9).number_format = "dd/mm/yyyy"
    ws2.freeze_panes = "A4"
    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=f"ventas_perdidas_{f['desde']}_al_{f['hasta']}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def hoja_encabezado(ws, columnas, fila):
    for i, (nombre, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=fila, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho


# -----------------------------
# Editar registro
# -----------------------------
@ventas_perdidas_bp.route("/editar_venta_perdida/<int:id>", methods=["GET", "POST"])
@requiere_acceso("editar_venta_perdida")
def editar_venta_perdida(id):
    if request.method == "POST":
        r, errores = leer_formulario(request.form)
        if errores:
            for e in errores:
                flash(f"⚠️ {e}", "warning")
            return redirect(url_for("ventas_perdidas.editar_venta_perdida", id=id, volver=request.form.get("volver")))
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute(f"UPDATE ventas_perdidas SET {', '.join(f'{c} = %s' for c in COLUMNAS)} WHERE id = %s",
                           [r[c] for c in COLUMNAS] + [id])
            conn.commit()
        finally:
            cursor.close()
            conn.close()
        flash("✅ Registro actualizado correctamente", "success")
        return redirect(volver_seguro(request.form.get("volver")))

    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("SELECT * FROM ventas_perdidas WHERE id=%s", (id,))
        registro = cursor.fetchone()
    finally:
        cursor.close()
        conn.close()
    if not registro:
        flash("El registro no existe.", "warning")
        return redirect(url_for("ventas_perdidas.ventas_perdidas", vista="registros"))
    return render_template("ventas_perdidas_edit.html", registro=registro, motivos=MOTIVOS, volver=volver_seguro(request.args.get("volver")))


# -----------------------------
# Eliminar registro
# -----------------------------
@ventas_perdidas_bp.route("/eliminar_venta_perdida/<int:id>", methods=["GET", "POST"])
@requiere_acceso("eliminar_venta_perdida")
def eliminar_venta_perdida(id):
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM ventas_perdidas WHERE id=%s", (id,))
        conn.commit()
    finally:
        cursor.close()
        conn.close()
    flash("🗑️ Registro eliminado correctamente", "success")
    return redirect(volver_seguro(request.form.get("volver") or request.args.get("volver")))


# -----------------------------
# Importar desde Excel
# -----------------------------
@ventas_perdidas_bp.route("/importar_ventas_perdidas", methods=["POST"])
@requiere_acceso("importar_ventas_perdidas")
def importar_ventas_perdidas():
    archivo = request.files.get("archivo_excel")
    if not archivo or not archivo.filename:
        flash("No se seleccionó archivo", "danger")
        return redirect(url_for("ventas_perdidas.ventas_perdidas", vista="registros"))

    filepath = os.path.join("uploads", secure_filename(archivo.filename))
    archivo.save(filepath)
    try:
        df = pd.read_excel(filepath, dtype=object, keep_default_na=False, na_values=[])
        obligatorias = ["nomb_vendedor", "nomb_cliente", "fecha", "numero_parte", "descripcion", "cantidad"]
        faltan = [c for c in obligatorias if c not in df.columns]
        if faltan:
            flash(f"Columnas faltantes en el Excel: {', '.join(faltan)} (opcionales: marca, motivo, nota, codigo_producto)", "danger")
            return redirect(url_for("ventas_perdidas.ventas_perdidas", vista="registros"))
        validos, malos = [], []
        for i, fila in df.iterrows():
            datos = {k: fila.get(k, "") for k in COLUMNAS}
            fecha = fila.get("fecha")
            datos["fecha"] = fecha.strftime("%Y-%m-%d") if hasattr(fecha, "strftime") else str(fecha or "")[:10]
            r, errores = leer_formulario(datos)
            if errores:
                malos.append(f"fila {i + 2}: {' '.join(errores)}")
            else:
                validos.append(r)
        if malos:   # nada se guarda si hay filas con error
            flash("No se importó nada. Corrija: " + "; ".join(malos[:10]) + (" …" if len(malos) > 10 else ""), "danger")
            return redirect(url_for("ventas_perdidas.ventas_perdidas", vista="registros"))
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.executemany(f"INSERT INTO ventas_perdidas ({', '.join(COLUMNAS)}) VALUES ({', '.join(['%s'] * len(COLUMNAS))})",
                               [[r[c] for c in COLUMNAS] for r in validos])
            conn.commit()
        finally:
            cursor.close()
            conn.close()
        flash(f"✅ Se importaron {len(validos):,} ventas perdidas", "success")
    except Exception as e:
        flash(f"Error al procesar archivo: {e}", "danger")
    finally:
        if os.path.exists(filepath):
            os.remove(filepath)
    return redirect(url_for("ventas_perdidas.ventas_perdidas", vista="registros"))
