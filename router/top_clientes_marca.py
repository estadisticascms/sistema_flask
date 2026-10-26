"""Top de clientes por marca: cuánto compró cada cliente de una o varias marcas, mes a mes, y su vendedor.
Fuente: ventas_por_producto (detalle por producto), venta neta sin IVA y sin facturas anuladas."""
import io
from datetime import date, datetime

from flask import Blueprint, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.db import get_connection
from router.accesos import requiere_acceso
from router.dashboard import ultima_fecha_con_ventas
from router.ventas_bodegas import en_cache
from router.ventas_vendedor import rangos_rapidos

top_clientes_marca_bp = Blueprint("top_clientes_marca", __name__)

TAMANOS = [10, 20, 50, 100, 0]   # 0 = todos
MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
VERDE, NARANJA = "027266", "FE6412"


def fecha_param(valor, defecto):
    try:
        return datetime.strptime(valor or "", "%Y-%m-%d").date()
    except ValueError:
        return defecto


def leer_filtros(args):
    ultima = ultima_fecha_con_ventas() or date.today()
    if isinstance(ultima, datetime):
        ultima = ultima.date()
    hasta = fecha_param(args.get("hasta"), ultima)
    desde = min(fecha_param(args.get("desde"), date(hasta.year, 1, 1)), hasta)
    try:
        top = int(args.get("top", 20))
    except ValueError:
        top = 20
    return {"desde": desde, "hasta": hasta, "ultima": ultima,
            "marcas": sorted({m.strip() for m in args.getlist("marca") if m.strip()}),
            "vendedor": (args.get("vendedor") or "").strip(), "q": (args.get("q") or "").strip(),
            "top": top if top in TAMANOS else 20}


def meses_entre(desde, hasta):
    meses, a, m = [], desde.year, desde.month
    while (a, m) <= (hasta.year, hasta.month):
        meses.append((f"{a}-{m:02d}", f"{MESES[m - 1]} {str(a)[2:]}"))
        a, m = (a + 1, 1) if m == 12 else (a, m + 1)
    return meses


def calcular(f):
    where = ["Fecha BETWEEN %s AND %s", "Estatus <> 'Anulado'"]
    params = [f["desde"].isoformat(), f["hasta"].isoformat()]   # Fecha es texto 'YYYY-MM-DD'
    if f["marcas"]:
        where.append(f"Marca IN ({', '.join(['%s'] * len(f['marcas']))})")
        params += f["marcas"]
    conn = get_connection()
    cursor = conn.cursor()
    try:
        # Una sola lectura de ventas_por_producto (no tiene índices); se agrupa por cliente, vendedor y mes
        cursor.execute(f"""
            SELECT `Cod. CLiente`, TRIM(Cliente), TRIM(Vendedor), LEFT(Fecha, 7), SUM(`Sub-Total 2`),
                   SUM(`Cantidad Vendida`), COUNT(DISTINCT `No. Factura`), MAX(Fecha)
            FROM ventas_por_producto WHERE {' AND '.join(where)}
            GROUP BY `Cod. CLiente`, TRIM(Cliente), TRIM(Vendedor), LEFT(Fecha, 7)
        """, params)
        filas = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()

    clientes = {}
    for cod, nombre, vendedor, mes, monto, unidades, facturas, ultima in filas:
        clave = cod if cod not in (None, "", 0) else nombre
        c = clientes.setdefault(clave, {"codigo": cod, "cliente": nombre, "meses": {}, "total": 0.0, "unidades": 0.0,
                                        "facturas": 0, "vendedores": {}, "ultima": None})
        monto = float(monto or 0)
        c["meses"][mes] = c["meses"].get(mes, 0.0) + monto
        c["total"] += monto
        c["unidades"] += float(unidades or 0)
        c["facturas"] += int(facturas or 0)
        if vendedor:
            c["vendedores"][vendedor] = c["vendedores"].get(vendedor, 0.0) + monto
        if ultima and (c["ultima"] is None or str(ultima) > str(c["ultima"])):
            c["ultima"] = str(ultima)[:10]

    lista = []
    for c in clientes.values():
        vend = sorted(c["vendedores"].items(), key=lambda x: -x[1])
        c["vendedor"] = vend[0][0] if vend else ""
        c["otros_vendedores"] = [v for v, _ in vend[1:]]
        lista.append(c)
    total_marca = sum(c["total"] for c in lista)
    vendedores = sorted({v for c in lista for v in c["vendedores"]})

    # Filtros que no cambian el total de la marca (para que el % siga siendo sobre toda la marca)
    if f["vendedor"]:
        lista = [c for c in lista if f["vendedor"] in c["vendedores"]]
    for palabra in f["q"].upper().split():
        lista = [c for c in lista if palabra in (c["cliente"] or "").upper()]
    lista.sort(key=lambda c: -c["total"])
    for i, c in enumerate(lista, start=1):
        c["posicion"] = i
        c["pct"] = c["total"] / total_marca * 100 if total_marca else 0
    seleccion = lista[:f["top"]] if f["top"] else lista
    meses = meses_entre(f["desde"], f["hasta"])
    return {
        "clientes": seleccion, "total_clientes": len(lista), "total_marca": total_marca,
        "total_top": sum(c["total"] for c in seleccion), "meses": meses, "vendedores": vendedores,
        "por_mes": {m: sum(c["meses"].get(m, 0) for c in seleccion) for m, _ in meses},
        "maximo_celda": max((v for c in seleccion for v in c["meses"].values()), default=0),
    }


def lista_marcas():
    def leer():
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT DISTINCT Marca FROM catalogo_productos WHERE Marca IS NOT NULL AND Marca <> '' ORDER BY Marca")
            return [r[0] for r in cursor.fetchall()]
        finally:
            cursor.close()
            conn.close()
    return en_cache("marcas_catalogo", leer)


@top_clientes_marca_bp.route("/top_clientes_marca")
@requiere_acceso("top_clientes_marca")
def reporte():
    f = leer_filtros(request.args)
    datos = calcular(f)
    return render_template("top_clientes_marca.html", f=f, d=datos, marcas=lista_marcas(), tamanos=TAMANOS,
                           rangos=rangos_rapidos(f["ultima"]))


@top_clientes_marca_bp.route("/top_clientes_marca/excel")
@requiere_acceso("top_clientes_marca")
def descargar_excel():
    f = leer_filtros(request.args)
    d = calcular(f)
    wb = Workbook()
    ws = wb.active
    ws.title = "Top clientes"
    marcas = ", ".join(f["marcas"]) if f["marcas"] else "Todas las marcas"
    ws["A1"] = f"Top de clientes · {marcas}"
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    ws["A2"] = (f"Del {f['desde']:%d/%m/%Y} al {f['hasta']:%d/%m/%Y} · Venta neta sin IVA (después de descuentos, sin facturas "
                f"anuladas) · {'Vendedor: ' + f['vendedor'] + ' · ' if f['vendedor'] else ''}Generado {datetime.now():%d/%m/%Y %H:%M}")
    ws["A2"].font = Font(size=9, color="52514E")
    columnas = ([("#", 5), ("Cód. cliente", 10), ("Cliente", 40), ("Vendedor", 28), ("Otros vendedores", 28)]
                + [(etiqueta, 12) for _, etiqueta in d["meses"]]
                + [("Total C$", 14), ("% de la marca", 10), ("Facturas", 9), ("Unidades", 10), ("Última compra", 12)])
    ft = 4
    for i, (nombre, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=ft, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NARANJA if nombre == "Total C$" else VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    n_meses = len(d["meses"])
    col_total = 6 + n_meses
    for n, c in enumerate(d["clientes"], start=ft + 1):
        valores = ([c["posicion"], c["codigo"], c["cliente"], c["vendedor"], ", ".join(c["otros_vendedores"]) or None]
                   + [round(c["meses"].get(m, 0), 2) or None for m, _ in d["meses"]]
                   + [round(c["total"], 2), c["pct"] / 100, c["facturas"], c["unidades"],
                      datetime.strptime(c["ultima"], "%Y-%m-%d").date() if c["ultima"] else None])
        for i, v in enumerate(valores, start=1):
            ws.cell(row=n, column=i, value=v)
        for i in range(6, col_total + 1):
            ws.cell(row=n, column=i).number_format = "#,##0.00"
        ws.cell(row=n, column=col_total).font = Font(bold=True)
        ws.cell(row=n, column=col_total + 1).number_format = "0.0%"
        ws.cell(row=n, column=col_total + 4).number_format = "dd/mm/yyyy"
    if d["clientes"]:
        tot = ft + len(d["clientes"]) + 1
        ws.cell(row=tot, column=3, value=f"TOTAL ({len(d['clientes'])} clientes)").font = Font(bold=True)
        for i in list(range(6, col_total + 1)) + [col_total + 2, col_total + 3]:
            l = get_column_letter(i)
            c = ws.cell(row=tot, column=i, value=f"=SUM({l}{ft + 1}:{l}{tot - 1})")
            c.font = Font(bold=True)
            c.number_format = "#,##0.00" if i <= col_total else "#,##0"
        ws.cell(row=tot + 1, column=3, value="Total de la marca en el período (todos los clientes)")
        ws.cell(row=tot + 1, column=col_total, value=round(d["total_marca"], 2)).number_format = "#,##0.00"
    ws.freeze_panes = ws.cell(row=ft + 1, column=4)
    ws.auto_filter.ref = f"A{ft}:{get_column_letter(len(columnas))}{ft + max(len(d['clientes']), 1)}"
    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    nombre = ("_".join(f["marcas"])[:40] if f["marcas"] else "todas").replace(" ", "_").replace("/", "-")
    return send_file(salida, as_attachment=True, download_name=f"top_clientes_{nombre}_{f['desde']}_al_{f['hasta']}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
