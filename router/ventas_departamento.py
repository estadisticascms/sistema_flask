"""Ventas por departamento: venta neta (sin IVA) de productos por departamento del cliente, con segmento,
familia y subfamilia (Clasificación de productos) y mapa de Nicaragua.

El departamento sale de la tabla clientes (cruce por nombre, como los reportes de zonas). Ese campo trae
municipios, regiones con nombres viejos y errores de escritura, así que se normaliza a los 17 departamentos."""
import difflib
import io
import json
import os
import re
from datetime import date, datetime

from flask import Blueprint, current_app, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.db import get_connection
from router.accesos import requiere_acceso
from router.cargar_catalogo import sin_tildes
from router.dashboard import ultima_fecha_con_ventas
from router.ventas_bodegas import en_cache
from router.ventas_vendedor import rangos_rapidos

ventas_departamento_bp = Blueprint("ventas_departamento", __name__)

SIN_DEPTO = "Sin departamento"       # está en la tabla clientes pero sin departamento ni municipio reconocible
OCASIONAL = "Cliente ocasional"      # el comprador no está en la tabla clientes (mostrador, nombre escrito distinto)
FUERA_DEL_MAPA = (SIN_DEPTO, OCASIONAL)
SIN_SEGMENTO = "Sin clasificar"
VERDE, NARANJA = "027266", "FE6412"

DEPARTAMENTOS = ["BOACO", "CARAZO", "CHINANDEGA", "CHONTALES", "ESTELI", "GRANADA", "JINOTEGA", "LEON", "MADRIZ", "MANAGUA",
                 "MASAYA", "MATAGALPA", "NUEVA SEGOVIA", "RACCN", "RACCS", "RIO SAN JUAN", "RIVAS"]
ALIAS = {"ATLANTICO NORTE": "RACCN", "RAAN": "RACCN", "COSTA CARIBE NORTE": "RACCN", "REGION AUTONOMA ATLANTICO NORTE": "RACCN",
         "ATLANTICO SUR": "RACCS", "RAAS": "RACCS", "COSTA CARIBE SUR": "RACCS", "REGION AUTONOMA ATLANTICO SUR": "RACCS",
         "SAN JUAN": "RIO SAN JUAN", "RIO SANJUAN": "RIO SAN JUAN", "CHON": "CHONTALES", "MN": "MANAGUA"}
# Municipios (y lugares conocidos o mal escritos que aparecen en la ficha) -> departamento
MUNICIPIOS = {
    "BOACO": "BOACO, CAMOAPA, SAN JOSE DE LOS REMATES, SAN LORENZO, SANTA LUCIA, TEUSTEPE",
    "CARAZO": "JINOTEPE, DIRIAMBA, SAN MARCOS, DOLORES, EL ROSARIO, LA CONQUISTA, LA PAZ DE CARAZO, SANTA TERESA",
    "CHINANDEGA": "CHINANDEGA, CHICHIGALPA, CORINTO, EL REALEJO, EL VIEJO, VIEJO, POSOLTEGA, PUERTO MORAZAN, "
                  "SANTO TOMAS DEL NORTE, SOMOTILLO, VILLANUEVA, CINCO PINOS, SAN FRANCISCO DEL NORTE, SAN PEDRO DEL NORTE, GUANACASTILLO",
    "CHONTALES": "JUIGALPA, ACOYAPA, COMALAPA, CUAPA, EL CORAL, LA LIBERTAD, SAN PEDRO DE LOVAGO, SANTO DOMINGO, SANTO TOMAS, VILLA SANDINO",
    "ESTELI": "ESTELI, CONDEGA, LA TRINIDAD, PUEBLO NUEVO, SAN JUAN DE LIMAY, SAN NICOLAS",
    "GRANADA": "GRANADA, DIRIA, DIRIOMO, NANDAIME",
    "JINOTEGA": "JINOTEGA, EL CUA, LA CONCORDIA, SAN JOSE DE BOCAY, SAN RAFAEL DEL NORTE, SAN SEBASTIAN DE YALI, "
                "SANTA MARIA DE PANTASMA, WIWILI DE JINOTEGA",
    "LEON": "LEON, ACHUAPA, EL JICARAL, EL SAUCE, LA PAZ CENTRO, LARREYNAGA, MALPAISILLO, NAGAROTE, QUEZALGUAQUE, QUETZALGUAQUE, "
            "SANTA ROSA DEL PENON, TELICA, PUERTO SANDINO",
    "MADRIZ": "MADRIZ, SOMOTO, LAS SABANAS, PALACAGUINA, SAN JOSE DE CUSMAPA, SAN JUAN DE RIO COCO, SAN JUAN RIO DE COCO, "
              "SAN J RIO COCOC, SAN LUCAS, TELPANECA, TOTOGALPA, YALAGUINA",
    "MANAGUA": "MANAGUA, CIUDAD SANDINO, EL CRUCERO, MATEARE, SAN FRANCISCO LIBRE, SAN RAFAEL DEL SUR, SAN RAFAEL SUR, "
               "TICUANTEPE, TIPITAPA, VILLA EL CARMEN",
    "MASAYA": "MASAYA, CATARINA, LA CONCEPCION, LA CONCHA, MASATEPE, NANDASMO, NINDIRI, NIQUINOHOMO, SAN JUAN DE ORIENTE, TISMA",
    "MATAGALPA": "MATAGALPA, CIUDAD DARIO, DARIO, EL TUMA LA DALIA, LA DALIA, ESQUIPULAS, MATIGUAS, MUY MUY, RANCHO GRANDE, "
                 "RIO BLANCO, SAN DIONISIO, SAN ISIDRO, SAN RAMON, SEBACO, TERRABONA",
    "NUEVA SEGOVIA": "NUEVA SEGOVIA, OCOTAL, CIUDAD ANTIGUA, DIPILTO, EL JICARO, JALAPA, MACUELIZO, MOZONTE, MURRA, QUILALI, "
                     "SAN FERNANDO, WIWILI DE NUEVA SEGOVIA",
    "RACCN": "PUERTO CABEZAS, PUERTO CABEZA, BILWI, BONANZA, MULUKUKU, PRINZAPOLKA, ROSITA, SIUNA, WASLALA, WASPAM, WASPAN",
    "RACCS": "BLUEFIELDS, CORN ISLAND, DESEMBOCADURA DE RIO GRANDE, EL AYOTE, EL RAMA, RAMA, EL TORTUGUERO, KUKRA HILL, "
             "LA CRUZ DE RIO GRANDE, LAGUNA DE PERLAS, MUELLE DE LOS BUEYES, NUEVA GUINEA, PAIWAS",
    "RIO SAN JUAN": "SAN CARLOS, EL ALMENDRO, EL CASTILLO, MORRITO, SAN JUAN DEL NORTE, SAN MIGUELITO",
    "RIVAS": "RIVAS, ALTAGRACIA, BELEN, BUENOS AIRES, CARDENAS, MOYOGALPA, POTOSI, SAN JORGE, SAN JUAN DEL SUR, TOLA, OMETEPE",
}
LUGARES = {d: d for d in DEPARTAMENTOS}
LUGARES.update(ALIAS)
for _depto, _lista in MUNICIPIOS.items():
    for _m in _lista.split(","):
        LUGARES.setdefault(_m.strip(), _depto)


def normal(texto):
    t = sin_tildes(texto or "").upper()
    t = re.sub(r"[^A-Z ]", " ", t)
    return " ".join(t.split())


def departamento_de(*textos):
    """Primer texto que se pueda reconocer como departamento (o municipio de uno)."""
    for texto in textos:
        t = normal(texto)
        if not t or t in ("XXX", "NA", "N A"):
            continue
        if t in LUGARES:
            return LUGARES[t]
        # empieza con un lugar conocido ("MANAGUA89213002", "MATAGALPA CENTRO")
        for lugar in sorted(LUGARES, key=len, reverse=True):
            if len(lugar) >= 4 and t.startswith(lugar):
                return LUGARES[lugar]
        # errores de escritura (MANAGAU, JUGALPA, DIRAMBA) o nombres cortados (CHINANDE)
        parecido = difflib.get_close_matches(t, LUGARES.keys(), n=1, cutoff=0.82)
        if parecido:
            return LUGARES[parecido[0]]
        cortado = [l for l in LUGARES if len(t) >= 5 and l.startswith(t)]
        if len(cortado) == 1:
            return LUGARES[cortado[0]]
    return None


def departamentos_clientes():
    """Nombre de cliente normalizado -> departamento, o SIN_DEPTO si está en la tabla sin ubicación (10 minutos en memoria).
    Quien no aparece aquí no está en la tabla clientes: es un cliente ocasional."""
    def leer():
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT nombre, departamento, municipio FROM clientes")
            resultado = {}
            for nombre, depto, municipio in cursor.fetchall():
                clave = " ".join((nombre or "").upper().split())
                d = departamento_de(depto, municipio) or SIN_DEPTO
                if clave and (clave not in resultado or resultado[clave] == SIN_DEPTO):
                    resultado[clave] = d
            return resultado
        finally:
            cursor.close()
            conn.close()
    return en_cache("vd_deptos_clientes", leer)


def cartera_por_departamento():
    """Clientes de la tabla clientes por departamento (nombres distintos), total y activos."""
    def leer():
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT nombre, departamento, municipio, estado FROM clientes")
            por_nombre = {}
            for nombre, depto, municipio, estado in cursor.fetchall():
                clave = " ".join((nombre or "").upper().split())
                if not clave:
                    continue
                d = departamento_de(depto, municipio) or SIN_DEPTO
                actual = por_nombre.get(clave)
                # un mismo nombre puede estar dos veces (activo e inactivo): cuenta una vez, activo si alguno lo está
                if actual is None or (actual[0] == SIN_DEPTO and d != SIN_DEPTO):
                    por_nombre[clave] = (d, estado == "Activo" or (actual is not None and actual[1]))
                elif estado == "Activo":
                    por_nombre[clave] = (actual[0], True)
            cartera = {}
            for d, activo in por_nombre.values():
                c = cartera.setdefault(d, {"total": 0, "activos": 0})
                c["total"] += 1
                c["activos"] += activo
            return cartera
        finally:
            cursor.close()
            conn.close()
    return en_cache("vd_cartera", leer)


def clasificacion():
    def leer():
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT codigo, segmento, familia, subfamilia FROM clase_prod")
            return {c: ((s or "").strip() or SIN_SEGMENTO, (f or "").strip() or SIN_SEGMENTO, (sf or "").strip() or SIN_SEGMENTO)
                    for c, s, f, sf in cursor.fetchall()}
        finally:
            cursor.close()
            conn.close()
    return en_cache("vd_clasificacion", leer)


def mapa():
    def leer():
        with open(os.path.join(current_app.static_folder, "mapa_nicaragua.json"), encoding="utf-8") as f:
            return json.load(f)
    return en_cache("vd_mapa", leer, segundos=86400)


def fecha_param(valor, defecto):
    try:
        return datetime.strptime(valor or "", "%Y-%m-%d").date()
    except ValueError:
        return defecto


def leer_filtros(args):
    ultima = ultima_fecha_con_ventas() or date.today()
    if isinstance(ultima, datetime):
        ultima = ultima.date()
    hasta = fecha_param(args.get("hasta"), ultima)
    desde = min(fecha_param(args.get("desde"), ultima.replace(day=1)), hasta)
    return {"desde": desde, "hasta": hasta, "ultima": ultima,
            "depto": (args.get("depto") or "").strip().upper(), "segmento": (args.get("segmento") or "").strip()}


def calcular(f):
    conn = get_connection()
    cursor = conn.cursor()
    try:
        # Una sola lectura de ventas_por_producto (no tiene índices)
        cursor.execute("""
            SELECT TRIM(Cliente), Codigo, SUM(`Sub-Total 2`), SUM(`Cantidad Vendida`)
            FROM ventas_por_producto WHERE Fecha BETWEEN %s AND %s AND Estatus <> 'Anulado'
            GROUP BY TRIM(Cliente), Codigo
        """, (f["desde"].isoformat(), f["hasta"].isoformat()))
        filas = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()

    deptos, clase = departamentos_clientes(), clasificacion()
    por_depto = {}      # depto -> {total, segmentos:{}, clientes:set}
    detalle = {}        # (depto, segmento, familia, subfamilia) -> [monto, unidades]
    total = 0.0
    for cliente, codigo, monto, unidades in filas:
        monto, unidades = float(monto or 0), float(unidades or 0)
        depto = deptos.get(" ".join((cliente or "").upper().split()), OCASIONAL)
        seg, fam, sub = clase.get(codigo, (SIN_SEGMENTO, SIN_SEGMENTO, SIN_SEGMENTO))
        d = por_depto.setdefault(depto, {"total": 0.0, "segmentos": {}, "clientes": set()})
        d["total"] += monto
        d["segmentos"][seg] = d["segmentos"].get(seg, 0.0) + monto
        d["clientes"].add(cliente)
        x = detalle.setdefault((depto, seg, fam, sub), [0.0, 0.0])
        x[0] += monto
        x[1] += unidades
        total += monto

    totales_seg = {}
    for d in por_depto.values():
        for s, v in d["segmentos"].items():
            totales_seg[s] = totales_seg.get(s, 0.0) + v
    segmentos = [s for s, _ in sorted(totales_seg.items(), key=lambda x: -x[1])]

    cartera = cartera_por_departamento()
    tabla = []
    for nombre, d in por_depto.items():
        c = cartera.get(nombre, {"total": 0, "activos": 0})
        tabla.append({"depto": nombre, "total": d["total"], "segmentos": d["segmentos"], "clientes": len(d["clientes"]),
                      "cartera": c["total"], "activos": c["activos"],
                      "pct": d["total"] / total * 100 if total else 0,
                      "principal": max(d["segmentos"].items(), key=lambda x: x[1])[0] if d["segmentos"] else None})
    for nombre, c in cartera.items():
        if nombre not in por_depto:
            tabla.append({"depto": nombre, "total": 0.0, "segmentos": {}, "clientes": 0, "cartera": c["total"],
                          "activos": c["activos"], "pct": 0, "principal": None})
    tabla.sort(key=lambda x: (x["depto"] in FUERA_DEL_MAPA, x["depto"] == OCASIONAL, -x["total"]))
    principales = segmentos[:5]   # columnas de la tabla; el resto va en "Otros"
    for x in tabla:
        x["otros"] = x["total"] - sum(x["segmentos"].get(sg, 0) for sg in principales)
    otros_total = total - sum(totales_seg.get(sg, 0) for sg in principales)

    # Familia / subfamilia (del departamento y segmento elegidos, o de todo)
    familias = {}
    for (depto, seg, fam, sub), (monto, unidades) in detalle.items():
        if (f["depto"] and depto.upper() != f["depto"]) or (f["segmento"] and seg != f["segmento"]):
            continue
        fa = familias.setdefault(fam, {"familia": fam, "monto": 0.0, "unidades": 0.0, "subfamilias": {}})
        fa["monto"] += monto
        fa["unidades"] += unidades
        sf = fa["subfamilias"].setdefault(sub, {"subfamilia": sub, "monto": 0.0, "unidades": 0.0})
        sf["monto"] += monto
        sf["unidades"] += unidades
    lista_fam = sorted(familias.values(), key=lambda x: -x["monto"])
    for fa in lista_fam:
        fa["subfamilias"] = sorted(fa["subfamilias"].values(), key=lambda x: -x["monto"])
    total_fam = sum(fa["monto"] for fa in lista_fam)

    sin = por_depto.get(SIN_DEPTO, {"total": 0.0})["total"]
    ocasional = por_depto.get(OCASIONAL, {"total": 0.0})["total"]
    return {"tabla": tabla, "segmentos": segmentos, "principales": principales, "otros_total": otros_total, "totales_seg": totales_seg, "total": total, "sin_depto": sin, "ocasional": ocasional,
            "familias": lista_fam, "total_familias": total_fam, "detalle": detalle,
            "mapa_datos": {d["depto"]: {"total": round(d["total"], 2), "segmentos": {s: round(v, 2) for s, v in d["segmentos"].items()},
                                        "clientes": d["clientes"], "cartera": d["cartera"], "activos": d["activos"]}
                           for d in tabla if d["depto"] not in FUERA_DEL_MAPA and d["total"] > 0},
            "cartera_total": sum(c["total"] for c in cartera.values()),
            "compraron_total": sum(x["clientes"] for x in tabla)}


@ventas_departamento_bp.route("/ventas_departamento")
@requiere_acceso("ventas_departamento")
def reporte():
    f = leer_filtros(request.args)
    d = calcular(f)
    return render_template("ventas_departamento.html", f=f, d=d, mapa=mapa(), rangos=rangos_rapidos(f["ultima"]),
                           departamentos=DEPARTAMENTOS, sin_depto=SIN_DEPTO, ocasional=OCASIONAL, fuera_del_mapa=FUERA_DEL_MAPA)


@ventas_departamento_bp.route("/ventas_departamento/excel")
@requiere_acceso("ventas_departamento")
def descargar_excel():
    f = leer_filtros(request.args)
    d = calcular(f)
    wb = Workbook()
    titulo = f"Ventas por departamento del {f['desde']:%d/%m/%Y} al {f['hasta']:%d/%m/%Y}"
    nota = "Venta neta sin IVA (después de descuentos, sin facturas anuladas). Departamento según la ficha del cliente."

    def encabezado(ws, columnas, fila=4):
        for i, (nombre, ancho) in enumerate(columnas, start=1):
            c = ws.cell(row=fila, column=i, value=nombre)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=NARANJA if nombre == "Total C$" else VERDE)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = ancho

    ws = wb.active
    ws.title = "Departamento y segmento"
    ws["A1"], ws["A2"] = titulo, nota
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    ws["A2"].font = Font(size=9, color="52514E")
    columnas = [("Departamento", 20)] + [(s, 14) for s in d["segmentos"]] + [("Total C$", 15), ("% del total", 10), ("Clientes que compraron", 11),
                                                                                ("Clientes en cartera", 10), ("Activos en cartera", 10), ("% que compró", 10)]
    encabezado(ws, columnas)
    for n, x in enumerate(d["tabla"], start=5):
        valores = ([x["depto"]] + [round(x["segmentos"].get(s, 0), 2) or None for s in d["segmentos"]]
                   + [round(x["total"], 2), x["pct"] / 100, x["clientes"], x["cartera"], x["activos"],
                      x["clientes"] / x["cartera"] if x["cartera"] else None])
        pct_cols = {len(d["segmentos"]) + 3, len(valores)}
        for i, v in enumerate(valores, start=1):
            c = ws.cell(row=n, column=i, value=v)
            c.number_format = "0.0%" if i in pct_cols else ("#,##0" if i > len(d["segmentos"]) + 3 else "#,##0.00")
        ws.cell(row=n, column=len(d["segmentos"]) + 2).font = Font(bold=True)
    tot = 5 + len(d["tabla"])
    ws.cell(row=tot, column=1, value="TOTAL").font = Font(bold=True)
    for i in range(2, len(d["segmentos"]) + 3):
        l = get_column_letter(i)
        c = ws.cell(row=tot, column=i, value=f"=SUM({l}5:{l}{tot - 1})")
        c.font = Font(bold=True)
        c.number_format = "#,##0.00"
    ws.freeze_panes = "B5"

    ws2 = wb.create_sheet("Familia y subfamilia")
    ws2["A1"], ws2["A2"] = titulo + " · detalle", nota
    ws2["A1"].font = Font(bold=True, size=14, color=VERDE)
    ws2["A2"].font = Font(size=9, color="52514E")
    encabezado(ws2, [("Departamento", 20), ("Segmento", 22), ("Familia", 30), ("Subfamilia", 36), ("Total C$", 15), ("Unidades", 11)])
    filas = sorted(d["detalle"].items(), key=lambda x: (x[0][0] in FUERA_DEL_MAPA, x[0][0], -x[1][0]))
    for n, ((depto, seg, fam, sub), (monto, unidades)) in enumerate(filas, start=5):
        for i, v in enumerate([depto, seg, fam, sub, round(monto, 2), unidades], start=1):
            ws2.cell(row=n, column=i, value=v)
        ws2.cell(row=n, column=5).number_format = "#,##0.00"
        ws2.cell(row=n, column=6).number_format = "#,##0.##"
    ws2.freeze_panes = "A5"
    ws2.auto_filter.ref = f"A4:F{4 + max(len(filas), 1)}"

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=f"ventas_departamento_{f['desde']}_al_{f['hasta']}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
