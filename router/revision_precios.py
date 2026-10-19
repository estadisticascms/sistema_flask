"""Revisión de precios de productos con existencia:
- Margen bajo: el precio de referencia deja un margen sobre la venta igual o menor al umbral (30%).
- Sin movimiento con margen alto: no se venden desde hace N meses y su margen supera el umbral (50%),
  candidatos a rebajar el precio."""
import io
from datetime import datetime, timedelta

from flask import Blueprint, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.db import get_connection
from router.accesos import requiere_acceso
from router.sugerido_compras import a_fecha
from router.productos_agotarse import lista_marcas
from router.productos_sin_movimiento import DIAS_MES, MESES_SIN_VENTA, PRECIOS, TC_DEFECTO, margen

revision_precios_bp = Blueprint("revision_precios", __name__)

VISTAS = {"bajo": "Margen bajo", "alto": "Sin movimiento con margen alto"}
VERDE, NARANJA = "027266", "FE6412"
POR_PAGINA = 200
# Márgenes fuera de este rango casi siempre son un error de costo o precio en el catálogo
# (costo por caja y precio por unidad, costo en otra moneda, precio sin actualizar...)
DUDOSO_MIN, DUDOSO_MAX = -1.0, 0.95


def numero(valor, defecto, minimo, maximo):
    try:
        v = float(valor)
    except (TypeError, ValueError):
        return defecto
    return v if minimo <= v <= maximo else defecto


def leer_filtros(args):
    vista = args.get("vista")
    meses = int(numero(args.get("meses"), 6, 1, 120))
    return {"vista": vista if vista in VISTAS else "bajo",
            "margen_bajo": numero(args.get("margen_bajo"), 30, -100, 100),
            "margen_alto": numero(args.get("margen_alto"), 50, 0, 100),
            "precio": int(numero(args.get("precio"), 1, 1, 4)),
            "meses": meses if meses in MESES_SIN_VENTA else 6,
            "tc": numero(args.get("tc"), TC_DEFECTO, 1, 1000),
            "marca": (args.get("marca") or "").strip(), "q": (args.get("q") or "").strip()}


def calcular(f):
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT Codigo_producto, SUM(Cant_Final) FROM existencias_bodega "
                       "GROUP BY Codigo_producto HAVING SUM(Cant_Final) > 0")
        existencias = {r[0]: float(r[1] or 0) for r in cursor.fetchall()}

        # Última venta (salidas por factura del inventario, usa el índice producto/fecha/tipo)
        cursor.execute("SELECT /*+ INDEX(transacciones_inventario idx_transacciones_producto_fecha) */ "
                       "codigo_producto, MAX(fecha) FROM transacciones_inventario "
                       "WHERE tipo_transaccion = 'SALIDA' GROUP BY codigo_producto")
        salidas = [(cod, a_fecha(fecha)) for cod, fecha in cursor.fetchall() if fecha]
        ultima_venta = {cod: fecha for cod, fecha in salidas if cod in existencias}
        fin = max((fecha for _, fecha in salidas), default=None) or datetime.today().date()
        corte = fin - timedelta(days=round(f["meses"] * DIAS_MES))

        # Ventas reales del período: unidades, venta neta (después de descuento, sin IVA) y su costo
        cursor.execute("""
            SELECT Codigo, SUM(`Cantidad Vendida`), SUM(`Sub-Total 2`), SUM(`Cantidad Vendida` * `Costo Unit.`)
            FROM ventas_por_producto WHERE Fecha > %s AND Estatus <> 'Anulado' GROUP BY Codigo
        """, (corte.isoformat(),))
        ventas = {r[0]: (float(r[1] or 0), float(r[2] or 0), float(r[3] or 0)) for r in cursor.fetchall() if r[0] in existencias}

        where, params = ["p.`Codigo Producto` IS NOT NULL"], []
        if f["marca"]:
            where.append("p.Marca = %s"); params.append(f["marca"])
        for palabra in f["q"].split():
            where.append("(p.`Codigo Producto` LIKE %s OR p.Nombre LIKE %s OR p.`Numero Parte` LIKE %s)")
            params += [f"%{palabra}%"] * 3
        cursor.execute(f"""
            SELECT /*+ NO_INDEX(c) */ p.`Codigo Producto`, p.Nombre, p.`Numero Parte`, p.Marca, p.Moneda,
                   p.`Costo Promedio`, p.`Ultimo Costo`, p.`Precio 1`, p.`Precio 2`, p.`Precio 3`, p.`Precio 4`, c.familia
            FROM catalogo_productos p LEFT JOIN clase_prod c ON c.codigo = p.`Codigo Producto`
            WHERE {' AND '.join(where)}
        """, params)
        catalogo = [r for r in cursor.fetchall() if r[0] in existencias]

        # Último ingreso por compra: fecha, cantidad (suma de ese día) y costo unitario de esa compra
        cursor.execute("SELECT codigo_producto, fecha, cantidad, costo_unitario_mn FROM transacciones_inventario "
                       "WHERE tipo_transaccion = 'ENTRADA' AND tipo_documento_origen = 'ORDEN COMPRA'")
        ingresos = {}
        for cod, fecha, cant, costo_u in cursor.fetchall():
            if cod not in existencias or not fecha:
                continue
            fecha, cant, costo_u = a_fecha(fecha), float(cant or 0), float(costo_u or 0)
            actual = ingresos.get(cod)
            if actual is None or fecha > actual["fecha"]:
                ingresos[cod] = {"fecha": fecha, "cantidad": cant, "monto": cant * costo_u}
            elif fecha == actual["fecha"]:
                actual["cantidad"] += cant
                actual["monto"] += cant * costo_u
    finally:
        cursor.close()
        conn.close()

    bajo, alto, sin_datos = [], [], 0
    ref = f["precio"] - 1
    for cod, nombre, parte, marca, moneda, costo_prom, ultimo_costo, p1, p2, p3, p4, familia in catalogo:
        costo_prom, ultimo_costo = float(costo_prom or 0), float(ultimo_costo or 0)
        costo = ultimo_costo or costo_prom
        factor = f["tc"] if (moneda or "U$") != "C$" else 1.0
        precios = []
        for nombre_precio, p in zip(PRECIOS, (p1, p2, p3, p4)):
            cs = float(p or 0) * factor
            precios.append({"nombre": nombre_precio, "usd": float(p or 0) if factor != 1 else None, "cs": cs,
                            "margen": margen(cs, costo) if cs > 0 and costo > 0 else None})
        m = precios[ref]["margen"]
        if m is None:   # sin costo o sin precio de referencia: no se puede revisar el margen
            sin_datos += 1
            continue
        uv = ultima_venta.get(cod)
        ing = ingresos.get(cod)
        unidades, venta, costo_venta = ventas.get(cod, (0.0, 0.0, 0.0))
        x = {"codigo": cod, "nombre": nombre, "numero_parte": parte, "marca": marca, "familia": familia,
             "existencia": existencias[cod], "costo": costo, "costo_promedio": costo_prom, "ultimo_costo": ultimo_costo,
             "valor": existencias[cod] * (costo_prom or ultimo_costo), "precios": precios, "precio_ref": precios[ref],
             "margen": m, "ultima_venta": uv, "meses_sin_venta": (fin - uv).days / DIAS_MES if uv else None,
             "unidades": unidades, "venta": venta,
             "precio_real": venta / unidades if unidades > 0 else None,
             "margen_real": (venta - costo_venta) / venta if venta > 0 else None,
             "factor": factor, "dudoso": not (DUDOSO_MIN <= m <= DUDOSO_MAX),
             "ultimo_ingreso": ing["fecha"] if ing else None, "cantidad_ingreso": ing["cantidad"] if ing else None,
             "costo_ingreso": ing["monto"] / ing["cantidad"] if ing and ing["cantidad"] else None,
             "meses_ingreso": (fin - ing["fecha"]).days / DIAS_MES if ing else None}
        sin_movimiento = uv is None or uv <= corte
        if m <= f["margen_bajo"] / 100:
            # Precio de referencia que daría el margen mínimo: costo / (1 − margen)
            objetivo = f["margen_bajo"] / 100
            x["precio_minimo"] = costo / (1 - objetivo) if objetivo < 1 else None
            bajo.append(x)
        if sin_movimiento and m > f["margen_alto"] / 100:
            alto.append(x)

    # Los posibles errores de catálogo van al final para no tapar los casos reales
    bajo.sort(key=lambda x: (x["dudoso"], x["margen"], -x["unidades"]))
    alto.sort(key=lambda x: (x["dudoso"], -x["margen"], -x["valor"]))
    return {"bajo": bajo, "alto": alto, "sin_datos": sin_datos, "fin": fin, "corte": corte}


@revision_precios_bp.route("/revision_precios")
@requiere_acceso("revision_precios")
def reporte():
    f = leer_filtros(request.args)
    d = calcular(f)
    filas = d[f["vista"]]
    paginas = max(1, -(-len(filas) // POR_PAGINA))
    try:
        pagina = min(max(1, int(request.args.get("pagina", 1))), paginas)
    except ValueError:
        pagina = 1
    bajo = d["bajo"]
    return render_template(
        "revision_precios.html", f=f, d=d, vistas=VISTAS, precios=PRECIOS, opciones_meses=MESES_SIN_VENTA,
        marcas=lista_marcas(), filas=filas[(pagina - 1) * POR_PAGINA: pagina * POR_PAGINA], total=len(filas),
        pagina=pagina, paginas=paginas, desde=(pagina - 1) * POR_PAGINA,
        negativos=sum(1 for x in bajo if x["margen"] < 0),
        bajo_con_venta=sum(1 for x in bajo if x["unidades"] > 0),
        dudosos=sum(1 for x in filas if x["dudoso"]),
        valor_alto=sum(x["valor"] for x in d["alto"]))


def hoja(ws, f, d, vista):
    bajo = vista == "bajo"
    ref = PRECIOS[f["precio"] - 1]
    if bajo:
        ws["A1"] = f"Productos con margen igual o menor a {f['margen_bajo']:g}% ({ref})"
    else:
        ws["A1"] = f"Productos sin ventas en {f['meses']} meses o más con margen mayor a {f['margen_alto']:g}% ({ref})"
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    ws["A2"] = (f"Margen = (precio − último costo) / precio · Precios del catálogo en U$ × TC (celda B3) · "
                f"Ventas del {d['corte'] + timedelta(days=1):%d/%m/%Y} al {d['fin']:%d/%m/%Y} · Generado {datetime.now():%d/%m/%Y %H:%M}")
    ws["A2"].font = Font(size=9, color="52514E")
    ws["A3"], ws["B3"] = "Tipo de cambio", f["tc"]
    ws["A3"].font = Font(bold=True)

    columnas = [("Código", 11), ("Nº parte", 20), ("Descripción", 36), ("Marca", 14), ("Familia", 20),
                ("Existencia", 10), ("Costo promedio C$", 13), ("Último costo C$", 13), ("Valor inventario C$", 15)]
    for p in PRECIOS:
        columnas += [(f"{p} U$", 10), (f"{p} C$", 11), (f"Margen {p}", 10)]
    if bajo:
        columnas += [(f"{ref} mínimo U$ ({f['margen_bajo']:g}%)", 13), (f"{ref} mínimo C$ ({f['margen_bajo']:g}%)", 13)]
    if bajo:
        columnas += [(f"Unid. vendidas {f['meses']}m", 10), ("Precio promedio facturado C$", 13), ("Margen real", 10)]
    columnas += [("Último ingreso", 12), ("Cant. último ingreso", 10), ("Costo unit. último ingreso C$", 13),
                 ("Última venta", 12), ("Meses sin venta", 10), ("Observación", 22)]
    ft = 5
    for i, (nombre, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=ft, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NARANJA if nombre.startswith(f"Margen {ref}") or "mínimo" in nombre else VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[ft].height = 45

    filas = d[vista]
    for n, x in enumerate(filas, start=ft + 1):
        valores = [x["codigo"], x["numero_parte"], x["nombre"], x["marca"], x["familia"], x["existencia"],
                   round(x["costo_promedio"], 4), round(x["ultimo_costo"], 4), round(x["valor"], 2)]
        for col, v in enumerate(valores, start=1):
            ws.cell(row=n, column=col, value=v)
        for col in (7, 8, 9):
            ws.cell(row=n, column=col).number_format = "#,##0.00"
        costo = f"IF(H{n}>0,H{n},G{n})"
        col = len(valores) + 1
        for p in x["precios"]:
            usd, cs = get_column_letter(col), get_column_letter(col + 1)
            if p["usd"] is not None:
                ws.cell(row=n, column=col, value=round(p["usd"], 4)).number_format = "#,##0.00"
                ws.cell(row=n, column=col + 1, value=f"={usd}{n}*$B$3")
            else:
                ws.cell(row=n, column=col + 1, value=round(p["cs"], 2))
            ws.cell(row=n, column=col + 1).number_format = "#,##0.00"
            c = ws.cell(row=n, column=col + 2, value=f'=IF({cs}{n}>0,({cs}{n}-{costo})/{cs}{n},"")')
            c.number_format = "0.0%"
            if p is x["precio_ref"]:
                c.font = Font(bold=True)
            col += 3
        if bajo:
            objetivo = f["margen_bajo"] / 100
            minimo_cs = f"{costo}/(1-{objetivo})" if objetivo < 1 else '""'
            ws.cell(row=n, column=col + 1, value=f"={minimo_cs}").number_format = "#,##0.00"
            ws.cell(row=n, column=col, value=f"={get_column_letter(col + 1)}{n}/$B$3" if x["factor"] != 1
                    else None).number_format = "#,##0.00"
            col += 2
        reales = ((x["unidades"], "#,##0"), (round(x["precio_real"], 2) if x["precio_real"] else None, "#,##0.00"),
                  (x["margen_real"], "0.0%")) if bajo else ()
        ingreso = ((x["ultimo_ingreso"], "dd/mm/yyyy"), (x["cantidad_ingreso"], "#,##0"),
                   (round(x["costo_ingreso"], 4) if x["costo_ingreso"] is not None else None, "#,##0.00"))
        for v, fmt in reales + ingreso + ((x["ultima_venta"], "dd/mm/yyyy"),
                       (round(x["meses_sin_venta"], 1) if x["meses_sin_venta"] is not None else "Sin venta", "0.0"),
                       ("Revisar costo/precio" if x["dudoso"] else None, "General")):
            ws.cell(row=n, column=col, value=v).number_format = fmt
            col += 1
    ws.freeze_panes = ws.cell(row=ft + 1, column=4)
    ws.auto_filter.ref = f"A{ft}:{get_column_letter(len(columnas))}{ft + max(len(filas), 1)}"


@revision_precios_bp.route("/revision_precios/excel")
@requiere_acceso("revision_precios")
def descargar_excel():
    f = leer_filtros(request.args)
    d = calcular(f)
    wb = Workbook()
    hoja(wb.active, f, d, "bajo")
    wb.active.title = f"Margen hasta {f['margen_bajo']:g}%"
    hoja(wb.create_sheet(f"Sin mov. margen > {f['margen_alto']:g}%"), f, d, "alto")
    if f["vista"] == "alto":
        wb.active = 1
    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    nombre = f"revision_precios_{f['marca'].replace(' ', '_') + '_' if f['marca'] else ''}{datetime.now():%Y-%m-%d}.xlsx"
    return send_file(salida, as_attachment=True, download_name=nombre,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
