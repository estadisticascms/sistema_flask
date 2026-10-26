"""Análisis de cartera: ventas de crédito contra recuperación, por mes y por cliente, con antigüedad de saldos,
límite y días de crédito, comportamiento de pago y recomendación (aumentar, mantener, revisar o suspender crédito).

Fuentes: ventas_historicas (facturas de crédito, total con IVA y fecha de vencimiento), recuperacion (pagos por
factura y recibo), devoluciones (de crédito, cerradas) y clientes (límite, términos y retenciones).
Saldo de una factura = total − pagado − devuelto; si lo que falta cabe en la retención (2% IR / 1% alcaldía,
hasta 3% del subtotal) se considera pagada: el cliente retuvo, no debe. La recuperación empieza en 2026,
por eso el análisis toma facturas desde enero 2026 (las de 2025 solo cuentan como pagos recibidos)."""
import io
import re
from datetime import date, datetime, timedelta

from flask import Blueprint, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.db import get_connection
from router.accesos import requiere_acceso
from router.ventas_bodegas import en_cache

analisis_cartera_bp = Blueprint("analisis_cartera", __name__)

INICIO_DATOS = date(2026, 1, 1)      # primer mes con recuperación cargada
RETENCION_MAX = 0.031                # 2% IR + 1% alcaldía (+ redondeo) sobre subtotal − descuento
TOLERANCIA_DIAS = 5                  # días de gracia para considerar un pago "a tiempo"
MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
TRAMOS = [("Por vencer", None, 0), ("1–30 días", 1, 30), ("31–60 días", 31, 60), ("61–90 días", 61, 90), ("Más de 90 días", 91, 10 ** 6)]
RECOMENDACIONES = {
    "suspender": ("Suspender crédito", "rojo", "bi-slash-circle"),
    "revisar": ("Reducir / revisar", "naranja", "bi-exclamation-triangle"),
    "aumentar": ("Puede aumentar", "verde", "bi-arrow-up-circle"),
    "sin_limite": ("Sin límite asignado", "gris", "bi-question-circle"),
    "sin_uso": ("No usa el crédito", "gris", "bi-moon"),
    "mantener": ("Mantener", "azul", "bi-check-circle"),
}
VERDE, NARANJA = "027266", "FE6412"


def clave(nombre):
    return " ".join((nombre or "").upper().split())


def plazo_de(terminos):
    m = re.search(r"(\d+)", terminos or "")
    return int(m.group(1)) if m else None


def fecha_param(valor, defecto):
    try:
        return datetime.strptime(valor or "", "%Y-%m-%d").date()
    except ValueError:
        return defecto


def leer_filtros(args):
    datos = fecha_datos()
    hasta = fecha_param(args.get("hasta"), datos)
    desde = max(min(fecha_param(args.get("desde"), INICIO_DATOS), hasta), INICIO_DATOS)
    reco = args.get("reco") if args.get("reco") in RECOMENDACIONES else ""
    return {"desde": desde, "hasta": hasta, "corte": datos, "vendedor": (args.get("vendedor") or "").strip(),
            "q": (args.get("q") or "").strip(), "reco": reco}


def fecha_datos():
    """Fecha de corte: el último día con pagos o ventas cargados."""
    def leer():
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT GREATEST((SELECT MAX(fechapago) FROM recuperacion), (SELECT MAX(Fecha) FROM ventas_historicas))")
            return cursor.fetchone()[0] or date.today()
        finally:
            cursor.close()
            conn.close()
    return en_cache("cartera_corte", leer, segundos=300)


def leer_datos(f):
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("""
            SELECT factura, Fecha AS fecha, fecha_vencimiento AS vence, TRIM(cliente) AS cliente, TRIM(vendedor) AS vendedor,
                   Total AS total, SubTotal - Descuento AS base, estado_factura AS estado
            FROM ventas_historicas
            WHERE tipo_venta = 'CREDITO' AND Fecha BETWEEN %s AND %s AND COALESCE(estado_factura, '') <> 'Anulado'
        """, (f["desde"], f["hasta"]))
        facturas = cursor.fetchall()
        # Pagos recibidos en el período (también de facturas anteriores) y todos los pagos de las facturas del período
        cursor.execute("""
            SELECT nfactura, nrecibo, fechafact, fechapago, drecuperado, mpagado, TRIM(cliente) AS cliente, TRIM(vendedor) AS vendedor
            FROM recuperacion WHERE fechapago >= %s
        """, (f["desde"],))
        pagos = cursor.fetchall()
        cursor.execute("SELECT factura, SUM(totaldv) AS dv FROM devoluciones WHERE estado = 'Cerrado' AND tipofact = 'CREDITO' "
                       "AND fechadv >= %s GROUP BY factura", (f["desde"],))
        devoluciones = {r["factura"]: float(r["dv"] or 0) for r in cursor.fetchall()}
        cursor.execute("SELECT factura, TRIM(cliente) AS cliente FROM ventas_historicas WHERE tipo_venta = 'CREDITO' AND Fecha >= %s",
                       (f["desde"] - timedelta(days=730),))
        cliente_de_factura = {r["factura"]: r["cliente"] for r in cursor.fetchall()}
        cursor.execute("SELECT nombre, estado, terminos_pago, limite_credito, vendedor, ret_dgi, ret_alma FROM clientes")
        fichas = {}
        for r in cursor.fetchall():
            k = clave(r["nombre"])
            if k and (k not in fichas or r["estado"] == "Activo"):
                fichas[k] = r
    finally:
        cursor.close()
        conn.close()
    return facturas, pagos, devoluciones, cliente_de_factura, fichas


def calcular(f):
    facturas, pagos, devoluciones, cliente_de_factura, fichas = leer_datos(f)
    corte = f["corte"]
    pagado_factura = {}
    for p in pagos:
        pagado_factura[p["nfactura"]] = pagado_factura.get(p["nfactura"], 0.0) + float(p["mpagado"] or 0)

    # ---------- Facturas: saldo, vencimiento y antigüedad ----------
    lista = []
    for x in facturas:
        total, base = float(x["total"] or 0), float(x["base"] or 0)
        pagado, devuelto = pagado_factura.get(x["factura"], 0.0), devoluciones.get(x["factura"], 0.0)
        falta = total - pagado - devuelto
        ficha = fichas.get(clave(x["cliente"]), {})
        ret = (float(ficha.get("ret_dgi") or 0) + float(ficha.get("ret_alma") or 0)) / 100
        retenido = 0.0
        if x["estado"] == "Pagado" or falta <= max(ret, RETENCION_MAX) * base + 1:
            retenido, falta = max(falta, 0.0), 0.0   # pagada (lo que falta es retención o redondeo)
        vence = x["vence"] or x["fecha"]
        dias_vencida = (corte - vence).days if falta > 0 else 0
        lista.append({**x, "total": total, "pagado": pagado, "devuelto": devuelto, "retenido": retenido, "saldo": max(falta, 0.0),
                      "vence": vence, "dias_vencida": dias_vencida, "plazo": (vence - x["fecha"]).days})
    if f["vendedor"]:
        lista = [x for x in lista if x["vendedor"] == f["vendedor"]]
        pagos = [p for p in pagos if p["vendedor"] == f["vendedor"]]

    # ---------- Por mes ----------
    meses = {}
    def mes(d):
        return meses.setdefault(f"{d.year}-{d.month:02d}", {"mes": f"{MESES[d.month - 1]} {str(d.year)[2:]}", "vendido": 0.0, "facturas": 0,
                                                              "cobrado": 0.0, "dias_x_monto": 0.0, "saldo": 0.0, "vencido": 0.0,
                                                              "devuelto": 0.0, "pagadas": 0})
    for x in lista:
        m = mes(x["fecha"])
        m["vendido"] += x["total"]; m["facturas"] += 1; m["saldo"] += x["saldo"]; m["devuelto"] += x["devuelto"]
        m["vencido"] += x["saldo"] if x["dias_vencida"] > 0 else 0
        m["pagadas"] += x["saldo"] == 0
    for p in pagos:
        if p["fechapago"] and f["desde"] <= p["fechapago"] <= f["hasta"]:
            m = mes(p["fechapago"])
            monto = float(p["mpagado"] or 0)
            m["cobrado"] += monto
            m["dias_x_monto"] += monto * float(p["drecuperado"] or 0)
    for m in meses.values():
        m["pct"] = m["cobrado"] / m["vendido"] * 100 if m["vendido"] else None
        m["dias"] = m["dias_x_monto"] / m["cobrado"] if m["cobrado"] else None
    por_mes = [meses[k] for k in sorted(meses)]

    # ---------- Antigüedad de la cartera ----------
    tramos = [{"nombre": n, "monto": 0.0, "facturas": 0} for n, _, _ in TRAMOS]
    for x in lista:
        if x["saldo"] <= 0:
            continue
        for t, (_, desde, hasta) in zip(tramos, TRAMOS):
            if (desde is None and x["dias_vencida"] <= 0) or (desde is not None and desde <= x["dias_vencida"] <= hasta):
                t["monto"] += x["saldo"]; t["facturas"] += 1
                break

    # ---------- Por cliente ----------
    clientes = {}
    def cli(nombre, vendedor=None):
        k = clave(nombre)
        c = clientes.get(k)
        if c is None:
            ficha = fichas.get(k, {})
            c = clientes[k] = {"cliente": nombre, "vendedor": vendedor or ficha.get("vendedor") or "", "terminos": ficha.get("terminos_pago"),
                               "plazo": plazo_de(ficha.get("terminos_pago")), "limite": float(ficha.get("limite_credito") or 0),
                               "activo": ficha.get("estado") == "Activo" if ficha else None, "en_tabla": bool(ficha),
                               "comprado": 0.0, "facturas": 0, "saldo": 0.0, "vencido": 0.0, "max_vencida": 0,
                               "vencido_60": 0.0, "vencido_90": 0.0, "cobrado": 0.0, "dias_x_monto": 0.0, "a_tiempo": 0.0,
                               "plazos": {}, "ultima_compra": None, "ultimo_pago": None, "facturas_pagadas": 0}
        return c
    for x in lista:
        c = cli(x["cliente"], x["vendedor"])
        c["vendedor"] = x["vendedor"] or c["vendedor"]
        c["comprado"] += x["total"]; c["facturas"] += 1; c["saldo"] += x["saldo"]; c["facturas_pagadas"] += x["saldo"] == 0
        c["plazos"][x["plazo"]] = c["plazos"].get(x["plazo"], 0) + 1
        if x["dias_vencida"] > 0:
            c["vencido"] += x["saldo"]
            c["max_vencida"] = max(c["max_vencida"], x["dias_vencida"])
            c["vencido_60"] += x["saldo"] if x["dias_vencida"] > 60 else 0
            c["vencido_90"] += x["saldo"] if x["dias_vencida"] > 90 else 0
        if c["ultima_compra"] is None or x["fecha"] > c["ultima_compra"]:
            c["ultima_compra"] = x["fecha"]
    for p in pagos:
        if not (p["fechapago"] and f["desde"] <= p["fechapago"] <= f["hasta"]):
            continue
        nombre = cliente_de_factura.get(p["nfactura"]) or p["cliente"]
        c = cli(nombre, p["vendedor"])
        monto = float(p["mpagado"] or 0)
        c["cobrado"] += monto
        c["dias_x_monto"] += monto * float(p["drecuperado"] or 0)
        if c["ultimo_pago"] is None or p["fechapago"] > c["ultimo_pago"]:
            c["ultimo_pago"] = p["fechapago"]
        c.setdefault("_pagos", []).append((monto, float(p["drecuperado"] or 0)))

    if not f["vendedor"]:
        for k, ficha in fichas.items():
            if k not in clientes and ficha.get("estado") == "Activo" and float(ficha.get("limite_credito") or 0) > 0:
                cli(ficha["nombre"].strip())
    meses_periodo = max(1, (f["hasta"].year - f["desde"].year) * 12 + f["hasta"].month - f["desde"].month + 1)
    resultado = []
    for c in clientes.values():
        if not c["plazo"]:   # sin términos en la ficha: el plazo más usado en sus facturas
            c["plazo"] = max(c["plazos"].items(), key=lambda x: x[1])[0] if c["plazos"] else 30
            c["plazo_estimado"] = True
        plazo = c["plazo"] or 30
        for monto, dias in c.pop("_pagos", []):
            c["a_tiempo"] += monto if dias <= plazo + TOLERANCIA_DIAS else 0
        c["dias_pago"] = c["dias_x_monto"] / c["cobrado"] if c["cobrado"] else None
        c["pct_a_tiempo"] = c["a_tiempo"] / c["cobrado"] * 100 if c["cobrado"] else None
        c["uso"] = c["saldo"] / c["limite"] * 100 if c["limite"] else None
        c["compra_mensual"] = c["comprado"] / meses_periodo
        recomendar(c, plazo)
        resultado.append(c)

    for palabra in f["q"].upper().split():
        resultado = [c for c in resultado if palabra in clave(c["cliente"]) or palabra in clave(c["vendedor"])]
    conteo_reco = {k: sum(1 for c in resultado if c["reco"] == k) for k in RECOMENDACIONES}
    if f["reco"]:
        resultado = [c for c in resultado if c["reco"] == f["reco"]]
    orden = {"suspender": 0, "revisar": 1, "sin_limite": 2, "aumentar": 3, "mantener": 4, "sin_uso": 5}
    resultado.sort(key=lambda c: (orden[c["reco"]], -c["vencido"], -c["saldo"], -c["comprado"]))

    vendido = sum(x["total"] for x in lista)
    cobrado = sum(m["cobrado"] for m in por_mes)
    saldo = sum(x["saldo"] for x in lista)
    vencido = sum(x["saldo"] for x in lista if x["dias_vencida"] > 0)
    dias = sum(m["dias_x_monto"] for m in por_mes) / cobrado if cobrado else None
    vendedores = sorted({x["vendedor"] for x in facturas if x["vendedor"]})
    return {"por_mes": por_mes, "tramos": tramos, "clientes": resultado, "conteo_reco": conteo_reco, "facturas": lista,
            "vendido": vendido, "cobrado": cobrado, "saldo": saldo, "vencido": vencido, "dias": dias,
            "retenido": sum(x["retenido"] for x in lista), "devuelto": sum(x["devuelto"] for x in lista),
            "vendedores": vendedores, "meses_periodo": meses_periodo}


def redondear(valor, a=5000):
    return max(a, int(round(valor / a)) * a)


def recomendar(c, plazo):
    """Recomendación de crédito con su motivo y un límite sugerido."""
    motivos = []
    dias, a_tiempo = c["dias_pago"], c["pct_a_tiempo"]
    # Vencido de más de 90 días que pesa: más de C$ 1,000 y (10% o más del saldo o más de C$ 50,000);
    # montos chicos en clientes grandes (posibles reclamos) van a "revisar"
    vencido_grave = c["vencido_90"] > 1000 and (c["vencido_90"] >= 0.10 * c["saldo"] or c["vencido_90"] > 50000)
    if vencido_grave or (dias is not None and dias > plazo + 60):
        c["reco"] = "suspender"
        if vencido_grave:
            motivos.append(f"C$ {c['vencido_90']:,.0f} vencidos hace más de 90 días")
        if dias is not None and dias > plazo + 60:
            motivos.append(f"paga en promedio a {dias:.0f} días con plazo de {plazo}")
        c["limite_sugerido"] = 0
    elif (c["vencido_60"] > 1000 or (dias is not None and dias > plazo + 30)
          or (a_tiempo is not None and a_tiempo < 40 and dias is not None and dias > plazo + 15)):
        c["reco"] = "revisar"
        if c["vencido_90"] > 1000:
            motivos.append(f"C$ {c['vencido_90']:,.0f} vencidos hace más de 90 días ({c['vencido_90'] / c['saldo'] * 100:.0f}% del saldo)")
        elif c["vencido_60"] > 1000:
            motivos.append(f"C$ {c['vencido_60']:,.0f} vencidos hace más de 60 días")
        if dias is not None and dias > plazo + 30:
            motivos.append(f"paga a {dias:.0f} días con plazo de {plazo}")
        if a_tiempo is not None and a_tiempo < 40 and dias is not None and dias > plazo + 15:
            motivos.append(f"solo {a_tiempo:.0f}% de lo pagado fue a tiempo (promedio {dias:.0f} días, plazo {plazo})")
        c["limite_sugerido"] = redondear(c["limite"] * 0.5) if c["limite"] else None
    elif c["limite"] <= 0 and c["comprado"] > 0:
        c["reco"] = "sin_limite"
        motivos.append(f"compró C$ {c['comprado']:,.0f} a crédito sin límite asignado")
        c["limite_sugerido"] = redondear(c["compra_mensual"] * max(plazo, 30) / 30 * 1.2) if c["compra_mensual"] else None
    elif c["comprado"] <= 0 and c["saldo"] <= 0:
        c["reco"] = "sin_uso"
        motivos.append("no compró a crédito en el período" + (f" (límite C$ {c['limite']:,.0f})" if c["limite"] else ""))
        c["limite_sugerido"] = None
    elif (a_tiempo is not None and a_tiempo >= 85 and dias is not None and dias <= plazo + TOLERANCIA_DIAS
          and c["vencido"] <= 0.1 * max(c["saldo"], 1) and c["facturas_pagadas"] >= 3
          and ((c["uso"] or 0) >= 70 or c["compra_mensual"] * max(plazo, 30) / 30 >= 0.8 * c["limite"])):
        c["reco"] = "aumentar"
        motivos.append(f"paga a tiempo ({a_tiempo:.0f}% del monto, promedio {dias:.0f} días)")
        motivos.append(f"usa {c['uso']:.0f}% del límite" if c["uso"] is not None else "compra cerca de su límite")
        c["limite_sugerido"] = redondear(max(c["limite"] * 1.25, c["compra_mensual"] * max(plazo, 30) / 30 * 1.3))
    else:
        c["reco"] = "mantener"
        if dias is not None:
            motivos.append(f"paga en promedio a {dias:.0f} días (plazo {plazo})")
        if c["vencido"] > 0:
            motivos.append(f"C$ {c['vencido']:,.0f} vencidos (máx. {c['max_vencida']} días)")
        c["limite_sugerido"] = None
    c["motivo"] = "; ".join(motivos)


@analisis_cartera_bp.route("/analisis_cartera")
@requiere_acceso("analisis_cartera")
def reporte():
    f = leer_filtros(request.args)
    d = calcular(f)
    return render_template("analisis_cartera.html", f=f, d=d, recomendaciones=RECOMENDACIONES, inicio_datos=INICIO_DATOS)


@analisis_cartera_bp.route("/analisis_cartera/cliente")
@requiere_acceso("analisis_cartera")
def detalle_cliente():
    """Facturas de crédito del cliente en el período con su saldo, y sus pagos."""
    f = leer_filtros(request.args)
    nombre = clave(request.args.get("cliente"))
    d = calcular(dict(f, q="", reco=""))
    facturas = [x for x in d["facturas"] if clave(x["cliente"]) == nombre]
    facturas.sort(key=lambda x: (x["saldo"] <= 0, x["fecha"]))
    numeros = {x["factura"] for x in facturas}
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("SELECT nfactura, nrecibo, fechafact, fechapago, drecuperado, mpagado FROM recuperacion "
                       "WHERE fechapago BETWEEN %s AND %s ORDER BY fechapago DESC", (f["desde"], f["hasta"]))
        pagos = [p for p in cursor.fetchall() if p["nfactura"] in numeros]
    finally:
        cursor.close()
        conn.close()
    fmt = lambda v: v.isoformat() if v else ""
    return jsonify({
        "facturas": [{"factura": x["factura"], "fecha": fmt(x["fecha"]), "vence": fmt(x["vence"]), "total": x["total"],
                      "pagado": x["pagado"], "devuelto": x["devuelto"], "retenido": x["retenido"], "saldo": x["saldo"],
                      "dias_vencida": x["dias_vencida"], "vendedor": x["vendedor"]} for x in facturas],
        "pagos": [{"factura": p["nfactura"], "recibo": p["nrecibo"], "fecha": fmt(p["fechapago"]), "dias": int(p["drecuperado"] or 0),
                   "monto": float(p["mpagado"] or 0)} for p in pagos]})


@analisis_cartera_bp.route("/analisis_cartera/excel")
@requiere_acceso("analisis_cartera")
def descargar_excel():
    f = leer_filtros(request.args)
    d = calcular(f)
    wb = Workbook()
    nota = (f"Facturas de crédito del {f['desde']:%d/%m/%Y} al {f['hasta']:%d/%m/%Y} · saldos al {f['corte']:%d/%m/%Y} · "
            f"saldo = total − pagado − devuelto (diferencias por retención IR/alcaldía se consideran pagadas)")

    def hoja(ws, titulo, columnas):
        ws["A1"], ws["A2"] = titulo, nota
        ws["A1"].font = Font(bold=True, size=14, color=VERDE)
        ws["A2"].font = Font(size=9, color="52514E")
        for i, (nombre, ancho) in enumerate(columnas, start=1):
            c = ws.cell(row=4, column=i, value=nombre)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=NARANJA if nombre in ("Recomendación", "Saldo C$") else VERDE)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = ancho
        ws.freeze_panes = "B5"
        ws.auto_filter.ref = f"A4:{get_column_letter(len(columnas))}4"

    def fila(ws, n, valores, formatos):
        for i, v in enumerate(valores, start=1):
            c = ws.cell(row=n, column=i, value=v)
            if i in formatos:
                c.number_format = formatos[i]

    ws = wb.active
    ws.title = "Clientes"
    hoja(ws, "Análisis de cartera por cliente", [("Cliente", 38), ("Vendedor", 24), ("Términos", 10), ("Límite C$", 13), ("Comprado C$", 14),
         ("Facturas", 8), ("Saldo C$", 13), ("Vencido C$", 13), ("Máx. días vencida", 9), ("Vencido +60", 12), ("Vencido +90", 12),
         ("Cobrado C$", 14), ("Días prom. pago", 9), ("% a tiempo", 9), ("Uso del límite", 9), ("Última compra", 12), ("Último pago", 12),
         ("Recomendación", 20), ("Límite sugerido C$", 13), ("Motivo", 60)])
    money = "#,##0.00"
    for n, c in enumerate(d["clientes"], start=5):
        fila(ws, n, [c["cliente"], c["vendedor"], c["terminos"] or f"{c['plazo']} días*", c["limite"], c["comprado"], c["facturas"],
                     c["saldo"], c["vencido"], c["max_vencida"] or None, c["vencido_60"] or None, c["vencido_90"] or None, c["cobrado"],
                     round(c["dias_pago"], 1) if c["dias_pago"] is not None else None,
                     c["pct_a_tiempo"] / 100 if c["pct_a_tiempo"] is not None else None,
                     c["uso"] / 100 if c["uso"] is not None else None, c["ultima_compra"], c["ultimo_pago"],
                     RECOMENDACIONES[c["reco"]][0], c["limite_sugerido"], c["motivo"]],
             {4: money, 5: money, 7: money, 8: money, 10: money, 11: money, 12: money, 14: "0%", 15: "0%", 16: "dd/mm/yyyy",
              17: "dd/mm/yyyy", 19: money})

    ws2 = wb.create_sheet("Por mes")
    hoja(ws2, "Crédito vendido contra recuperado por mes", [("Mes", 10), ("Vendido crédito C$", 15), ("Facturas", 9), ("Cobrado C$", 15),
         ("% cobrado / vendido", 10), ("Días prom. de pago", 10), ("Devuelto C$", 12), ("Saldo pendiente de esas facturas", 15),
         ("Vencido de esas facturas", 14), ("Facturas pagadas", 9)])
    for n, m in enumerate(d["por_mes"], start=5):
        fila(ws2, n, [m["mes"], m["vendido"], m["facturas"], m["cobrado"], m["pct"] / 100 if m["pct"] is not None else None,
                      round(m["dias"], 1) if m["dias"] else None, m["devuelto"], m["saldo"], m["vencido"], m["pagadas"]],
             {2: money, 4: money, 5: "0%", 7: money, 8: money, 9: money})

    ws3 = wb.create_sheet("Facturas pendientes")
    hoja(ws3, "Facturas de crédito con saldo", [("Factura", 10), ("Cliente", 38), ("Vendedor", 24), ("Fecha", 11), ("Vence", 11),
         ("Total C$", 13), ("Pagado C$", 13), ("Devuelto C$", 12), ("Saldo C$", 13), ("Días vencida", 9)])
    pendientes = sorted((x for x in d["facturas"] if x["saldo"] > 0), key=lambda x: -x["dias_vencida"])
    for n, x in enumerate(pendientes, start=5):
        fila(ws3, n, [x["factura"], x["cliente"], x["vendedor"], x["fecha"], x["vence"], x["total"], x["pagado"], x["devuelto"],
                      x["saldo"], x["dias_vencida"] if x["dias_vencida"] > 0 else None],
             {4: "dd/mm/yyyy", 5: "dd/mm/yyyy", 6: money, 7: money, 8: money, 9: money})
    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=f"analisis_cartera_{f['desde']}_al_{f['hasta']}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
