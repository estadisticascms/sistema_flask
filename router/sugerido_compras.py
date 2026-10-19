"""Sugerido de compras: promedio de ventas × meses a cubrir − existencia disponible (sin bodegas NHE)."""
import io
import json
import math
from datetime import date, datetime, timedelta

from flask import jsonify, Blueprint, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from utils.db import get_connection
from router.accesos import requiere_acceso
from router.ventas_perdidas import perdidas_por_codigo, detalle_por_codigo

sugerido_bp = Blueprint("sugerido_compras", __name__)

MESES_COBERTURA = [1, 2, 3, 4, 6, 9, 12]
MESES_HISTORIAL = [3, 6, 12, 24]
DIAS_RECIENTE = [0, 15, 30, 60, 90]
MAX_FILAS_PANTALLA = 2000          # la descarga a Excel no tiene límite
LOTE_IN = 2000                     # tamaño de listas IN (...) por consulta
VERDE, NARANJA = "027266", "FE6412"


# -----------------------------
# Parámetros
# -----------------------------
def entero(valor, permitidos, defecto):
    try:
        v = int(valor)
    except (TypeError, ValueError):
        return defecto
    return v if v in permitidos else defecto


def leer_parametros(fuente):
    marcas = sorted({m.strip() for m in fuente.getlist("marca") if m.strip()})
    return {
        "marcas": marcas,
        "meses_cobertura": entero(fuente.get("meses_cobertura"), MESES_COBERTURA, 3),
        "meses_historial": entero(fuente.get("meses_historial"), MESES_HISTORIAL, 6),
        "dias_reciente": entero(fuente.get("dias_reciente"), DIAS_RECIENTE, 0),
        "q": (fuente.get("q") or "").strip(),
        # por defecto solo productos con sugerido; el checkbox enviado vacío significa "todos"
        "solo_sugerido": fuente.get("solo_sugerido", "1") == "1",
        "ocultar_cubiertos": fuente.get("ocultar_cubiertos") == "1",
        # por defecto las ventas perdidas se suman a la demanda; el checkbox desmarcado manda "0"
        "con_perdidas": fuente.get("con_perdidas", "1") == "1",
        "calculado": fuente.get("calcular") == "1",
    }


def restar_meses(d, meses):
    mes = d.month - meses
    anio = d.year + (mes - 1) // 12
    mes = (mes - 1) % 12 + 1
    dias_mes = [31, 29 if anio % 4 == 0 and (anio % 100 or anio % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    return date(anio, mes, min(d.day, dias_mes[mes - 1]))


def a_fecha(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return datetime.strptime(str(v)[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


# -----------------------------
# Datos
# -----------------------------
def en_lotes(cursor, sql_con_marca, codigos, params_extra=()):
    """Ejecuta una consulta con `IN (...)` por lotes y devuelve todas las filas."""
    filas = []
    for i in range(0, len(codigos), LOTE_IN):
        lote = codigos[i:i + LOTE_IN]
        cursor.execute(sql_con_marca.replace("{IN}", ", ".join(["%s"] * len(lote))), list(params_extra) + lote)
        filas.extend(cursor.fetchall())
    return filas


def calcular(p):
    conn = get_connection()
    cursor = conn.cursor()
    try:
        # Fin del período: último día con ventas cargadas
        cursor.execute("SELECT MAX(Fecha) FROM ventas_historicas")
        fin = a_fecha(cursor.fetchone()[0]) or date.today()
        inicio = restar_meses(fin, p["meses_historial"]) + timedelta(days=1)

        # 1) Productos del catálogo (filtro por marcas y búsqueda)
        where, params = ["p.`Codigo Producto` IS NOT NULL"], []
        if p["marcas"]:
            where.append(f"p.Marca IN ({', '.join(['%s'] * len(p['marcas']))})")
            params += p["marcas"]
        for palabra in p["q"].split():
            like = f"%{palabra}%"
            where.append("(p.`Codigo Producto` LIKE %s OR p.Nombre LIKE %s OR p.`Numero Parte` LIKE %s)")
            params += [like] * 3
        cursor.execute(f"""
            SELECT /*+ NO_INDEX(c) */ p.`Codigo Producto`, p.Nombre, p.`Numero Parte`, p.Marca, p.`Ultimo Costo`, c.segmento,
                   c.grupo_equival, c.familia, c.subfamilia
            FROM catalogo_productos p LEFT JOIN clase_prod c ON c.codigo = p.`Codigo Producto`
            WHERE {' AND '.join(where)}
        """, params)
        productos = {r[0]: {"codigo": r[0], "nombre": r[1], "numero_parte": r[2], "marca": r[3],
                            "costo": float(r[4] or 0), "segmento": r[5], "grupo": (r[6] or "").strip() or None,
                            "familia": r[7], "subfamilia": r[8]}
                     for r in cursor.fetchall()}
        codigos = sorted(productos)
        if not codigos:
            return [], {"inicio": inicio, "fin": fin}

        filtrar = bool(p["marcas"] or p["q"])  # con filtros se consulta solo esos códigos

        # 2) Ventas del período (unidades, sin anuladas). ventas_por_producto no tiene índices:
        #    se recorre una sola vez y se filtra en Python.
        cursor.execute(
            "SELECT Codigo, SUM(`Cantidad Vendida`) FROM ventas_por_producto "
            "WHERE Fecha BETWEEN %s AND %s AND Estatus <> 'Anulado' GROUP BY Codigo",
            (inicio.isoformat(), fin.isoformat()))
        ventas_todas = {cod: float(n or 0) for cod, n in cursor.fetchall()}
        ventas = {cod: n for cod, n in ventas_todas.items() if cod in productos}

        # 3) Existencias por bodega: total y lo que está en bodegas NHE (mercadería faltante)
        sql_exist = ("SELECT Codigo_producto, SUM(Cant_Final), "
                     "SUM(CASE WHEN UPPER(Bodega) LIKE '%NHE%' THEN Cant_Final ELSE 0 END) FROM existencias_bodega {W} "
                     "GROUP BY Codigo_producto")
        if filtrar:
            filas_exist = en_lotes(cursor, sql_exist.replace("{W}", "WHERE Codigo_producto IN ({IN})"), codigos)
        else:
            cursor.execute(sql_exist.replace("{W}", ""))
            filas_exist = cursor.fetchall()
        existencias = {r[0]: (float(r[1] or 0), float(r[2] or 0)) for r in filas_exist}

        # 4) Sugerido por producto. Demanda = lo vendido + (opcional) lo que se perdió por no tenerlo.
        perdidas = perdidas_por_codigo(inicio, fin)
        filas = []
        for cod in codigos:
            total, nhe = existencias.get(cod, (0.0, 0.0))
            disponible = total - nhe
            vendido = ventas.get(cod, 0.0)
            promedio = vendido / p["meses_historial"]
            x = perdidas.get(cod)
            perdido = x["unidades"] if x else 0.0
            def sugerir(demanda):
                return max(0, math.ceil(demanda / p["meses_historial"] * p["meses_cobertura"] - max(disponible, 0) - 1e-9))
            sin_perdidas = sugerir(vendido)
            filas.append({**productos[cod], "existencia": total, "nhe": nhe, "disponible": disponible,
                          "vendido": vendido, "promedio": promedio,
                          "cobertura": (disponible / promedio) if promedio > 0 else None,
                          "sugerido": sugerir(vendido + perdido) if p["con_perdidas"] else sin_perdidas,
                          "sugerido_sin_perdidas": sin_perdidas,
                          "perdidas": perdido, "perdidas_clientes": len(x["clientes"]) if x else 0,
                          "perdidas_ultima": x["ultima"] if x else None,
                          "ingresos": [], "dias_compra": None, "ultima_venta": None})
        if p["solo_sugerido"]:   # también se dejan los que tienen ventas perdidas aunque no tengan sugerido
            filas = [f for f in filas if f["sugerido"] > 0 or f["perdidas"] > 0]

        # 5) Últimos 3 ingresos por compra y última venta, solo de los productos que quedan
        #    (transacciones_inventario tiene índice por producto + fecha)
        resultado = [f["codigo"] for f in filas]
        ingresos = {}
        for cod, fecha, cant in en_lotes(cursor, """
            WITH compras AS (
                SELECT codigo_producto AS cod, fecha, id_documento_origen AS doc, SUM(cantidad) AS cant
                FROM transacciones_inventario
                WHERE codigo_producto IN ({IN}) AND tipo_transaccion = 'ENTRADA' AND tipo_documento_origen = 'ORDEN COMPRA'
                GROUP BY codigo_producto, fecha, id_documento_origen
            ), orden AS (
                SELECT cod, fecha, cant, ROW_NUMBER() OVER (PARTITION BY cod ORDER BY fecha DESC, doc DESC) AS rn FROM compras
            )
            SELECT cod, fecha, cant FROM orden WHERE rn <= 3 ORDER BY cod, rn
        """, resultado):
            ingresos.setdefault(cod, []).append({"fecha": a_fecha(fecha), "cantidad": float(cant or 0)})
        equivalentes = buscar_equivalentes(cursor, filas, existencias, ventas_todas, p["meses_historial"])
        ultima_venta = dict(en_lotes(cursor,
            "SELECT codigo_producto, MAX(fecha) FROM transacciones_inventario "
            "WHERE codigo_producto IN ({IN}) AND tipo_transaccion = 'SALIDA' AND tipo_documento_origen LIKE 'FACTURA%%' "
            "GROUP BY codigo_producto", resultado))
    finally:
        cursor.close()
        conn.close()

    for f in filas:
        f["equivalentes"] = equivalentes.get(f["codigo"], [])
        f["eq_disponible"] = sum(max(e["disponible"], 0) for e in f["equivalentes"])
        f["eq_marcas"] = len({e["marca"] for e in f["equivalentes"] if e["marca"] != f["marca"]})
        f["ingresos"] = ingresos.get(f["codigo"], [])
        ultima_compra = f["ingresos"][0]["fecha"] if f["ingresos"] else None
        f["dias_compra"] = (fin - ultima_compra).days if ultima_compra else None
        f["ultima_venta"] = a_fecha(ultima_venta.get(f["codigo"]))

    # "Tomar en cuenta las últimas compras": opción de ocultar lo comprado recientemente
    if p["dias_reciente"]:
        filas = [f for f in filas if f["dias_compra"] is None or f["dias_compra"] > p["dias_reciente"]]
    if p["ocultar_cubiertos"]:  # lo que pide ya está disponible en productos equivalentes
        filas = [f for f in filas if not (f["sugerido"] > 0 and f["eq_disponible"] >= f["sugerido"])]
    filas.sort(key=lambda f: ((f["marca"] or "").upper(), -f["sugerido"] * f["costo"], f["nombre"] or ""))
    return filas, {"inicio": inicio, "fin": fin}


def buscar_equivalentes(cursor, filas, existencias, ventas_todas, meses_historial):
    """Otros productos del mismo grupo de equivalencia (clase_prod.grupo_equival), con su disponible sin NHE."""
    grupos = sorted({f["grupo"] for f in filas if f["grupo"]})
    if not grupos:
        return {}
    miembros = {}
    for cod, grupo, marca, nombre, parte in en_lotes(cursor, """
        SELECT c.codigo, c.grupo_equival, p.Marca, p.Nombre, p.`Numero Parte`
        FROM clase_prod c JOIN catalogo_productos p ON p.`Codigo Producto` = c.codigo
        WHERE c.grupo_equival IN ({IN})
    """, grupos):
        miembros.setdefault(grupo, []).append({"codigo": cod, "marca": marca, "nombre": nombre, "numero_parte": parte})

    # Existencias de los equivalentes que no estaban en el alcance del reporte
    faltan = sorted({m["codigo"] for ms in miembros.values() for m in ms} - set(existencias))
    for cod, total, nhe in en_lotes(cursor,
            "SELECT Codigo_producto, SUM(Cant_Final), SUM(CASE WHEN UPPER(Bodega) LIKE '%%NHE%%' THEN Cant_Final ELSE 0 END) "
            "FROM existencias_bodega WHERE Codigo_producto IN ({IN}) GROUP BY Codigo_producto", faltan):
        existencias[cod] = (float(total or 0), float(nhe or 0))

    resultado = {}
    for f in filas:
        otros = []
        for m in miembros.get(f["grupo"], []):
            if m["codigo"] == f["codigo"]:
                continue
            total, nhe = existencias.get(m["codigo"], (0.0, 0.0))
            otros.append({**m, "disponible": total - nhe, "promedio": ventas_todas.get(m["codigo"], 0.0) / meses_historial})
        # primero otras marcas y lo que hay disponible
        otros.sort(key=lambda e: (e["marca"] == f["marca"], -e["disponible"], e["codigo"]))
        if otros:
            resultado[f["codigo"]] = otros
    return resultado


def lista_marcas():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT Marca, COUNT(*) FROM catalogo_productos WHERE Marca IS NOT NULL AND Marca <> '' "
                   "GROUP BY Marca ORDER BY Marca")
    marcas = cursor.fetchall()
    cursor.close()
    conn.close()
    return marcas


# -----------------------------
# Vistas
# -----------------------------
@sugerido_bp.route("/sugerido_compras", methods=["GET"])
@requiere_acceso("sugerido_compras")
def sugerido_compras():
    p = leer_parametros(request.args)
    filas, periodo = ([], None)
    if p["calculado"]:
        filas, periodo = calcular(p)
    return render_template(
        "sugerido_compras.html", p=p, filas=filas[:MAX_FILAS_PANTALLA], total_filas=len(filas),
        recortado=len(filas) > MAX_FILAS_PANTALLA, periodo=periodo, marcas=lista_marcas(),
        opciones_cobertura=MESES_COBERTURA, opciones_historial=MESES_HISTORIAL, opciones_reciente=DIAS_RECIENTE,
        total_nhe=sum(f["nhe"] for f in filas), con_sugerido=sum(1 for f in filas if f["sugerido"] > 0),
        con_equivalente=sum(1 for f in filas if f["sugerido"] > 0 and f["eq_disponible"] > 0),
        con_perdidas=sum(1 for f in filas if f["perdidas"] > 0),
        por_perdidas=sum(f["sugerido"] - f["sugerido_sin_perdidas"] for f in filas),
        perdidas_sin_sugerido=sum(1 for f in filas if f["perdidas"] > 0 and f["sugerido"] == 0),
    )


@sugerido_bp.route("/sugerido_compras/perdidas/<codigo>")
@requiere_acceso("sugerido_compras")
def detalle_perdidas(codigo):
    """Detalle de las ventas perdidas de un producto (se carga al abrir la ventana)."""
    desde, hasta = a_fecha(request.args.get("desde")), a_fecha(request.args.get("hasta"))
    if not desde or not hasta:
        return jsonify([]), 400
    return jsonify(detalle_por_codigo(codigo.strip(), desde, hasta))


@sugerido_bp.route("/sugerido_compras/excel", methods=["POST"])
@requiere_acceso("sugerido_compras")
def descargar_excel():
    p = leer_parametros(request.form)
    try:
        cambios = {str(k): max(0, int(v)) for k, v in json.loads(request.form.get("cambios") or "{}").items()}
    except (ValueError, TypeError, AttributeError):
        cambios = {}
    solo_pedido = request.form.get("solo_pedido") == "1"
    filas, periodo = calcular(p)
    for f in filas:
        f["pedido"] = cambios.get(f["codigo"], f["sugerido"])
        f["editado"] = f["codigo"] in cambios and cambios[f["codigo"]] != f["sugerido"]
    if solo_pedido:
        filas = [f for f in filas if f["pedido"] > 0]

    contenido = generar_excel(filas, p, periodo)
    nombre = f"sugerido_compras_{(p['marcas'][0] if len(p['marcas']) == 1 else 'marcas').replace(' ', '_')}_{date.today():%Y-%m-%d}.xlsx"
    return send_file(io.BytesIO(contenido), as_attachment=True, download_name=nombre,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def generar_excel(filas, p, periodo):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sugerido de compras"
    fino = Side(style="thin", color="E6E4DF")
    borde = Border(bottom=fino)

    ws["A1"] = "Sugerido de compras"
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    marcas = ", ".join(p["marcas"]) if p["marcas"] else "Todas"
    ws["A2"] = (f"Marcas: {marcas}  ·  Cobertura: {p['meses_cobertura']} mes(es)  ·  "
                f"Promedio de ventas: {periodo['inicio']:%d/%m/%Y} al {periodo['fin']:%d/%m/%Y} ({p['meses_historial']} meses)")
    demanda = "(vendido + ventas perdidas)" if p["con_perdidas"] else "vendido"
    ws["A3"] = ("Disponible = existencia total − bodegas NHE (mercadería faltante).  "
                f"Sugerido = {demanda} ÷ meses × meses de cobertura − disponible.  Generado {datetime.now():%d/%m/%Y %H:%M}")
    ws["A2"].font = ws["A3"].font = Font(size=9, color="52514E")

    columnas = [
        ("Código", 11), ("Nº parte", 18), ("Descripción", 38), ("Marca", 14), ("Segmento", 14), ("Familia", 20),
        ("Subfamilia", 22), ("Existencia total", 11), ("En NHE", 9), ("Disponible", 11), ("Vendido período", 11),
        ("Ventas perdidas (unid.)", 11), ("Prom. mensual", 11), ("Cobertura (meses)", 11), ("Ingreso 1 fecha", 12),
        ("Ingreso 1 cant.", 10), ("Ingreso 2 fecha", 12), ("Ingreso 2 cant.", 10), ("Ingreso 3 fecha", 12),
        ("Ingreso 3 cant.", 10), ("Última venta", 12), ("Sugerido sin ventas perdidas", 12), ("Sugerido sistema", 11),
        ("Pedido final", 11), ("Último costo C$", 12), ("Costo pedido C$", 14), ("Grupo equiv.", 11), ("Equivalentes", 11),
        ("Disponible en equivalentes", 13), ("Detalle equivalentes", 60), ("Clientes que lo pidieron", 11),
        ("Última venta perdida", 12),
    ]
    col = {titulo: i for i, (titulo, _) in enumerate(columnas, start=1)}
    letra = {titulo: get_column_letter(i) for titulo, i in col.items()}
    fila_tit = 5
    for i, (titulo, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=fila_tit, column=i, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NARANJA if titulo in ("Pedido final", "Ventas perdidas (unid.)") else VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[fila_tit].height = 42

    resaltado = PatternFill("solid", fgColor="FFF1E8")
    formatos = {"dd/mm/yyyy": ("Ingreso 1 fecha", "Ingreso 2 fecha", "Ingreso 3 fecha", "Última venta", "Última venta perdida"),
                "#,##0.##": ("Existencia total", "En NHE", "Disponible", "Vendido período", "Ventas perdidas (unid.)",
                             "Ingreso 1 cant.", "Ingreso 2 cant.", "Ingreso 3 cant.", "Sugerido sin ventas perdidas",
                             "Sugerido sistema", "Pedido final", "Equivalentes", "Disponible en equivalentes"),
                "#,##0.0": ("Prom. mensual", "Cobertura (meses)"),
                "#,##0.00": ("Último costo C$", "Costo pedido C$")}
    for n, f in enumerate(filas, start=fila_tit + 1):
        ing = f["ingresos"] + [{"fecha": None, "cantidad": None}] * (3 - len(f["ingresos"]))
        valores = {
            "Código": f["codigo"], "Nº parte": f["numero_parte"], "Descripción": f["nombre"], "Marca": f["marca"],
            "Segmento": f["segmento"], "Familia": f["familia"], "Subfamilia": f["subfamilia"],
            "Existencia total": f["existencia"], "En NHE": f["nhe"], "Disponible": f["disponible"],
            "Vendido período": f["vendido"], "Ventas perdidas (unid.)": f["perdidas"] or None,
            "Prom. mensual": round(f["promedio"], 2),
            "Cobertura (meses)": round(f["cobertura"], 1) if f["cobertura"] is not None else None,
            "Ingreso 1 fecha": ing[0]["fecha"], "Ingreso 1 cant.": ing[0]["cantidad"],
            "Ingreso 2 fecha": ing[1]["fecha"], "Ingreso 2 cant.": ing[1]["cantidad"],
            "Ingreso 3 fecha": ing[2]["fecha"], "Ingreso 3 cant.": ing[2]["cantidad"],
            "Última venta": f["ultima_venta"], "Sugerido sin ventas perdidas": f["sugerido_sin_perdidas"],
            "Sugerido sistema": f["sugerido"], "Pedido final": f["pedido"], "Último costo C$": f["costo"],
            "Costo pedido C$": f"={letra['Pedido final']}{n}*{letra['Último costo C$']}{n}",
            "Grupo equiv.": f["grupo"], "Equivalentes": len(f["equivalentes"]) or None,
            "Disponible en equivalentes": f["eq_disponible"] if f["equivalentes"] else None,
            "Detalle equivalentes": "; ".join(f"{e['marca']} {e['codigo']} ({e['numero_parte'] or 's/n'}): disp. {e['disponible']:,.0f}"
                                              for e in f["equivalentes"]) or None,
            "Clientes que lo pidieron": f["perdidas_clientes"] or None, "Última venta perdida": f["perdidas_ultima"],
        }
        for titulo, v in valores.items():
            c = ws.cell(row=n, column=col[titulo], value=v)
            c.border = borde
        for fmt, titulos in formatos.items():
            for t in titulos:
                ws.cell(row=n, column=col[t]).number_format = fmt
        if f["perdidas"]:
            ws.cell(row=n, column=col["Ventas perdidas (unid.)"]).fill = resaltado
        if f["eq_disponible"] > 0:
            ws.cell(row=n, column=col["Disponible en equivalentes"]).fill = PatternFill("solid", fgColor="E6F2F0")
        ws.cell(row=n, column=col["Pedido final"]).font = Font(bold=True)
        if f.get("editado"):
            ws.cell(row=n, column=col["Pedido final"]).fill = resaltado

    ultima = fila_tit + len(filas)
    tot = ultima + 1
    ws.cell(row=tot, column=1, value="TOTAL").font = Font(bold=True)
    if filas:
        for titulo in ("Ventas perdidas (unid.)", "Sugerido sin ventas perdidas", "Sugerido sistema", "Pedido final", "Costo pedido C$"):
            l = letra[titulo]
            c = ws.cell(row=tot, column=col[titulo], value=f"=SUM({l}{fila_tit + 1}:{l}{ultima})")
            c.font = Font(bold=True)
            c.number_format = "#,##0.00" if titulo == "Costo pedido C$" else "#,##0"
    ws.freeze_panes = ws.cell(row=fila_tit + 1, column=4)
    ws.auto_filter.ref = f"A{fila_tit}:{get_column_letter(len(columnas))}{max(ultima, fila_tit)}"
    ws.sheet_view.zoomScale = 90

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
