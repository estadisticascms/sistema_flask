"""Bitácora de accesos (tabla bitacora_accesos, docs/sql/bitacora_accesos.sql).

Se registra sola: iniciar_bitacora(app) engancha un after_request que anota cada pantalla abierta, cada formulario enviado
(guardar, cargar, eliminar) y cada descarga de un usuario con sesión. El login, los intentos fallidos, la salida y los
accesos denegados los anotan login.py y requiere_acceso. El nombre del módulo sale del menú (router/menu.py).
El reporte está en /bitacora_accesos."""
import io
from datetime import date, datetime, timedelta

from flask import Blueprint, current_app, render_template, request, send_file, session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from router.accesos import catalogo_vistas, requiere_acceso
from utils.db import get_connection

bitacora_bp = Blueprint("bitacora", __name__)

EVENTOS = {
    "login": "Inició sesión",
    "vista": "Abrió",
    "accion": "Guardó / cargó",
    "descarga": "Descargó",
    "salida": "Cerró sesión",
    "denegado": "Sin permiso",
    "fallido": "Clave incorrecta",
}
NO_REGISTRAR = ("/static/", "/favicon")
LIMITE_DETALLE = 1000


def _modulos():
    """(url -> (módulo, área), vista -> (módulo, área)), armado una vez desde el menú."""
    cache = current_app.config.get("_BITACORA_MODULOS")
    if cache is None:
        from router.menu import SECCIONES
        por_url = {}
        for s in sorted(SECCIONES, key=lambda x: bool(x.get("acceso_rapido"))):   # el área real antes que el atajo de Gerencia
            for g in s["grupos"]:
                for nombre, url, *_ in g["items"]:
                    por_url.setdefault(url, (nombre, s["titulo"]))
        por_vista = {i["vista"]: (i["nombre"], g["titulo"]) for g in catalogo_vistas() for i in g["items"]}
        cache = current_app.config["_BITACORA_MODULOS"] = (por_url, por_vista)
    return cache


def modulo_actual():
    if request.path == "/":
        return "Inicio", ""
    if request.path in ("/login", "/logout"):
        return "Ingreso al sistema", ""
    por_url, por_vista = _modulos()
    if request.path in por_url:
        return por_url[request.path]
    vista = getattr(current_app.view_functions.get(request.endpoint), "vista_acceso", None)
    if vista in por_vista:
        return por_vista[vista]
    return (request.endpoint or request.path).split(".")[-1].replace("_", " ").capitalize(), ""


def registrar(evento, usuario=None):
    """Anota un evento. Nunca interrumpe la pantalla: si la base falla, se ignora."""
    usuario = usuario or session.get("usuario")
    if not usuario:
        return
    try:
        modulo, area = modulo_actual()
        ip = (request.headers.get("X-Forwarded-For") or request.remote_addr or "").split(",")[0].strip()
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("INSERT INTO bitacora_accesos (fecha_hora, usuario, evento, modulo, area, ruta, metodo, ip, navegador) "
                           "VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)",
                           (datetime.now().replace(microsecond=0), usuario[:100], evento, modulo[:150], area[:60],
                            request.full_path.rstrip("?")[:255], request.method, ip[:45],
                            (request.headers.get("User-Agent") or "")[:255]))
            conn.commit()
        finally:
            cursor.close()
            conn.close()
    except Exception:
        current_app.logger.exception("No se pudo registrar en la bitácora")


def _despues(respuesta):
    if "usuario" not in session or request.path.startswith(NO_REGISTRAR) or request.path in ("/login", "/logout"):
        return respuesta
    if request.method == "POST":
        registrar("accion")
    elif request.method == "GET" and respuesta.status_code == 200:
        if "attachment" in (respuesta.headers.get("Content-Disposition") or ""):
            registrar("descarga")
        elif respuesta.mimetype == "text/html":   # las consultas internas (JSON) de una pantalla no cuentan
            registrar("vista")
    return respuesta


def iniciar_bitacora(app):
    app.after_request(_despues)


# ---------------------------------------------------------------- Reporte

def leer_filtros(args):
    hoy = date.today()
    try:
        desde = date.fromisoformat(args.get("desde", ""))
    except ValueError:
        desde = hoy - timedelta(days=6)
    try:
        hasta = date.fromisoformat(args.get("hasta", ""))
    except ValueError:
        hasta = hoy
    if hasta < desde:
        desde, hasta = hasta, desde
    return {"desde": desde, "hasta": hasta, "usuario": args.get("usuario", ""), "modulo": args.get("modulo", ""),
            "evento": args.get("evento", "") if args.get("evento") in EVENTOS else ""}


def _where(f):
    sql = ["fecha_hora >= %s AND fecha_hora < %s"]
    valores = [f["desde"], f["hasta"] + timedelta(days=1)]
    for campo in ("usuario", "modulo", "evento"):
        if f[campo]:
            sql.append(f"{campo} = %s")
            valores.append(f[campo])
    return " AND ".join(sql), valores


def calcular(f):
    where, valores = _where(f)
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute(f"SELECT evento, COUNT(*) n, COUNT(DISTINCT usuario) usuarios FROM bitacora_accesos WHERE {where} GROUP BY evento", valores)
        por_evento = {r["evento"]: r for r in cursor.fetchall()}
        cursor.execute(f"SELECT COUNT(DISTINCT usuario) usuarios, COUNT(DISTINCT modulo) modulos FROM bitacora_accesos "
                       f"WHERE {where} AND evento IN ('login','vista','accion','descarga')", valores)
        totales = cursor.fetchone()

        # Un renglón por usuario y día: a qué hora entró, hasta qué hora trabajó y qué módulos usó
        cursor.execute(f"""
            SELECT usuario, DATE(fecha_hora) dia, MIN(fecha_hora) entrada, MAX(fecha_hora) ultima,
                   SUM(evento = 'login') logins, SUM(evento = 'vista') vistas, SUM(evento = 'accion') acciones,
                   SUM(evento = 'descarga') descargas, SUM(evento IN ('fallido','denegado')) alertas,
                   GROUP_CONCAT(DISTINCT CASE WHEN evento IN ('vista','accion','descarga') AND modulo NOT IN ('Inicio','Ingreso al sistema')
                                THEN modulo END ORDER BY modulo SEPARATOR ' · ') modulos
            FROM bitacora_accesos WHERE {where}
            GROUP BY usuario, DATE(fecha_hora) ORDER BY dia DESC, entrada""", valores)
        jornadas = cursor.fetchall()
        for j in jornadas:
            j["minutos"] = int((j["ultima"] - j["entrada"]).total_seconds() // 60)

        cursor.execute(f"""
            SELECT modulo, area, COUNT(*) usos, COUNT(DISTINCT usuario) usuarios, MAX(fecha_hora) ultimo,
                   GROUP_CONCAT(DISTINCT usuario ORDER BY usuario SEPARATOR ', ') quienes
            FROM bitacora_accesos WHERE {where} AND evento IN ('vista','accion','descarga')
              AND modulo NOT IN ('Inicio','Ingreso al sistema')
            GROUP BY modulo, area ORDER BY usos DESC""", valores)
        modulos = cursor.fetchall()

        cursor.execute(f"SELECT HOUR(fecha_hora) hora, COUNT(*) n FROM bitacora_accesos WHERE {where} "
                       f"AND evento IN ('vista','accion','descarga') GROUP BY HOUR(fecha_hora)", valores)
        por_hora = {r["hora"]: r["n"] for r in cursor.fetchall()}

        cursor.execute(f"SELECT * FROM bitacora_accesos WHERE {where} ORDER BY fecha_hora DESC, id DESC LIMIT %s",
                       valores + [LIMITE_DETALLE])
        detalle = cursor.fetchall()

        # Usuarios que no entraron en el período (y cuándo fue la última vez)
        cursor.execute("""
            SELECT u.usuario, u.rol, MAX(b.fecha_hora) ultima
            FROM usuarios u LEFT JOIN bitacora_accesos b ON b.usuario = u.usuario AND b.evento <> 'fallido'
            GROUP BY u.usuario, u.rol ORDER BY u.usuario""")
        todos = cursor.fetchall()

        cursor.execute("SELECT DISTINCT usuario FROM bitacora_accesos ORDER BY usuario")
        lista_usuarios = sorted({r["usuario"] for r in cursor.fetchall()} | {u["usuario"] for u in todos}, key=str.lower)
        cursor.execute("SELECT DISTINCT modulo FROM bitacora_accesos WHERE modulo IS NOT NULL ORDER BY modulo")
        lista_modulos = [r["modulo"] for r in cursor.fetchall()]
    finally:
        cursor.close()
        conn.close()

    activos = {j["usuario"].lower() for j in jornadas}
    sin_entrar = [u for u in todos if u["usuario"].lower() not in activos] if not (f["modulo"] or f["evento"]) else []
    return {
        "eventos": {k: por_evento.get(k, {}).get("n", 0) for k in EVENTOS},
        "usuarios_activos": totales["usuarios"] or 0, "modulos_usados": totales["modulos"] or 0,
        "jornadas": jornadas, "modulos": modulos, "detalle": detalle, "sin_entrar": sin_entrar,
        "horas": [por_hora.get(h, 0) for h in range(24)],
        "lista_usuarios": lista_usuarios, "lista_modulos": lista_modulos,
    }


@bitacora_bp.route("/bitacora_accesos")
@requiere_acceso("bitacora_accesos")
def reporte():
    f = leer_filtros(request.args)
    d = calcular(f)
    hoy = date.today()
    rangos = [("Hoy", hoy, hoy), ("Ayer", hoy - timedelta(days=1), hoy - timedelta(days=1)),
              ("Últimos 7 días", hoy - timedelta(days=6), hoy), ("Mes en curso", hoy.replace(day=1), hoy)]
    params = {k: (v.isoformat() if isinstance(v, date) else v) for k, v in f.items() if v}
    return render_template("bitacora_accesos.html", f=f, d=d, eventos=EVENTOS, rangos=rangos, params=params, limite=LIMITE_DETALLE)


@bitacora_bp.route("/bitacora_accesos/excel")
@requiere_acceso("bitacora_accesos")
def descargar_excel():
    f = leer_filtros(request.args)
    where, valores = _where(f)
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute(f"SELECT * FROM bitacora_accesos WHERE {where} ORDER BY fecha_hora, id", valores)
        filas = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Bitácora"
    ws["A1"] = "Bitácora de accesos al sistema"
    ws["A1"].font = Font(bold=True, size=14, color="027266")
    ws["A2"] = f"Del {f['desde']:%d/%m/%Y} al {f['hasta']:%d/%m/%Y}" + "".join(f" · {k}: {f[k]}" for k in ("usuario", "modulo") if f[k])
    ws["A2"].font = Font(size=9, color="52514E")
    encabezados = ["Fecha", "Hora", "Usuario", "Evento", "Módulo", "Área", "Ruta", "IP", "Navegador"]
    for i, t in enumerate(encabezados, start=1):
        c = ws.cell(row=4, column=i, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="027266")
    rojo = Font(bold=True, color="B42318")
    for n, r in enumerate(filas, start=5):
        valores_fila = [r["fecha_hora"].date(), r["fecha_hora"].strftime("%H:%M:%S"), r["usuario"], EVENTOS.get(r["evento"], r["evento"]),
                        r["modulo"], r["area"], r["ruta"], r["ip"], r["navegador"]]
        for i, v in enumerate(valores_fila, start=1):
            ws.cell(row=n, column=i, value=v)
        ws.cell(row=n, column=1).number_format = "dd/mm/yyyy"
        if r["evento"] in ("fallido", "denegado"):
            ws.cell(row=n, column=4).font = rojo
    for i, ancho in enumerate([11, 9, 16, 16, 38, 22, 40, 15, 50], start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{max(4, len(filas) + 4)}"

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=f"bitacora_accesos_{f['desde']:%Y%m%d}_{f['hasta']:%Y%m%d}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
