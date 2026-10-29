"""Listado de existencias para clientes: el vendedor lo genera por marca, segmento, familia y subfamilia
y lo entrega en Excel o PDF, con la aplicación de cada producto (Descripción Larga del catálogo). La existencia es la disponible para vender: no cuenta las bodegas NHE
(mercadería faltante), Consumo ni Ejército."""
import io
from datetime import datetime

from flask import Blueprint, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.db import get_connection
from router.accesos import requiere_acceso
from router.productos_incompletos import tiene_valor
from router.sugerido_compras import lista_marcas

listado_existencias_bp = Blueprint("listado_existencias", __name__)

# Bodegas que no se ofrecen al cliente (se buscan por nombre: "NHE B1", "CONSUMO", "EJERCITO DE NIC.", ...)
BODEGAS_EXCLUIDAS = ["NHE", "CONSUMO", "EJERCITO"]
PRECIOS = {1: "Precio 1", 2: "Precio 2", 3: "Precio 3", 4: "Precio 4"}
AGRUPAR = {"familia": "Familia / subfamilia", "marca": "Marca", "ninguno": "Sin agrupar"}
MAX_FILAS_PANTALLA = 2000        # el Excel y el PDF llevan todo
VERDE, NARANJA = "027266", "FE6412"
SIN_CLASIFICAR = "Sin clasificar"

COND_EXCLUIDA = " OR ".join(f"UPPER(Bodega) LIKE '%{b}%'" for b in BODEGAS_EXCLUIDAS)


def leer_filtros(fuente):
    try:
        precio = int(fuente.get("precio") or 0)
    except ValueError:
        precio = 0
    return {
        "marcas": sorted({m.strip() for m in fuente.getlist("marca") if m.strip()}),
        "segmentos": sorted({s.strip() for s in fuente.getlist("segmento") if s.strip()}),
        "familia": (fuente.get("familia") or "").strip(),
        "subfamilia": (fuente.get("subfamilia") or "").strip(),
        "q": (fuente.get("q") or "").strip(),
        "precio": precio if precio in PRECIOS else 0,
        "agrupar": fuente.get("agrupar") if fuente.get("agrupar") in AGRUPAR else "familia",
        # por defecto solo lo que hay para vender; la casilla desmarcada manda "0"
        "solo_con_existencia": fuente.get("solo_con_existencia", "1") == "1",
    }


def hay_filtro(f):
    return bool(f["marcas"] or f["segmentos"] or f["familia"] or f["subfamilia"] or f["q"])


def params_url(f):
    """Filtros para armar los enlaces de Excel / PDF (marca va repetida)."""
    p = [("marca", m) for m in f["marcas"]] + [("segmento", s) for s in f["segmentos"]]
    p += [(k, f[k]) for k in ("familia", "subfamilia", "q", "agrupar") if f[k]]
    p += [("precio", f["precio"]), ("solo_con_existencia", "1" if f["solo_con_existencia"] else "0")]
    return p


def descripcion_filtros(f):
    partes = []
    if f["marcas"]:
        partes.append(("Marca" if len(f["marcas"]) == 1 else "Marcas") + ": " + ", ".join(f["marcas"]))
    if f["segmentos"]:
        partes.append(("Segmento" if len(f["segmentos"]) == 1 else "Segmentos") + ": " + ", ".join(f["segmentos"]))
    for k, t in (("familia", "Familia"), ("subfamilia", "Subfamilia")):
        if f[k]:
            partes.append(f"{t}: {f[k]}")
    if f["q"]:
        partes.append(f"Búsqueda: {f['q']}")
    return " · ".join(partes) or "Todos los productos"


def opciones():
    """Segmentos y árbol familia -> subfamilias de clase_prod, para los filtros."""
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT DISTINCT segmento FROM clase_prod WHERE segmento IS NOT NULL AND segmento <> '' ORDER BY segmento")
        segmentos = [r[0] for r in cursor.fetchall()]
        cursor.execute("SELECT DISTINCT familia, subfamilia FROM clase_prod WHERE familia IS NOT NULL AND familia <> '' "
                       "ORDER BY familia, subfamilia")
        arbol = {}
        for familia, sub in cursor.fetchall():
            subs = arbol.setdefault(familia, [])
            if sub:
                subs.append(sub)
        cursor.execute(f"SELECT DISTINCT Bodega FROM existencias_bodega WHERE {COND_EXCLUIDA} ORDER BY Bodega")
        excluidas = [r[0] for r in cursor.fetchall()]
        return segmentos, arbol, excluidas
    finally:
        cursor.close()
        conn.close()


def calcular(f):
    where, params = ["p.`Codigo Producto` IS NOT NULL"], []
    if f["marcas"]:
        where.append(f"p.Marca IN ({', '.join(['%s'] * len(f['marcas']))})")
        params += f["marcas"]
    if f["segmentos"]:
        where.append(f"c.segmento IN ({', '.join(['%s'] * len(f['segmentos']))})")
        params += f["segmentos"]
    for campo in ("familia", "subfamilia"):
        if f[campo]:
            where.append(f"c.{campo} = %s")
            params.append(f[campo])
    if f["q"]:
        where.append("(p.Nombre LIKE %s OR p.`Codigo Producto` LIKE %s OR p.`Numero Parte` LIKE %s OR p.`Descripcion Larga` LIKE %s)")
        params += [f"%{f['q']}%"] * 4
    precio = f"p.`Precio {f['precio']}`" if f["precio"] else "NULL"

    conn = get_connection()
    cursor = conn.cursor()
    try:
        # Existencia por producto en una sola pasada: total y lo que está en bodegas excluidas
        cursor.execute(f"""
            SELECT p.`Codigo Producto`, p.Nombre, p.`Numero Parte`, p.Medida, p.Marca, p.`Unidad Medida`, p.`Descripcion Larga`,
                   c.segmento, c.familia, c.subfamilia, {precio}, p.Moneda,
                   COALESCE(e.total, 0), COALESCE(e.excluida, 0)
            FROM catalogo_productos p
            LEFT JOIN clase_prod c ON c.codigo = p.`Codigo Producto`
            LEFT JOIN (SELECT Codigo_producto, SUM(Cant_Final) AS total,
                              SUM(CASE WHEN {COND_EXCLUIDA} THEN Cant_Final ELSE 0 END) AS excluida
                       FROM existencias_bodega GROUP BY Codigo_producto) e ON e.Codigo_producto = p.`Codigo Producto`
            WHERE {' AND '.join(where)}
        """.replace("%'", "%%'"), params)
        filas = []
        for (cod, nombre, parte, medida, marca, uom, aplicacion, segmento, familia, subfamilia,
             valor_precio, moneda, total, excluida) in cursor.fetchall():
            disponible = float(total) - float(excluida)
            if f["solo_con_existencia"] and disponible <= 0:
                continue
            filas.append({"codigo": cod, "nombre": (nombre or "").strip(), "numero_parte": (parte or "").strip(),
                          "medida": (medida or "").strip(), "aplicacion": aplicacion.strip() if tiene_valor(aplicacion) else "",
                          "marca": (marca or "").strip() or "Sin marca", "uom": uom or "",
                          "segmento": segmento or SIN_CLASIFICAR, "familia": familia or SIN_CLASIFICAR,
                          "subfamilia": subfamilia or "", "precio": float(valor_precio) if valor_precio is not None else None,
                          "moneda": moneda or "", "total": float(total), "excluida": float(excluida),
                          "existencia": max(disponible, 0.0)})
    finally:
        cursor.close()
        conn.close()

    if f["agrupar"] == "marca":
        filas.sort(key=lambda x: (x["marca"], x["nombre"], x["numero_parte"]))
    elif f["agrupar"] == "familia":
        # lo sin clasificar va al final
        filas.sort(key=lambda x: (x["familia"] == SIN_CLASIFICAR, x["familia"], x["subfamilia"], x["nombre"], x["numero_parte"]))
    else:
        filas.sort(key=lambda x: (x["nombre"], x["numero_parte"]))
    for x in filas:
        x["grupo"] = (x["marca"] if f["agrupar"] == "marca"
                      else (x["familia"] + (f" / {x['subfamilia']}" if x["subfamilia"] else "")) if f["agrupar"] == "familia" else "")
    return filas


def resumen(filas):
    return {"productos": len(filas), "unidades": sum(x["existencia"] for x in filas),
            "excluida": sum(max(x["excluida"], 0) for x in filas),
            "marcas": len({x["marca"] for x in filas}), "grupos": len({x["grupo"] for x in filas})}


def titulo_listado(f):
    if len(f["marcas"]) == 1 and not (f["segmentos"] or f["familia"] or f["subfamilia"]):
        return f"Existencias {f['marcas'][0]}"
    detalle = f["subfamilia"] or f["familia"] or (" y ".join(f["segmentos"]) if 0 < len(f["segmentos"]) <= 2 else "")
    return f"Existencias {detalle}" if detalle else "Listado de existencias"


def nombre_archivo(f, ext):
    base = titulo_listado(f).replace("/", "-").replace(" ", "_")
    return f"{base}_{datetime.now():%Y-%m-%d}.{ext}"


# -----------------------------
# Vistas
# -----------------------------
@listado_existencias_bp.route("/listado_existencias")
@requiere_acceso("listado_existencias")
def listado():
    f = leer_filtros(request.args)
    segmentos, arbol, excluidas = opciones()
    filas = calcular(f) if hay_filtro(f) else []
    return render_template("listado_existencias.html", f=f, filas=filas[:MAX_FILAS_PANTALLA], total_filas=len(filas),
                           r=resumen(filas), buscado=hay_filtro(f), segmentos=segmentos, arbol=arbol, excluidas=excluidas,
                           marcas=lista_marcas(), precios=PRECIOS, agrupar=AGRUPAR, params=params_url(f),
                           titulo=titulo_listado(f), filtros_txt=descripcion_filtros(f))


@listado_existencias_bp.route("/listado_existencias/datos")
@requiere_acceso("listado_existencias")
def datos():
    """Todo el listado en JSON para armar el PDF en el navegador."""
    f = leer_filtros(request.args)
    filas = calcular(f) if hay_filtro(f) else []
    return jsonify({"titulo": titulo_listado(f), "filtros": descripcion_filtros(f), "archivo": nombre_archivo(f, "pdf"),
                    "precio": PRECIOS.get(f["precio"]), "agrupar": f["agrupar"], "generado": f"{datetime.now():%d/%m/%Y %H:%M}",
                    "filas": [{k: x[k] for k in ("codigo", "numero_parte", "nombre", "aplicacion", "marca", "uom", "grupo",
                                                  "existencia", "precio", "moneda")} for x in filas]})


@listado_existencias_bp.route("/listado_existencias/excel")
@requiere_acceso("listado_existencias")
def excel():
    f = leer_filtros(request.args)
    filas = calcular(f) if hay_filtro(f) else []
    wb = Workbook()
    ws = wb.active
    ws.title = "Existencias"
    ws["A1"] = f"Grupo Comersa · {titulo_listado(f)}"
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    ws["A2"] = f"{descripcion_filtros(f)} · Existencias al {datetime.now():%d/%m/%Y %H:%M}"
    ws["A2"].font = Font(size=9, color="52514E")

    columnas = [("Código", 12), ("Nº parte", 20), ("Descripción", 40), ("Aplicación", 50), ("Marca", 16),
                ("Segmento", 16), ("Familia", 20), ("Subfamilia", 22), ("Unidad", 8), ("Existencia", 11)]
    if f["precio"]:
        columnas.append((f"{PRECIOS[f['precio']]} U$", 12))
    ft = 4
    for i, (nombre, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=ft, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NARANJA if nombre == "Existencia" else VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[ft].height = 22
    for n, x in enumerate(filas, start=ft + 1):
        valores = [x["codigo"], x["numero_parte"] or None, x["nombre"], x["aplicacion"] or None, x["marca"], x["segmento"],
                   x["familia"], x["subfamilia"] or None, x["uom"] or None, x["existencia"]]
        if f["precio"]:
            valores.append(round(x["precio"], 2) if x["precio"] is not None else None)
        for i, v in enumerate(valores, start=1):
            ws.cell(row=n, column=i, value=v)
        ws.cell(row=n, column=4).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=n, column=10).number_format = "#,##0.##"
        if f["precio"]:
            ws.cell(row=n, column=11).number_format = "#,##0.00"
    if filas:
        tot = ft + len(filas) + 1
        ws.cell(row=tot, column=1, value=f"{len(filas):,} productos").font = Font(bold=True)
        c = ws.cell(row=tot, column=10, value=f"=SUM(J{ft + 1}:J{tot - 1})")
        c.font = Font(bold=True)
        c.number_format = "#,##0.##"
    ws.freeze_panes = ws.cell(row=ft + 1, column=4)
    ws.auto_filter.ref = f"A{ft}:{get_column_letter(len(columnas))}{ft + max(len(filas), 1)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{ft}:{ft}"

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=nombre_archivo(f, "xlsx"),
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
