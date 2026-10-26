"""Revisión de precios para comisionistas: revisa si el precio neto facturado de cada línea respetó el precio establecido para el
cliente, para saber si se le paga la comisión (línea debajo del mínimo permitido = no se paga).

- Precio neto unitario = Sub-Total 2 (después de descuento, sin IVA) / cantidad vendida (ventas_por_producto).
- Precios del sistema = Precio 1..4 del catálogo (en U$ se convierten con el tipo de cambio).
- Cada cliente tiene asignada una lista de precio (clientes.precio) y un descuento máximo (clientes.descuento):
  mínimo permitido = precio de su lista × (1 − descuento / 100). Debajo de eso la línea se marca.
- Si la ficha no trae descuento se puede indicar un descuento general permitido (se usa el mayor de los dos).
- Ojo: el catálogo tiene los precios vigentes hoy; para facturas viejas el precio del sistema pudo ser otro."""
import io
from datetime import date, datetime

from flask import Blueprint, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.db import get_connection
from router.accesos import requiere_acceso
from router.dashboard import ultima_fecha_con_ventas
from router.productos_sin_movimiento import TC_DEFECTO
from router.ventas_bodegas import en_cache

revision_precios_venta_bp = Blueprint("revision_precios_venta", __name__)

ESTADOS = {"": "Todas las líneas", "debajo": "Debajo del mínimo permitido", "costo": "Debajo del costo",
           "descuento": "Con descuento permitido", "lista": "Precio de lista o mayor", "sin_precio": "Sin precio en catálogo"}
TIPOS = {"": "Crédito y contado", "CREDITO": "Crédito", "CONTADO": "Contado"}
# Comisión según el resultado: debajo del mínimo no se paga; sin precio en catálogo hay que revisarla a mano
COMISION = {"si": "Sí paga", "no": "No paga", "revisar": "Revisar"}
VERDE, NARANJA, ROJO_CLARO = "027266", "FE6412", "F8C9C4"
POR_PAGINA = 300


def clave(nombre):
    return " ".join((nombre or "").upper().split())


def leer_filtros(args):
    ultima = ultima_fecha_con_ventas() or date.today()
    def fecha(valor, defecto):
        try:
            return datetime.strptime(valor or "", "%Y-%m-%d").date()
        except ValueError:
            return defecto
    hasta = fecha(args.get("hasta"), ultima)
    desde = fecha(args.get("desde"), hasta.replace(day=1))
    if desde > hasta:
        desde, hasta = hasta, desde
    try:
        tc = float(args.get("tc") or TC_DEFECTO)
    except ValueError:
        tc = TC_DEFECTO
    try:
        tolerancia = min(max(float(args.get("tolerancia") or 0.5), 0), 20)
    except ValueError:
        tolerancia = 0.5
    try:   # descuento que se permite aunque la ficha del cliente diga 0 (se usa el mayor de los dos)
        desc_general = min(max(float(args.get("desc_general") or 0), 0), 90)
    except ValueError:
        desc_general = 0.0
    return {"desde": desde, "hasta": hasta, "cliente": (args.get("cliente") or "").strip(),
            "factura": "".join(ch for ch in (args.get("factura") or "") if ch.isdigit()),
            "recibo": "".join(ch for ch in (args.get("recibo") or "") if ch.isdigit()),
            "vendedor": (args.get("vendedor") or "").strip(), "tipo": args.get("tipo") if args.get("tipo") in TIPOS else "",
            "estado": args.get("estado") if args.get("estado") in ESTADOS else "", "tc": tc if tc > 0 else TC_DEFECTO,
            "tolerancia": tolerancia, "desc_general": desc_general}


def calcular(f):
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        # Facturas pagadas con el recibo indicado
        facturas_recibo = None
        if f["recibo"]:
            cursor.execute("SELECT DISTINCT nfactura FROM recuperacion WHERE nrecibo = %s", (f["recibo"],))
            facturas_recibo = [r["nfactura"] for r in cursor.fetchall() if r["nfactura"]]
            if not facturas_recibo:
                return [], {"facturas_recibo": []}

        where, params = ["v.Estatus <> 'Anulado'"], []
        if f["factura"]:
            where.append("v.`No. Factura` = %s"); params.append(int(f["factura"]))
        elif facturas_recibo is not None:
            where.append(f"v.`No. Factura` IN ({', '.join(['%s'] * len(facturas_recibo))})"); params += [int(x) for x in facturas_recibo]
        else:   # sin factura ni recibo: rango de fechas (Fecha es texto 'YYYY-MM-DD')
            where.append("v.Fecha BETWEEN %s AND %s"); params += [str(f["desde"]), str(f["hasta"])]
        if f["tipo"]:
            where.append("v.`Contado/Credito` = %s"); params.append(f["tipo"])
        for palabra in f["cliente"].split():
            where.append("v.Cliente LIKE %s"); params.append(f"%{palabra}%")
        if f["vendedor"]:
            where.append("v.Vendedor = %s"); params.append(f["vendedor"])
        cursor.execute(f"""
            SELECT v.Fecha AS fecha, v.`No. Factura` AS factura, v.`Contado/Credito` AS tipo, TRIM(v.Cliente) AS cliente, v.Vendedor AS vendedor,
                   v.Codigo AS codigo, v.Numero_parte AS parte, v.`Descrip Prod.` AS descripcion, v.Marca AS marca,
                   v.`Cantidad Vendida` AS cantidad, v.`Precio Facturado` AS precio_facturado, v.`Sub-Total` AS subtotal,
                   v.Descuento AS descuento, v.`Sub-Total 2` AS neto, v.`Costo Unit.` AS costo, v.`Tipo Factura` AS tipo_factura
            FROM ventas_por_producto v WHERE {' AND '.join(where)}
            ORDER BY v.Fecha, v.`No. Factura`
        """, params)
        lineas = cursor.fetchall()
        codigos = sorted({x["codigo"] for x in lineas if x["codigo"]})
        catalogo = {}
        for i in range(0, len(codigos), 800):
            bloque = codigos[i:i + 800]
            cursor.execute(f"SELECT `Codigo Producto` AS codigo, Moneda AS moneda, `Precio 1` AS p1, `Precio 2` AS p2, `Precio 3` AS p3, `Precio 4` AS p4 "
                           f"FROM catalogo_productos WHERE `Codigo Producto` IN ({', '.join(['%s'] * len(bloque))})", bloque)
            catalogo.update({r["codigo"]: r for r in cursor.fetchall()})
        cursor.execute("SELECT nombre, estado, precio, descuento FROM clientes")
        fichas = {}
        for r in cursor.fetchall():
            k = clave(r["nombre"])
            if k and (k not in fichas or r["estado"] == "Activo"):
                fichas[k] = r
    finally:
        cursor.close()
        conn.close()

    tol = f["tolerancia"] / 100
    filas = []
    for x in lineas:
        cant = float(x["cantidad"] or 0)
        neto = float(x["neto"] or 0)
        unit = neto / cant if cant else 0.0
        subtotal = float(x["subtotal"] or 0)
        cat = catalogo.get(x["codigo"])
        factor = f["tc"] if cat and (cat["moneda"] or "U$") != "C$" else 1.0
        precios = [float(cat[f"p{i}"] or 0) * factor if cat else 0.0 for i in range(1, 5)]
        ficha = fichas.get(clave(x["cliente"]))
        nivel = int(float(ficha["precio"])) if ficha and ficha["precio"] and 1 <= float(ficha["precio"]) <= 4 else None
        desc_ficha = float(ficha["descuento"] or 0) if ficha else 0.0
        desc_cliente = max(desc_ficha, f["desc_general"])
        lista = precios[nivel - 1] if nivel else (precios[0] or None)   # sin lista asignada: se compara con Precio 1
        minimo = lista * (1 - desc_cliente / 100) if lista else None
        costo = float(x["costo"] or 0)
        if not lista:
            estado = "sin_precio"
        elif unit < minimo * (1 - tol):
            estado = "debajo"
        elif unit < lista * (1 - tol):
            estado = "descuento"
        else:
            estado = "lista"
        bajo_costo = costo > 0 and unit < costo
        filas.append({
            **x, "cantidad": cant, "precio_facturado": float(x["precio_facturado"] or 0), "descuento": float(x["descuento"] or 0),
            "desc_pct": float(x["descuento"] or 0) / subtotal * 100 if subtotal else 0.0, "neto": neto, "unit": unit,
            "precios": [{"cs": p, "dif": (unit / p - 1) * 100 if p else None} for p in precios],
            "nivel": nivel, "desc_cliente": desc_cliente, "desc_ficha": desc_ficha, "en_tabla": bool(ficha), "lista": lista, "minimo": minimo,
            "dif_lista": (unit / lista - 1) * 100 if lista else None, "dif_minimo": (unit / minimo - 1) * 100 if minimo else None,
            # lo que se dejó de cobrar respecto al mínimo permitido
            "faltante": max((minimo - unit) * cant, 0.0) if estado == "debajo" else 0.0,
            "costo": costo, "margen": (unit - costo) / unit * 100 if unit and costo else None, "bajo_costo": bajo_costo, "estado": estado,
            "comision": "no" if estado == "debajo" else ("revisar" if estado == "sin_precio" else "si"),
        })
    if f["estado"] == "costo":
        filas = [x for x in filas if x["bajo_costo"]]
    elif f["estado"]:
        filas = [x for x in filas if x["estado"] == f["estado"]]
    return filas, {"facturas_recibo": facturas_recibo}


def resumen(filas):
    venta = sum(x["neto"] for x in filas)
    debajo = [x for x in filas if x["estado"] == "debajo"]
    vendedores = {}
    for x in debajo:
        v = vendedores.setdefault(x["vendedor"] or "—", {"vendedor": x["vendedor"] or "—", "lineas": 0, "faltante": 0.0, "facturas": set()})
        v["lineas"] += 1
        v["faltante"] += x["faltante"]
        v["facturas"].add(x["factura"])
    por_vendedor = sorted(vendedores.values(), key=lambda v: -v["faltante"])
    for v in por_vendedor:
        v["facturas"] = len(v["facturas"])
    return {"lineas": len(filas), "facturas": len({x["factura"] for x in filas}), "clientes": len({clave(x["cliente"]) for x in filas}),
            "venta": venta, "debajo": len(debajo), "faltante": sum(x["faltante"] for x in debajo),
            "costo": sum(1 for x in filas if x["bajo_costo"]), "descuento": sum(1 for x in filas if x["estado"] == "descuento"),
            "lista": sum(1 for x in filas if x["estado"] == "lista"), "sin_precio": sum(1 for x in filas if x["estado"] == "sin_precio"),
            "sin_ficha": len({clave(x["cliente"]) for x in filas if not x["en_tabla"]}), "por_vendedor": por_vendedor,
            "facturas_sin_comision": len({x["factura"] for x in debajo})}


def lista_vendedores():
    """Vendedores tal como aparecen en ventas_por_producto (el filtro compara contra esa tabla)."""
    def leer():
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT DISTINCT Vendedor FROM ventas_por_producto WHERE Fecha >= %s AND Vendedor <> '' ORDER BY Vendedor",
                           (str(date.today().replace(year=date.today().year - 1)),))
            return [r[0] for r in cursor.fetchall()]
        finally:
            cursor.close()
            conn.close()
    return en_cache("precios_venta_vendedores", leer)


@revision_precios_venta_bp.route("/revision_precios_venta")
@requiere_acceso("revision_precios_venta")
def reporte():
    f = leer_filtros(request.args)
    filas, extra = calcular(f)
    paginas = max(1, -(-len(filas) // POR_PAGINA))
    try:
        pagina = min(max(1, int(request.args.get("pagina", 1))), paginas)
    except ValueError:
        pagina = 1
    return render_template("revision_precios_venta.html", f=f, r=resumen(filas), extra=extra,
                           filas=filas[(pagina - 1) * POR_PAGINA: pagina * POR_PAGINA], pagina=pagina, paginas=paginas,
                           estados=ESTADOS, tipos=TIPOS, comision=COMISION, vendedores=lista_vendedores())


@revision_precios_venta_bp.route("/revision_precios_venta/excel")
@requiere_acceso("revision_precios_venta")
def descargar_excel():
    f = leer_filtros(request.args)
    filas, _ = calcular(f)
    wb = Workbook()
    ws = wb.active
    ws.title = "Precios comisionistas"
    if f["factura"]:
        titulo = f"factura {f['factura']}"
    elif f["recibo"]:
        titulo = f"facturas del recibo {f['recibo']}"
    else:
        titulo = f"del {f['desde']:%d/%m/%Y} al {f['hasta']:%d/%m/%Y}"
    ws["A1"] = f"Revisión de precios para comisionistas · {titulo}"
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    ws["A2"] = (f"Precio unitario neto = Sub-Total 2 / cantidad · Precios del catálogo vigentes (U$ × TC {f['tc']:g}) · "
                f"Mínimo = precio de la lista del cliente − su descuento · Generado {datetime.now():%d/%m/%Y %H:%M}")
    ws["A2"].font = Font(size=9, color="52514E")
    columnas = [("Fecha", 11), ("Contado/Crédito", 10), ("No. Factura", 10), ("Cliente", 30), ("Vendedor", 20), ("Número parte", 16),
                ("Descripción", 32), ("Cantidad", 8), ("Precio facturado", 11), ("Descuento", 10), ("% desc.", 7), ("Sub-Total 2", 12),
                ("Precio unit. venta", 11), ("Prec1 C$", 10), ("% P1", 7), ("Prec2 C$", 10), ("% P2", 7), ("Prec3 C$", 10), ("% P3", 7),
                ("Prec4 C$", 10), ("% P4", 7), ("Lista cliente", 8), ("Precio del cliente C$", 11), ("Desc. cliente %", 8),
                ("Mínimo permitido", 11), ("% vs mínimo", 9), ("Dejado de cobrar C$", 12), ("Costo unit.", 10), ("Margen", 8),
                ("Estado", 22), ("Comisión", 10), ("Estatus factura", 14)]
    ft = 4
    for i, (nombre, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=ft, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NARANJA if i in (13, 23, 25, 26, 27, 30, 31) else VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[ft].height = 32
    rojo = PatternFill("solid", fgColor=ROJO_CLARO)
    for n, x in enumerate(filas, start=ft + 1):
        valores = [datetime.strptime(x["fecha"], "%Y-%m-%d").date() if x["fecha"] else None, x["tipo"], int(x["factura"]) if x["factura"] else None,
                   x["cliente"], x["vendedor"], x["parte"], x["descripcion"], x["cantidad"], x["precio_facturado"], x["descuento"],
                   x["desc_pct"] / 100, x["neto"], round(x["unit"], 2)]
        for p in x["precios"]:
            valores += [round(p["cs"], 2) if p["cs"] else None, p["dif"] / 100 if p["dif"] is not None else None]
        valores += [x["nivel"] or "—", round(x["lista"], 2) if x["lista"] else None, x["desc_cliente"],
                    round(x["minimo"], 2) if x["minimo"] else None,
                    x["dif_minimo"] / 100 if x["dif_minimo"] is not None else None, round(x["faltante"], 2) or None,
                    round(x["costo"], 2), x["margen"] / 100 if x["margen"] is not None else None,
                    ESTADOS[x["estado"]] + (" · bajo costo" if x["bajo_costo"] else ""), COMISION[x["comision"]], x["tipo_factura"]]
        for col, v in enumerate(valores, start=1):
            ws.cell(row=n, column=col, value=v)
        ws.cell(row=n, column=1).number_format = "dd/mm/yyyy"
        for col in (9, 10, 12, 13, 14, 16, 18, 20, 23, 25, 27, 28):
            ws.cell(row=n, column=col).number_format = "#,##0.00"
        for col in (11, 15, 17, 19, 21, 26, 29):
            ws.cell(row=n, column=col).number_format = "0%"
        for k, col in enumerate((15, 17, 19, 21)):
            d = x["precios"][k]["dif"]
            if d is not None and d < -f["tolerancia"]:
                ws.cell(row=n, column=col).fill = rojo
        if x["comision"] == "no":
            ws.cell(row=n, column=31).fill = rojo
            ws.cell(row=n, column=31).font = Font(bold=True, color="B42318")
        if x["estado"] == "debajo" or x["bajo_costo"]:
            for col in (26, 27, 30):
                ws.cell(row=n, column=col).fill = rojo
                ws.cell(row=n, column=col).font = Font(bold=True, color="B42318")
    ws.freeze_panes = ws.cell(row=ft + 1, column=5)
    ws.auto_filter.ref = f"A{ft}:{get_column_letter(len(columnas))}{ft + max(len(filas), 1)}"
    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    sufijo = f"factura_{f['factura']}" if f["factura"] else (f"recibo_{f['recibo']}" if f["recibo"] else f"{f['desde']:%Y%m%d}_{f['hasta']:%Y%m%d}")
    return send_file(salida, as_attachment=True, download_name=f"precios_comisionistas_{sufijo}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
