"""Ventas por día con desglose de impuestos, para cuadrar con el área financiera.

Por día y tipo de venta (contado / crédito):
- Facturado: subtotal, descuento, venta neta (gravada y exenta), IVA débito y total. Sin facturas anuladas.
  Gravada = IVA ÷ 15% de cada factura (tope: su venta neta); exenta = el resto de la venta neta.
- Devoluciones cerradas (por fecha de la devolución): el total de la devolución trae IVA; se separa con la proporción de IVA
  de la factura original (exenta = sin IVA). Si no se encuentra la factura se usa 15%.
- Neto = facturado − devoluciones (base, IVA y total).
Abre y cierra su propia conexión: no usa la conexión global de app.py."""
import io
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP

from flask import Blueprint, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from router.accesos import requiere_acceso
from utils.db import get_connection

ventas_dia_bp = Blueprint("ventas_dia", __name__)

TASA_IVA = Decimal("0.15")
TC_DEFECTO = Decimal("36.6243")
CERO = Decimal(0)
TIPOS = ("CONTADO", "CREDITO")
DIAS = ["lun", "mar", "mié", "jue", "vie", "sáb", "dom"]
VERDE, NARANJA = "027266", "FE6412"
CAMPOS = ("facturas", "subtotal", "descuento", "neto", "gravado", "exento", "iva", "total",
          "devoluciones", "dev_base", "dev_iva", "dev_total", "neto_final", "iva_final", "total_final")


def d2(v):
    return Decimal(v or 0).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def leer_filtros(args):
    hoy = date.today()

    def fecha(nombre, defecto):
        try:
            return datetime.strptime(args.get(nombre) or "", "%Y-%m-%d").date()
        except ValueError:
            return defecto
    hasta = fecha("fecha_fin", None)
    desde = fecha("fecha_inicio", None)
    if not desde and not hasta:   # por defecto: el último día con ventas cargadas
        desde = hasta = ultimo_dia() or hoy
    desde, hasta = desde or hasta, hasta or desde
    if hasta < desde:
        desde, hasta = hasta, desde
    try:
        tc = Decimal(str(args.get("tipo_cambio") or TC_DEFECTO).replace(",", ""))
        tc = tc if tc > 0 else TC_DEFECTO
    except Exception:
        tc = TC_DEFECTO
    return {"desde": desde, "hasta": hasta, "tc": tc}


def ultimo_dia():
    conn = get_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT MAX(Fecha) FROM ventas_historicas")
        v = cursor.fetchone()[0]
        return v if isinstance(v, date) else (datetime.strptime(str(v)[:10], "%Y-%m-%d").date() if v else None)
    finally:
        cursor.close()
        conn.close()


def vacio():
    return {k: (0 if k == "facturas" else CERO) for k in CAMPOS}


def sumar(destino, origen):
    for k in CAMPOS:
        destino[k] += origen[k]


def calcular(f):
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        cursor.execute("""
            SELECT Fecha AS fecha, UPPER(tipo_venta) AS tipo, factura, SubTotal AS subtotal, Descuento AS descuento,
                   neto, IVA AS iva, Total AS total
            FROM ventas_historicas WHERE Fecha BETWEEN %s AND %s AND estado_factura <> 'Anulado'
        """, (f["desde"], f["hasta"]))
        facturas = cursor.fetchall()
        cursor.execute("SELECT COUNT(*) AS n, COALESCE(SUM(Total), 0) AS total FROM ventas_historicas "
                       "WHERE Fecha BETWEEN %s AND %s AND estado_factura = 'Anulado'", (f["desde"], f["hasta"]))
        anuladas = cursor.fetchone()
        # Devoluciones con la proporción de IVA de su factura original
        cursor.execute("""
            SELECT d.fechadv AS fecha, UPPER(d.tipofact) AS tipo, d.numerodv, d.factura, d.totaldv AS total,
                   v.neto AS f_neto, v.iva AS f_iva
            FROM devoluciones d
            LEFT JOIN (SELECT factura, UPPER(tipo_venta) AS tipo, SUM(neto) AS neto, SUM(IVA) AS iva
                       FROM ventas_historicas GROUP BY factura, UPPER(tipo_venta)) v
                   ON v.factura = d.factura AND v.tipo = UPPER(d.tipofact)
            WHERE d.fechadv BETWEEN %s AND %s AND d.estado = 'Cerrado'
        """, (f["desde"], f["hasta"]))
        devoluciones = cursor.fetchall()
    finally:
        cursor.close()
        conn.close()

    dias = defaultdict(lambda: {t: vacio() for t in TIPOS})
    vistas = defaultdict(set)   # una factura puede venir en varias filas (una por bodega): se cuenta una vez
    for x in facturas:
        tipo = x["tipo"] if x["tipo"] in TIPOS else "CONTADO"
        r = dias[x["fecha"]][tipo]
        neto, iva = d2(x["neto"]), d2(x["iva"])
        gravado = min(d2(iva / TASA_IVA), neto) if iva > 0 else CERO
        if x["factura"] not in vistas[(x["fecha"], tipo)]:
            vistas[(x["fecha"], tipo)].add(x["factura"])
            r["facturas"] += 1
        r["subtotal"] += d2(x["subtotal"]); r["descuento"] += d2(x["descuento"]); r["neto"] += neto
        r["gravado"] += gravado; r["exento"] += neto - gravado; r["iva"] += iva; r["total"] += d2(x["total"])
    sin_factura = 0
    for x in devoluciones:
        tipo = x["tipo"] if x["tipo"] in TIPOS else "CONTADO"
        r = dias[x["fecha"]][tipo]
        total = d2(x["total"])
        if x["f_neto"]:
            proporcion = Decimal(x["f_iva"] or 0) / Decimal(x["f_neto"])
        else:
            proporcion, sin_factura = TASA_IVA, sin_factura + 1
        base = d2(total / (1 + proporcion))
        r["devoluciones"] += 1; r["dev_base"] += base; r["dev_iva"] += total - base; r["dev_total"] += total

    filas, totales = [], {t: vacio() for t in TIPOS}
    for dia in sorted(dias):
        fila = {"fecha": dia, "tipos": dias[dia], "total": vacio()}
        for t in TIPOS:
            r = dias[dia][t]
            r["neto_final"] = r["neto"] - r["dev_base"]
            r["iva_final"] = r["iva"] - r["dev_iva"]
            r["total_final"] = r["total"] - r["dev_total"]
            sumar(fila["total"], r)
            sumar(totales[t], r)
        filas.append(fila)
    general = vacio()
    for t in TIPOS:
        sumar(general, totales[t])
    return {"filas": filas, "totales": totales, "general": general, "anuladas": anuladas,
            "dev_sin_factura": sin_factura, "dias_con_venta": len(filas)}


# Renglones del cuadre para Finanzas: (clave, etiqueta, nivel) — nivel 0 = total del bloque
CUADRE = [
    ("titulo", "Facturado", None),
    ("subtotal", "Ventas brutas (subtotal)", 1), ("descuento", "(−) Descuentos", 1),
    ("neto", "Venta neta", 0), ("gravado", "   Gravada 15%", 2), ("exento", "   Exenta", 2),
    ("iva", "IVA débito fiscal", 0), ("total", "Total facturado", 0),
    ("titulo", "Devoluciones (cerradas)", None),
    ("dev_base", "Base de devoluciones", 1), ("dev_iva", "IVA de devoluciones", 1), ("dev_total", "Total devuelto", 0),
    ("titulo", "Neto después de devoluciones", None),
    ("neto_final", "Venta neta final", 0), ("iva_final", "IVA neto", 0), ("total_final", "Total neto", 0),
]


@ventas_dia_bp.route("/reporte_ventas_dia", methods=["GET"])
@requiere_acceso("reporte_ventas_dia")
def reporte():
    f = leer_filtros(request.args)
    d = calcular(f)
    hoy = date.today()
    ultima = ultimo_dia() or hoy
    rangos = [("Último día", ultima, ultima), ("Ayer", ultima - timedelta(days=1), ultima - timedelta(days=1)),
              ("Mes en curso", ultima.replace(day=1), ultima),
              ("Mes anterior", (ultima.replace(day=1) - timedelta(days=1)).replace(day=1), ultima.replace(day=1) - timedelta(days=1))]
    grafico = {"labels": [f"{DIAS[x['fecha'].weekday()]} {x['fecha']:%d/%m}" for x in d["filas"]],
               "contado": [float(x["tipos"]["CONTADO"]["neto_final"]) for x in d["filas"]],
               "credito": [float(x["tipos"]["CREDITO"]["neto_final"]) for x in d["filas"]]}
    return render_template("reporte_ventas_dia.html", f=f, d=d, cuadre=CUADRE, tipos=TIPOS, dias=DIAS,
                           rangos=rangos, grafico=grafico, tasa=int(TASA_IVA * 100))


@ventas_dia_bp.route("/reporte_ventas_dia/excel")
@requiere_acceso("reporte_ventas_dia")
def excel():
    f = leer_filtros(request.args)
    d = calcular(f)
    wb = Workbook()
    fino = Side(style="thin", color="D9D6CE")
    titulo = f"Ventas del {f['desde']:%d/%m/%Y}" + (f" al {f['hasta']:%d/%m/%Y}" if f["hasta"] != f["desde"] else "")

    # Hoja 1: cuadre
    ws = wb.active
    ws.title = "Cuadre"
    ws["A1"] = f"Cuadre de ventas e IVA · {titulo}"
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    ws["A2"] = (f"Sin facturas anuladas ({d['anuladas']['n']} por C$ {d['anuladas']['total']:,.2f}). Devoluciones cerradas por fecha de "
                f"devolución; su IVA según la factura original. Generado {datetime.now():%d/%m/%Y %H:%M}")
    ws["A2"].font = Font(size=9, color="52514E")
    for i, t in enumerate(["Concepto", "Contado C$", "Crédito C$", "Total C$"], start=1):
        c = ws.cell(row=4, column=i, value=t)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=VERDE)
        c.alignment = Alignment(horizontal="center")
    n = 5
    for clave, etiqueta, nivel in CUADRE:
        if clave == "titulo":
            ws.cell(row=n, column=1, value=etiqueta).font = Font(bold=True, color=NARANJA)
        else:
            ws.cell(row=n, column=1, value=etiqueta.strip()).alignment = Alignment(indent=2 if nivel == 2 else 1 if nivel == 1 else 0)
            for col, v in enumerate([d["totales"]["CONTADO"][clave], d["totales"]["CREDITO"][clave], d["general"][clave]], start=2):
                c = ws.cell(row=n, column=col, value=float(v))
                c.number_format = "#,##0.00"
                if nivel == 0:
                    c.font = Font(bold=True)
            if nivel == 0:
                ws.cell(row=n, column=1).font = Font(bold=True)
                for col in range(1, 5):
                    ws.cell(row=n, column=col).border = Border(top=fino)
        n += 1
    ws.cell(row=n + 1, column=1, value="Facturas")
    for col, t in enumerate(["CONTADO", "CREDITO"], start=2):
        ws.cell(row=n + 1, column=col, value=d["totales"][t]["facturas"])
    ws.cell(row=n + 1, column=4, value=d["general"]["facturas"])
    ws.cell(row=n + 2, column=1, value="Tipo de cambio")
    ws.cell(row=n + 2, column=4, value=float(f["tc"])).number_format = "0.0000"
    ws.cell(row=n + 3, column=1, value="Venta neta final US$")
    ws.cell(row=n + 3, column=4, value=float(d["general"]["neto_final"] / f["tc"])).number_format = "#,##0.00"
    for col, ancho in zip("ABCD", (34, 16, 16, 16)):
        ws.column_dimensions[col].width = ancho

    # Hoja 2: por día
    ws2 = wb.create_sheet("Por día")
    columnas = [("Fecha", 11), ("Tipo", 10), ("Facturas", 9), ("Subtotal", 14), ("Descuento", 12), ("Venta neta", 14),
                ("Gravada", 14), ("Exenta", 12), ("IVA", 13), ("Total facturado", 14), ("Devol. #", 8), ("Devol. base", 13),
                ("Devol. IVA", 12), ("Devol. total", 13), ("Venta neta final", 15), ("IVA neto", 13), ("Total neto", 15), ("Neto US$", 12)]
    for i, (t, ancho) in enumerate(columnas, start=1):
        c = ws2.cell(row=1, column=i, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NARANJA if t.startswith("Devol") else VERDE)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws2.column_dimensions[get_column_letter(i)].width = ancho
    n = 2

    def fila(fecha, tipo, r, negrita=False):
        valores = [fecha, tipo, r["facturas"]] + [float(r[k]) for k in ("subtotal", "descuento", "neto", "gravado", "exento", "iva", "total")] \
            + [r["devoluciones"]] + [float(r[k]) for k in ("dev_base", "dev_iva", "dev_total", "neto_final", "iva_final", "total_final")] \
            + [float(r["neto_final"] / f["tc"])]
        for col, v in enumerate(valores, start=1):
            c = ws2.cell(row=n, column=col, value=v)
            if col == 1 and isinstance(v, date):
                c.number_format = "dd/mm/yyyy"
            elif col > 3 and col != 11:
                c.number_format = "#,##0.00"
            if negrita:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="F1F0EC")
    for x in d["filas"]:
        for t in TIPOS:
            if x["tipos"][t]["facturas"] or x["tipos"][t]["devoluciones"]:
                fila(x["fecha"], t.capitalize(), x["tipos"][t]); n += 1
        fila(x["fecha"], "Total día", x["total"], True); n += 1
    fila("TOTAL", "", d["general"], True)
    ws2.freeze_panes = "C2"

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    nombre = f"ventas_impuestos_{f['desde']:%Y%m%d}" + (f"_{f['hasta']:%Y%m%d}" if f["hasta"] != f["desde"] else "") + ".xlsx"
    return send_file(salida, as_attachment=True, download_name=nombre,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
