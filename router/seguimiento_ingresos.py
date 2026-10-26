"""Seguimiento de ingresos de mercadería: de lo que entró (órdenes de compra), cuánto ya se vendió, a quién y con qué vendedor.

- Ingresos: transacciones_inventario con ENTRADA / ORDEN COMPRA, agrupadas por fecha + documento + producto.
- Proveedor: el de la transacción; desde 2023 el sistema no lo trae, entonces se usa "Proveedor Ultima Compra" del catálogo.
- Ventas asignadas por PEPS (FIFO): primero se vende la existencia que ya había antes del primer ingreso del período,
  luego cada ingreso en orden de fecha. Las devoluciones de clientes se restan de las ventas del mismo día.
- El detalle (clic en una fila) trae cliente, factura, vendedor y monto desde ventas_por_producto."""
import io
from collections import defaultdict
from datetime import date, datetime, timedelta

from flask import Blueprint, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.db import get_connection
from router.accesos import requiere_acceso
from router.productos_agotarse import lista_marcas
from router.ventas_bodegas import en_cache

seguimiento_ingresos_bp = Blueprint("seguimiento_ingresos", __name__)

ESTADOS = {"": "Todos", "sin_vender": "Sin vender", "parcial": "Vendido en parte", "vendido": "Vendido todo"}
ORDENES = {"fecha": "Ingreso más reciente", "pendiente": "Mayor valor sin vender", "avance": "Menor % vendido"}
VERDE, NARANJA = "027266", "FE6412"
POR_PAGINA = 200
HINT = "/*+ INDEX(transacciones_inventario idx_transacciones_producto_fecha) */"


def leer_filtros(args):
    hoy = date.today()
    try:
        desde = datetime.strptime(args.get("desde", ""), "%Y-%m-%d").date()
    except ValueError:
        desde = hoy - timedelta(days=90)
    try:
        hasta = datetime.strptime(args.get("hasta", ""), "%Y-%m-%d").date()
    except ValueError:
        hasta = hoy
    if hasta < desde:
        desde, hasta = hasta, desde
    return {"desde": desde, "hasta": hasta, "marca": (args.get("marca") or "").strip(), "proveedor": (args.get("proveedor") or "").strip(),
            "q": (args.get("q") or "").strip(), "estado": args.get("estado") if args.get("estado") in ESTADOS else "",
            "orden": args.get("orden") if args.get("orden") in ORDENES else "fecha"}


def lista_proveedores():
    def leer():
        conn = get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT DISTINCT cliente_proveedor FROM transacciones_inventario WHERE tipo_documento_origen = 'ORDEN COMPRA' "
                           "AND cliente_proveedor IS NOT NULL AND cliente_proveedor <> ''")
            nombres = {r[0].strip() for r in cursor.fetchall()}
            cursor.execute("SELECT DISTINCT `Proveedor Ultima Compra` FROM catalogo_productos WHERE `Proveedor Ultima Compra` <> ''")
            nombres |= {r[0].strip() for r in cursor.fetchall() if r[0]}
            return sorted(nombres)
        finally:
            cursor.close()
            conn.close()
    return en_cache("seguimiento_ingresos_proveedores", leer)


def en_lista(cursor, sql, codigos, extra=()):
    """Ejecuta sql con un IN (%s, …) de códigos, por bloques."""
    filas = []
    codigos = list(codigos)
    for i in range(0, len(codigos), 800):
        bloque = codigos[i:i + 800]
        cursor.execute(sql.replace("{IN}", ", ".join(["%s"] * len(bloque))), tuple(bloque) + tuple(extra))
        filas += cursor.fetchall()
    return filas


def asignar_peps(lotes, previo, ventas):
    """Reparte las ventas (fecha, cantidad, venta) entre la existencia previa y los lotes (ingresos) en orden de fecha.
    Una venta solo puede tomar de lotes que ya habían ingresado ese día. Devuelve {indice_lote: [(venta, cantidad)]}."""
    asignado = defaultdict(list)
    restante_previo = previo
    pendiente = [l["cantidad"] for l in lotes]
    for v in ventas:
        cant = v["cantidad"]
        if restante_previo > 0:
            toma = min(cant, restante_previo)
            restante_previo -= toma
            cant -= toma
        for i, l in enumerate(lotes):
            if cant <= 0:
                break
            if l["fecha"] > v["fecha"] or pendiente[i] <= 0:
                continue
            toma = min(cant, pendiente[i])
            pendiente[i] -= toma
            cant -= toma
            asignado[i].append((v, toma))
    return asignado


def calcular(f, solo_codigo=None):
    """Devuelve las líneas de ingreso del período con lo vendido de cada una (y la lista de ventas asignadas)."""
    conn = get_connection()
    cursor = conn.cursor(dictionary=True)
    try:
        # 1) Ingresos por orden de compra en el período
        sql = ("SELECT fecha, id_documento_origen AS documento, MAX(referencia) AS referencia, MAX(cliente_proveedor) AS proveedor, "
               "codigo_producto AS codigo, MAX(producto) AS producto, MAX(marca) AS marca, SUM(cantidad) AS cantidad, "
               "SUM(total_mn) AS costo_total, GROUP_CONCAT(DISTINCT bodega SEPARATOR ', ') AS bodegas "
               "FROM transacciones_inventario WHERE tipo_transaccion = 'ENTRADA' AND tipo_documento_origen = 'ORDEN COMPRA' "
               "AND fecha BETWEEN %s AND %s" + (" AND codigo_producto = %s" if solo_codigo else "") +
               " GROUP BY fecha, id_documento_origen, codigo_producto HAVING SUM(cantidad) > 0")
        cursor.execute(sql, (f["desde"], f["hasta"]) + ((solo_codigo,) if solo_codigo else ()))
        ingresos = cursor.fetchall()
        if not ingresos:
            return []
        codigos = {x["codigo"] for x in ingresos}

        # 2) Catálogo y existencia actual
        catalogo = {r["codigo"]: r for r in en_lista(cursor, "SELECT `Codigo Producto` AS codigo, Nombre, `Numero Parte` AS parte, Marca, "
                                                            "`Proveedor Ultima Compra` AS proveedor FROM catalogo_productos "
                                                            "WHERE `Codigo Producto` IN ({IN})", codigos)}
        existencia = {r["codigo"]: float(r["e"] or 0) for r in en_lista(cursor, "SELECT Codigo_producto AS codigo, SUM(Cant_Final) AS e "
                                                                              "FROM existencias_bodega WHERE Codigo_producto IN ({IN}) "
                                                                              "GROUP BY Codigo_producto", codigos)}
        for x in ingresos:
            cat = catalogo.get(x["codigo"], {})
            x["cantidad"] = float(x["cantidad"])
            x["costo_total"] = float(x["costo_total"] or 0)
            x["costo_unit"] = x["costo_total"] / x["cantidad"] if x["cantidad"] else 0
            x["nombre"] = cat.get("Nombre") or x["producto"] or ""
            x["parte"] = cat.get("parte") or ""
            x["marca"] = (x["marca"] or cat.get("Marca") or "").strip()
            x["proveedor_origen"] = "ingreso" if (x["proveedor"] or "").strip() else ("catálogo" if cat.get("proveedor") else "")
            x["proveedor"] = ((x["proveedor"] or "").strip() or (cat.get("proveedor") or "").strip())
            x["existencia"] = existencia.get(x["codigo"], 0.0)

        # 3) Filtros de marca, proveedor y búsqueda (antes de leer movimientos)
        def pasa(x):
            # Marca exacta o la familia completa: "SKF" incluye SKF-GENA, SKF-USA, SKF-BANDAS…
            m = f["marca"].upper()
            if m and x["marca"].upper() != m and not x["marca"].upper().startswith(m + "-"):
                return False
            if f["proveedor"] and f["proveedor"].upper() not in x["proveedor"].upper():
                return False
            if f["q"]:
                q = f["q"].upper()
                return q in x["codigo"].upper() or q in x["nombre"].upper() or q in x["parte"].upper() or q in str(x["documento"]).upper() \
                    or q in (x["referencia"] or "").upper()
            return True
        ingresos = [x for x in ingresos if pasa(x)]
        if not ingresos:
            return []
        codigos = {x["codigo"] for x in ingresos}
        inicio = min(x["fecha"] for x in ingresos)

        # 4) Movimientos desde el primer ingreso: para la existencia previa, las ventas y los ingresos posteriores al período
        movs = en_lista(cursor, f"SELECT {HINT} codigo_producto AS codigo, fecha, tipo_transaccion AS tipo, tipo_documento_origen AS doc, "
                                "tipo_operacion AS signo, id_documento_origen AS documento, cliente_proveedor AS cliente, SUM(cantidad) AS cantidad "
                                "FROM transacciones_inventario WHERE codigo_producto IN ({IN}) AND fecha >= %s "
                                "GROUP BY codigo_producto, fecha, tipo_transaccion, tipo_documento_origen, tipo_operacion, id_documento_origen, "
                                "cliente_proveedor", codigos, (inicio,))
    finally:
        cursor.close()
        conn.close()

    por_codigo = defaultdict(list)
    for m in movs:
        por_codigo[m["codigo"]].append(m)
    lineas_por_codigo = defaultdict(list)
    for x in ingresos:
        lineas_por_codigo[x["codigo"]].append(x)

    for codigo, lineas in lineas_por_codigo.items():
        # Desde el primer ingreso de este producto en el período (igual en la tabla y en el detalle)
        primero = min(x["fecha"] for x in lineas)
        ms = [m for m in por_codigo.get(codigo, []) if m["fecha"] >= primero]
        neto = sum(float(m["cantidad"]) * (1 if m["signo"] == "+" else -1) for m in ms)
        previo = max(0.0, lineas[0]["existencia"] - neto)   # existencia que había antes de ese primer ingreso
        # Lotes: los ingresos del período y también los posteriores (para que las ventas no se asignen de más)
        por_lote = defaultdict(float)
        for m in ms:
            if m["tipo"] == "ENTRADA" and m["doc"] == "ORDEN COMPRA":
                por_lote[(m["fecha"], str(m["documento"]))] += float(m["cantidad"])
        lotes = [{"fecha": k[0], "documento": k[1], "cantidad": c} for k, c in sorted(por_lote.items()) if c > 0]
        # Ventas: salidas por factura menos devoluciones del mismo día
        devol = defaultdict(float)
        for m in ms:
            if m["tipo"] == "ENTRADA" and m["doc"] == "DEVOLUCION":
                devol[m["fecha"]] += float(m["cantidad"])
        ventas = []
        for m in sorted((m for m in ms if m["tipo"] == "SALIDA"), key=lambda m: (m["fecha"], str(m["documento"]))):
            cant = float(m["cantidad"])
            quita = min(cant, devol[m["fecha"]])
            devol[m["fecha"]] -= quita
            if cant - quita > 0:
                ventas.append({"fecha": m["fecha"], "cantidad": cant - quita, "factura": m["documento"], "cliente": (m["cliente"] or "").strip()})
        asignado = asignar_peps(lotes, previo, ventas)
        for x in lineas:
            # Lote que corresponde a esta línea (mismo día y documento)
            i = next((k for k, l in enumerate(lotes) if l["fecha"] == x["fecha"] and l["documento"] == str(x["documento"])), None)
            vendidas = asignado.get(i, []) if i is not None else []
            x["ventas"] = [{"fecha": v["fecha"], "factura": v["factura"], "cliente": v["cliente"], "cantidad": c} for v, c in vendidas]
            x["vendido"] = min(x["cantidad"], sum(c for _, c in vendidas))
            x["pendiente"] = x["cantidad"] - x["vendido"]
            x["avance"] = x["vendido"] / x["cantidad"] if x["cantidad"] else 0
            x["estado"] = "vendido" if x["pendiente"] <= 0.0001 else ("parcial" if x["vendido"] > 0 else "sin_vender")
            x["clientes"] = len({v["cliente"].upper() for v in x["ventas"]})
            x["primera_venta"] = x["ventas"][0]["fecha"] if x["ventas"] else None
            x["ultima_venta"] = x["ventas"][-1]["fecha"] if x["ventas"] else None
            x["dias_venta"] = (x["ultima_venta"] - x["fecha"]).days if x["estado"] == "vendido" else None
            x["previo"] = previo
            x["valor_pendiente"] = x["pendiente"] * x["costo_unit"]

    if f["estado"]:
        ingresos = [x for x in ingresos if x["estado"] == f["estado"]]
    if f["orden"] == "pendiente":
        ingresos.sort(key=lambda x: -x["valor_pendiente"])
    elif f["orden"] == "avance":
        ingresos.sort(key=lambda x: (x["avance"], -x["valor_pendiente"]))
    else:
        ingresos.sort(key=lambda x: (x["fecha"], str(x["documento"]), x["codigo"]), reverse=True)
    return ingresos


def resumen(filas):
    unidades = sum(x["cantidad"] for x in filas)
    vendidas = sum(x["vendido"] for x in filas)
    costo = sum(x["costo_total"] for x in filas)
    costo_vendido = sum(x["vendido"] * x["costo_unit"] for x in filas)
    return {"lineas": len(filas), "documentos": len({(x["fecha"], str(x["documento"])) for x in filas}),
            "productos": len({x["codigo"] for x in filas}), "unidades": unidades, "vendidas": vendidas,
            "avance": vendidas / unidades if unidades else 0, "costo": costo, "costo_vendido": costo_vendido,
            "avance_costo": costo_vendido / costo if costo else 0,
            "clientes": len({v["cliente"].upper() for x in filas for v in x["ventas"]}),
            "estados": {k: sum(1 for x in filas if x["estado"] == k) for k in ("sin_vender", "parcial", "vendido")}}


@seguimiento_ingresos_bp.route("/seguimiento_ingresos")
@requiere_acceso("seguimiento_ingresos")
def reporte():
    f = leer_filtros(request.args)
    filas = calcular(f)
    paginas = max(1, -(-len(filas) // POR_PAGINA))
    try:
        pagina = min(max(1, int(request.args.get("pagina", 1))), paginas)
    except ValueError:
        pagina = 1
    return render_template("seguimiento_ingresos.html", f=f, r=resumen(filas), filas=filas[(pagina - 1) * POR_PAGINA: pagina * POR_PAGINA],
                           pagina=pagina, paginas=paginas, marcas=lista_marcas(), proveedores=lista_proveedores(),
                           estados=ESTADOS, ordenes=ORDENES)


@seguimiento_ingresos_bp.route("/seguimiento_ingresos/detalle")
@requiere_acceso("seguimiento_ingresos")
def detalle():
    """Ventas asignadas a un ingreso (fecha + documento + producto), con vendedor y monto de la factura."""
    f = leer_filtros(request.args)
    codigo, documento, fecha = request.args.get("codigo", ""), request.args.get("documento", ""), request.args.get("fecha", "")
    linea = next((x for x in calcular({**f, "marca": "", "proveedor": "", "q": "", "estado": ""}, solo_codigo=codigo)
                  if str(x["documento"]) == documento and x["fecha"].isoformat() == fecha), None)
    if not linea:
        return jsonify(error="No se encontró el ingreso."), 404
    # Vendedor y monto desde ventas_por_producto (por factura + código)
    facturas = sorted({int(v["factura"]) for v in linea["ventas"] if str(v["factura"]).isdigit()})
    info = {}
    if facturas:
        conn = get_connection()
        cursor = conn.cursor(dictionary=True)
        try:
            for i in range(0, len(facturas), 800):
                bloque = facturas[i:i + 800]
                cursor.execute("SELECT `No. Factura` AS factura, MAX(Vendedor) AS vendedor, MAX(Cliente) AS cliente, SUM(`Cantidad Vendida`) AS cantidad, "
                               "SUM(`Sub-Total 2`) AS monto FROM ventas_por_producto WHERE Codigo = %s AND `No. Factura` IN ("
                               + ", ".join(["%s"] * len(bloque)) + ") GROUP BY `No. Factura`", (codigo, *bloque))
                for r in cursor.fetchall():
                    info[int(r["factura"])] = r
        finally:
            cursor.close()
            conn.close()
    ventas = []
    for v in linea["ventas"]:
        d = info.get(int(v["factura"])) if str(v["factura"]).isdigit() else None
        cant_factura = float(d["cantidad"] or 0) if d else 0
        ventas.append({"fecha": v["fecha"].strftime("%d/%m/%Y"), "factura": v["factura"],
                       "cliente": (d["cliente"] if d and d["cliente"] else v["cliente"]) or "—",
                       "vendedor": (d["vendedor"] if d else None) or "—", "cantidad": v["cantidad"],
                       # monto proporcional a las unidades que salieron de este ingreso
                       "monto": round(float(d["monto"] or 0) * v["cantidad"] / cant_factura, 2) if d and cant_factura else None})
    return jsonify(producto=f"{linea['codigo']} · {linea['nombre']}", marca=linea["marca"], proveedor=linea["proveedor"] or "—",
                   documento=str(linea["documento"]), referencia=linea["referencia"] or "", fecha=linea["fecha"].strftime("%d/%m/%Y"),
                   cantidad=linea["cantidad"], vendido=linea["vendido"], pendiente=linea["pendiente"], previo=linea["previo"],
                   existencia=linea["existencia"], ventas=ventas)


@seguimiento_ingresos_bp.route("/seguimiento_ingresos/excel")
@requiere_acceso("seguimiento_ingresos")
def descargar_excel():
    f = leer_filtros(request.args)
    filas = calcular(f)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ingresos"
    ws["A1"] = f"Seguimiento de ingresos de mercadería del {f['desde']:%d/%m/%Y} al {f['hasta']:%d/%m/%Y}"
    ws["A1"].font = Font(bold=True, size=14, color=VERDE)
    ws["A2"] = ("Ventas asignadas por PEPS: primero la existencia que ya había y luego cada ingreso en orden de fecha · "
                f"Filtros: marca {f['marca'] or 'todas'}, proveedor {f['proveedor'] or 'todos'} · Generado {datetime.now():%d/%m/%Y %H:%M}")
    ws["A2"].font = Font(size=9, color="52514E")
    columnas = [("Fecha ingreso", 12), ("Documento", 11), ("Referencia", 24), ("Proveedor", 28), ("Código", 11), ("Nº parte", 18),
                ("Descripción", 34), ("Marca", 14), ("Cant. ingresada", 11), ("Costo unit. C$", 12), ("Costo total C$", 14),
                ("Vendido", 10), ("% vendido", 10), ("Sin vender", 10), ("Costo sin vender C$", 14), ("Clientes", 9),
                ("Primera venta", 12), ("Última venta", 12), ("Días en venderse", 10), ("Existencia actual", 11)]
    ft = 4
    for i, (nombre, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=ft, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NARANJA if i in (12, 13, 14, 15) else VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[ft].height = 32
    for n, x in enumerate(filas, start=ft + 1):
        valores = [x["fecha"], x["documento"], x["referencia"], x["proveedor"], x["codigo"], x["parte"], x["nombre"], x["marca"],
                   x["cantidad"], round(x["costo_unit"], 2), round(x["costo_total"], 2), x["vendido"], f"=IF(I{n}>0,L{n}/I{n},0)",
                   f"=I{n}-L{n}", f"=N{n}*J{n}", x["clientes"], x["primera_venta"], x["ultima_venta"], x["dias_venta"], x["existencia"]]
        for col, v in enumerate(valores, start=1):
            ws.cell(row=n, column=col, value=v)
        for col in (1, 17, 18):
            ws.cell(row=n, column=col).number_format = "dd/mm/yyyy"
        for col in (10, 11, 15):
            ws.cell(row=n, column=col).number_format = "#,##0.00"
        ws.cell(row=n, column=13).number_format = "0%"
    if filas:
        tot = ft + len(filas) + 1
        ws.cell(row=tot, column=1, value="TOTAL").font = Font(bold=True)
        for col, fmt in ((9, "#,##0"), (11, "#,##0.00"), (12, "#,##0"), (14, "#,##0"), (15, "#,##0.00")):
            letra = get_column_letter(col)
            c = ws.cell(row=tot, column=col, value=f"=SUM({letra}{ft + 1}:{letra}{tot - 1})")
            c.font = Font(bold=True)
            c.number_format = fmt
        c = ws.cell(row=tot, column=13, value=f"=IF(I{tot}>0,L{tot}/I{tot},0)")
        c.font = Font(bold=True)
        c.number_format = "0%"
    ws.freeze_panes = ws.cell(row=ft + 1, column=6)
    ws.auto_filter.ref = f"A{ft}:{get_column_letter(len(columnas))}{ft + max(len(filas), 1)}"

    # Hoja con el detalle de a quién se vendió cada ingreso (sin vendedor: sale de transacciones)
    wd = wb.create_sheet("Ventas por ingreso")
    enc = [("Fecha ingreso", 12), ("Documento", 11), ("Código", 11), ("Descripción", 34), ("Fecha venta", 12), ("Factura", 10),
           ("Cliente", 36), ("Cantidad", 10)]
    for i, (nombre, ancho) in enumerate(enc, start=1):
        c = wd.cell(row=1, column=i, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VERDE)
        wd.column_dimensions[get_column_letter(i)].width = ancho
    n = 2
    for x in filas:
        for v in x["ventas"]:
            for col, val in enumerate([x["fecha"], x["documento"], x["codigo"], x["nombre"], v["fecha"], v["factura"], v["cliente"], v["cantidad"]], start=1):
                wd.cell(row=n, column=col, value=val)
            wd.cell(row=n, column=1).number_format = wd.cell(row=n, column=5).number_format = "dd/mm/yyyy"
            n += 1
    wd.freeze_panes = "A2"
    wd.auto_filter.ref = f"A1:H{max(n - 1, 1)}"

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=f"seguimiento_ingresos_{f['desde']:%Y%m%d}_{f['hasta']:%Y%m%d}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
